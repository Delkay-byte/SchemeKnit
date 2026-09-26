"""
Phase 16.5 — Real-World Content Correctness + AI Mode Wiring.

Regression tests for the five defects found during real teacher-style use:

  1. A COMPLETE multi-subject scheme must not collapse to 0 lessons when
     usable subject material is present (and special/revision/SBA weeks must
     not be mistaken for a parser failure).
  2. Changing AI mode must actually change backend behaviour, and the reported
     mode/provider must match reality.
  3. Lessons for different indicators must have substantively different phase
     content (no cloned instructional content).
  4. Teacher-entered keywords must survive the whole generation path.
  5. Keywords / TLRs / core competencies / assessment must be derived from the
     lesson's own context, not blindly cloned from one template.
"""

import asyncio
import os
import sys
from datetime import date

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from src.models import (
    Week, WeekType, SchemeOfWork, TermConfig, Subject, ClassLevel, AIMode,
)
from src.engines.generation_pipeline import GenerationPipeline
from src.engines.ai_provider import AIProvider


# ── Shared fixtures ─────────────────────────────────────────────────────────

def _config(subject, class_level=ClassLevel.BASIC_8, duration=60, ai_mode=AIMode.OFF,
            keywords=None):
    return TermConfig(
        scheme_of_work_id="s", academic_year="2026/2027", term="First Term",
        class_level=class_level, subject=subject,
        term_start_date=date(2026, 9, 7), term_end_date=date(2026, 12, 18),
        lessons_per_week=3, lesson_duration_minutes=duration,
        teaching_days=[0, 1, 2], holidays=[], ai_mode=ai_mode,
        keywords=list(keywords or []),
        school_name="Test School", teacher_name="Test Teacher",
    )


def _weeks(codes, strand, sub_strand, content_standard="B8.1.1.1 Standard",
           per_week=3):
    weeks = []
    counter = 1
    for code in codes:
        wn = (counter - 1) // per_week + 1
        if not weeks or weeks[-1].week_number != wn:
            weeks.append(Week(
                week_number=wn, start_date=date(2026, 9, 7),
                end_date=date(2026, 9, 11), week_type=WeekType.INSTRUCTION,
                strand=strand, sub_strand=sub_strand,
                content_standards=[content_standard], indicators=[],
                resources=[], scheme_of_work_id="s",
            ))
        weeks[-1].indicators.append(code)
        counter += 1
    return weeks


def _scheme(subject, weeks, class_level=ClassLevel.BASIC_8):
    return SchemeOfWork(
        id="s", filename="x.docx", class_level=class_level, subject=subject,
        term="First Term", academic_year="2026/2027", weeks=weeks,
        upload_date=date.today(),
    )


def _run(subject, codes, strand, sub_strand, **kw):
    weeks = _weeks(codes, strand, sub_strand)
    scheme = _scheme(subject, weeks, class_level=kw.pop("class_level", ClassLevel.BASIC_8))
    config = _config(subject, **kw)
    job = GenerationPipeline().generate_all(scheme, config)
    return sorted(job._lesson_plans, key=lambda lp: lp.lesson_sequence), job


# ── Defect 1 — complete multi-subject scheme → usable lessons ───────────────

HEADER = ["Week Ending", "Strand", "Sub-Strand", "Content Standard", "Indicators"]


def _multi_subject_docx(path, subjects, weeks_per_subject=4):
    from docx import Document

    doc = Document()
    doc.add_paragraph("BASIC 8 SCHEME OF LEARNING")
    for subject, prefix in subjects:
        doc.add_paragraph(subject.upper())
        rows = [HEADER]
        for n in range(1, weeks_per_subject + 1):
            rows.append([
                f"{n}\n{10 + n}/09/2026", "Diversity of Matter", "Elements",
                f"{prefix}.1.1.{n} Content standard",
                f"{prefix}.1.1.{n}.1 {prefix} indicator number {n}",
            ])
        table = doc.add_table(rows=len(rows), cols=len(HEADER))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    doc.save(str(path))
    return path


class TestMultiSubjectFullDocument:
    def test_full_multi_subject_docx_yields_usable_lessons(self, tmp_path):
        from src.parsers.docx_parser import DOCXParser
        from src.engines.allocation_engine import AllocationEngine
        from src.engines.calendar_engine import CalendarEngine

        path = _multi_subject_docx(
            tmp_path / "basic8_full.docx",
            [("English Language", "B8.1"), ("Mathematics", "B8.2"),
             ("Science", "B8.3")],
        )
        info = DOCXParser().analyze(path, original_filename="basic8_full.docx")
        assert info["detection_status"] == "multiple"
        assert len(info["detected_subjects"]) >= 2
        # Every section carries usable material.
        assert all(s["week_count"] > 0 for s in info["sections"])

        # Confirm ONE subject — never silently pick or mix.
        scheme = asyncio.run(
            DOCXParser().parse(path, target_subject="Science")
        )
        assert scheme.subject == Subject.SCIENCE
        assert len(scheme.weeks) > 0

        config = _config(Subject.SCIENCE)
        calendar = CalendarEngine().build_calendar(config, scheme.weeks, config.holidays)
        coverage = AllocationEngine().allocate(scheme.weeks, calendar, config)
        assert coverage.total_generated_lessons > 0
        assert all(a.indicator_code for a in coverage.allocations)

    def test_special_week_does_not_become_extraction_failure(self, tmp_path):
        """Revision/SBA weeks carry no instructional indicators; they must not
        be reported as a parser failure."""
        from docx import Document
        from src.parsers.docx_parser import DOCXParser

        doc = Document()
        doc.add_paragraph("BASIC 8 SCHEME OF LEARNING")
        doc.add_paragraph("SCIENCE")
        rows = [HEADER]
        for n in (1, 2):
            rows.append([
                f"{n}\n{10 + n}/09/2026", "Diversity of Matter", "Elements",
                f"B8.3.1.1.{n} Content standard",
                f"B8.3.1.1.{n} indicator number {n}",
            ])
        rows.append(["3\n30/09/2026", "REVISION", "", "", ""])
        table = doc.add_table(rows=len(rows), cols=len(HEADER))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        path = tmp_path / "with_revision.docx"
        doc.save(str(path))

        info = DOCXParser().analyze(path, original_filename="with_revision.docx")
        assert info["detection_status"] != "extraction_failed"
        scheme = asyncio.run(DOCXParser().parse(path, original_filename="with_revision.docx"))
        assert len(scheme.weeks) >= 2
        assert any(w.week_type != WeekType.INSTRUCTION for w in scheme.weeks)

        # Only instructional indicators are allocated by default.
        _, job = _run(Subject.SCIENCE,
                      ["B8.3.1.1.1 A", "B8.3.1.1.2 B"], "Diversity", "Elements")
        assert job.status.value == "completed"
        assert job.total_lessons == 2


# ── Defect 3 / 5 — indicator-specific deterministic content ─────────────────

TWELVE = [
    "B8.1.1.1.1 Solve word problems involving addition and subtraction",
    "B8.1.1.1.2 Identify the place value of digits in whole numbers",
    "B8.1.1.1.3 Compare and order whole numbers up to 10000",
    "B8.2.1.1.1 Investigate the effect of heat on different materials",
    "B8.2.1.1.2 Classify materials into solids, liquids and gases",
    "B8.2.1.1.3 Describe the characteristics of living things",
    "B8.3.1.1.1 Read and discuss a passage about the environment",
    "B8.3.1.1.2 Write a short paragraph using adjectives correctly",
    "B8.3.1.1.3 Create a simple pattern using local materials",
    "B8.4.1.1.1 Discuss the roles of members of the family",
    "B8.4.1.1.2 Compare the culture of two Ghanaian communities",
    "B8.5.1.1.1 Use a chart to analyse data on rainfall",
]


def _tokens(text):
    import re
    return {w for w in re.findall(r"[a-z]{4,}", (text or "").lower())}


class TestIndicatorSpecificContent:
    def test_twelve_indicators_produce_different_phase_content(self):
        plans, job = _run(Subject.SCIENCE, TWELVE, "Strand", "Sub-Strand")
        assert len(plans) == 12

        signatures = {
            (
                lp.starter_activity,
                tuple(a.description for a in lp.main_activities),
                lp.assessment,
                lp.conclusion,
            )
            for lp in plans
        }
        # Not a single cloned lesson repeated twelve times.
        assert len(signatures) >= 5, (
            f"only {len(signatures)} distinct lesson structures across 12 indicators"
        )

        # Every lesson must contain its own indicator's distinctive terminology.
        for lp in plans:
            own = _tokens(lp.indicators[0])
            combined = _tokens(
                " ".join([lp.starter_activity, lp.assessment, lp.conclusion]
                         + [a.description for a in lp.main_activities])
            )
            # At least one content word from the indicator appears in the content.
            assert own & combined, f"lesson {lp.indicator_codes} is not about its indicator"

    def test_assessment_is_indicator_specific(self):
        plans, _ = _run(Subject.SCIENCE, TWELVE, "Strand", "Sub-Strand")
        assessments = [lp.assessment for lp in plans]
        assert len(set(assessments)) >= 5

    def test_activities_differ_for_different_activity_types(self):
        solve = _run(Subject.MATHEMATICS, [TWELVE[0]], "Number", "Whole Numbers")[0][0]
        classify = _run(Subject.SCIENCE, [TWELVE[4]], "Diversity", "Matter")[0][0]
        discuss = _run(Subject.SOCIAL_STUDIES, [TWELVE[9]], "Family", "Roles")[0][0]
        solve_txt = " ".join(a.description for a in solve.main_activities).lower()
        classify_txt = " ".join(a.description for a in classify.main_activities).lower()
        discuss_txt = " ".join(a.description for a in discuss.main_activities).lower()
        assert solve_txt != classify_txt != discuss_txt
        assert "problem" in solve_txt or "worked" in solve_txt
        assert "sort" in classify_txt or "group" in classify_txt
        assert "discuss" in discuss_txt or "scenario" in discuss_txt


# ── Defect 4 — teacher keywords survive ─────────────────────────────────────

class TestTeacherKeywordPropagation:
    def test_keywords_in_every_deterministic_lesson(self):
        plans, _ = _run(
            Subject.SCIENCE, TWELVE[:4], "Strand", "Sub-Strand",
            keywords=["photosynthesis", "chlorophyll"],
        )
        assert plans, "expected lessons"
        for lp in plans:
            low = [k.lower() for k in lp.keywords]
            assert "photosynthesis" in low
            assert "chlorophyll" in low

    def test_keywords_survive_ai_enrichment_merge(self, monkeypatch):
        class FakeProvider(AIProvider):
            def is_available(self):
                return True

            def get_name(self):
                return "fake"

            def generate_lesson_v2(self, *, subject, class_level, strand,
                                   sub_strand, content_standard, indicator_code,
                                   indicator_text, class_size=35,
                                   duration_minutes=60, source_resources=None,
                                   previous_lesson_context=None,
                                   next_lesson_context=None, teaching_day=None,
                                   week_number=None, term=None,
                                   teaching_week=None, period=None,
                                   teacher_keywords=None):
                # AI returns its own vocabulary and does NOT include the teacher's.
                return {
                    "learning_objectives": [f"Learners can {indicator_text}"],
                    "key_vocabulary": ["ai_term_one"],
                    "starter": {"activity": f"AI starter about {indicator_text}"},
                    "main_learning": {"phase1": {
                        "name": "AI MAIN", "activity": f"AI main about {indicator_text}",
                        "duration_minutes": 30}},
                    "assessment": {"activity": f"AI assessment of {indicator_text}"},
                    "plenary": {"activity": f"AI plenary on {indicator_text}"},
                }

        monkeypatch.setattr(
            "src.engines.generation_pipeline.get_provider", lambda *a, **k: FakeProvider()
        )
        plans, job = _run(
            Subject.SCIENCE, TWELVE[:3], "Strand", "Sub-Strand",
            ai_mode=AIMode.BASIC, keywords=["keppler", "geostationary"],
        )
        assert job.ai_enrichment_succeeded == len(plans)
        for lp in plans:
            low = [k.lower() for k in lp.keywords]
            assert "keppler" in low and "geostationary" in low, (
                f"teacher keywords were lost in AI enrichment: {lp.keywords}"
            )
            assert "ai_term_one" in low


# ── Defect 2 — AI mode wiring / provider reporting ──────────────────────────

class TestAIModeWiring:
    def test_ai_status_off_is_inactive(self):
        from src.routers.settings import ai_status
        res = asyncio.run(ai_status("OFF"))
        assert res["active"] is False
        assert res["state"] == "OFF"
        assert res["provider"] is None

    def test_ai_status_reports_unavailable_provider(self, monkeypatch):
        import src.engines.ai_provider as ap

        class Unavailable(AIProvider):
            def is_available(self):
                return False

            def get_name(self):
                return "gemini"

        monkeypatch.setattr(ap, "get_provider", lambda *a, **k: Unavailable())
        monkeypatch.setattr(ap, "resolve_provider_mode", lambda *a, **k: "gemini")
        from src.routers.settings import ai_status
        res = asyncio.run(ai_status("ENHANCED"))
        assert res["active"] is False
        assert res["provider"] == "gemini"
        assert res["reason"] and "provider" in res["reason"].lower()

    def test_ai_status_reports_available_provider(self, monkeypatch):
        import src.engines.ai_provider as ap

        class Available(AIProvider):
            def is_available(self):
                return True

            def get_name(self):
                return "groq"

        monkeypatch.setattr(ap, "get_provider", lambda *a, **k: Available())
        monkeypatch.setattr(ap, "resolve_provider_mode", lambda *a, **k: "groq")
        from src.routers.settings import ai_status
        res = asyncio.run(ai_status("BASIC"))
        assert res["active"] is True
        assert res["provider"] == "groq"

    def test_ai_off_generates_deterministically(self):
        plans, job = _run(Subject.SCIENCE, TWELVE[:2], "S", "Sub",
                          ai_mode=AIMode.OFF)
        assert job.ai_enrichment_succeeded == 0
        assert all(lp.ai_generated is False for lp in plans)

    def test_ai_mode_without_provider_falls_back_and_is_honest(self, monkeypatch):
        import src.engines.generation_pipeline as gp

        class Unavailable(AIProvider):
            def is_available(self):
                return False

            def get_name(self):
                return "noprovider"

        monkeypatch.setattr(gp, "get_provider", lambda *a, **k: Unavailable())
        plans, job = _run(Subject.SCIENCE, TWELVE[:2], "S", "Sub",
                          ai_mode=AIMode.ENHANCED)
        # Deterministic content still produced; nothing falsely marked AI.
        assert len(plans) == 2
        assert job.ai_enrichment_succeeded == 0
        assert all(lp.ai_generated is False for lp in plans)
        assert all(lp.main_activities for lp in plans)

    def test_ai_mode_with_provider_marks_ai_generated(self, monkeypatch):
        class FakeProvider(AIProvider):
            def is_available(self):
                return True

            def get_name(self):
                return "fake"

            def generate_lesson_v2(self, *, subject, class_level, strand,
                                   sub_strand, content_standard, indicator_code,
                                   indicator_text, **kw):
                return {
                    "learning_objectives": [f"Learners can {indicator_text}"],
                    "key_vocabulary": ["ai_kw"],
                    "starter": {"activity": f"AI starter {indicator_text}"},
                    "main_learning": {"phase1": {
                        "name": "AI MAIN",
                        "activity": f"AI activity {indicator_text}",
                        "duration_minutes": 25}},
                    "assessment": {"activity": f"AI assessment {indicator_text}"},
                    "plenary": {"activity": f"AI plenary {indicator_text}"},
                }

        monkeypatch.setattr(
            "src.engines.generation_pipeline.get_provider",
            lambda *a, **k: FakeProvider(),
        )
        plans, job = _run(Subject.SCIENCE, TWELVE[:3], "S", "Sub",
                          ai_mode=AIMode.BASIC)
        assert job.ai_enrichment_succeeded == len(plans)
        assert all(lp.ai_generated for lp in plans)
        # AI content is indicator-specific per lesson.
        starters = {lp.starter_activity for lp in plans}
        assert len(starters) == len(plans)


# ── Defect 5 — TLRs / competencies / references derived from context ────────

class TestContextualFields:
    def test_resources_track_activity_type(self):
        solve, _ = _run(Subject.MATHEMATICS, [TWELVE[0]], "Number", "Whole")
        classify, _ = _run(Subject.SCIENCE, [TWELVE[4]], "Diversity", "Matter")
        assert solve[0].teaching_learning_resources != classify[0].teaching_learning_resources

    def test_competencies_vary_by_activity(self):
        solve, _ = _run(Subject.MATHEMATICS, [TWELVE[0]], "Number", "Whole")
        create, _ = _run(Subject.CREATIVE_ARTS, [TWELVE[8]], "Visual Arts", "Patterns")
        assert solve[0].core_competencies != create[0].core_competencies

    def test_references_are_teacher_entered(self):
        """PART L/M: the builder no longer fabricates curriculum references —
        they stay empty until the teacher enters them."""
        plans, _ = _run(Subject.SCIENCE, TWELVE[:2], "Diversity", "Matter")
        for lp in plans:
            assert lp.references == []
            assert lp.structured_references == []

    def test_teacher_resources_are_per_lesson_additions(self):
        """PART I: the batch TLR seed is NOT copied into per-lesson Other TLRs
        or the display union — teacher additions are entered per lesson."""
        config = _config(Subject.SCIENCE)
        config.teaching_learning_resources = ["School lab kit"]
        from src.curriculum.lesson_builder import build_lesson
        from src.models import AllocatedIndicator
        alloc = AllocatedIndicator(
            indicator_code="B8.1.1.1.1",
            indicator_description=TWELVE[0],
            content_standard_code="", content_standard_description="",
            strand="S", sub_strand="Sub", week_number=1, period_index=1,
            teaching_week=1,
        )
        lp = build_lesson(alloc, config, "s")
        assert lp.other_tlrs == []
        assert "School lab kit" not in lp.teaching_learning_resources
