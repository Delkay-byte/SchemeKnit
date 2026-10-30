"""Phase 3 — DoK assessment + analytic rubric tests.

Covers the acceptance list for the assessment/rubric enhancement:
appropriate DoK selection, indicator-aligned items, computed-answer
correctness, meaningful rubric criteria and totals, no invented official
grading claims, class-level appropriateness, ordinary lessons that need no
rubric, persistence + reload, WAPEF compatibility, DOCX/PDF fidelity and the
P4 gate staying green.
"""

from __future__ import annotations

import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest

from src.curriculum.assessment_design import (
    build_assessment_items,
    build_rubric,
    dok_level,
)
from src.models import (
    AIMode, AnalyticRubric, AssessmentItem, ClassLevel, SchemeOfWork, Subject,
    TermConfig, Week, WeekType,
)

B7 = ClassLevel.BASIC_7
B9 = ClassLevel.BASIC_9

# Reference you can reuse — the mission's illustrative DoK patterns.
B9_CREATIVE_INDICATOR = ("B9.1.3.1.1 Plan and execute design projects using "
                         "the creative process")


def _b(*, level: ClassLevel, code: str, indicator: str, subject: Subject,
       tlrs=("chart", "exercise books"), duration=60, teaching_days=(0,)):
    week = Week(
        week_number=1, start_date=date(2026, 9, 7), end_date=date(2026, 9, 11),
        week_type=WeekType.INSTRUCTION,
        strand="S", sub_strand="SS",
        content_standards=["CS " + code], indicators=[indicator],
        resources=list(tlrs), scheme_of_work_id="p3t",
    )
    scheme = SchemeOfWork(
        id="p3t", filename="p3.docx", class_level=level, subject=subject,
        term="First Term", academic_year="2026/2027",
        weeks=[week], upload_date=date(2026, 9, 1),
    )
    config = TermConfig(
        scheme_of_work_id="p3t", academic_year="2026/2027", term="First Term",
        class_level=level, subject=subject,
        term_start_date=date(2026, 9, 7), term_end_date=date(2026, 12, 18),
        lessons_per_week=1, lesson_duration_minutes=duration,
        teaching_days=list(teaching_days), holidays=[], ai_mode=AIMode.OFF,
    )
    return scheme, config


def _generate(scheme, config):
    from src.engines.generation_pipeline import GenerationPipeline
    job = GenerationPipeline().generate_all(scheme, config)
    return list(job._lesson_plans)[0]


# ── 1. DoK selection is action- and level-appropriate ───────────────────────


@pytest.mark.parametrize("indicator,level,duration,expect", [
    # routine computation stays a skills-and-concepts task (DoK 2)
    ("B9.1.1.1.1 Add whole numbers up to 10,000", B7, 60, 2),
    # pure recall is DoK 1
    ("B7.1.1.1.1 Identify hardware components", B7, 60, 1),
    # investigation reaches strategic thinking
    ("B9.2.1.1.1 Investigate how water pollution affects the community",
     B9, 60, 3),
    # evaluation/judgement is strategic thinking too
    ("B6.3.4.1.1 Assess the changes that the European presence brought "
     "to Ghana", ClassLevel.BASIC_6, 60, 3),
    # extended thinking only for a genuine design demand in a full JHS lesson
    ("B9.1.3.1.1 Plan and execute design projects using the creative "
     "process", B9, 60, 4),
    # a 30-minute lesson does NOT reach extended thinking
    ("B9.1.3.1.1 Plan and execute design projects using the creative "
     "process", B9, 30, 3),
])
def test_dok_level_selection(indicator, level, duration, expect):
    got = dok_level(indicator, level.value, duration)
    assert got == expect, (indicator, level, duration, got)


def test_not_every_level_four():
    """No lesson is forced upward — routine skills stay at their honest level."""
    assert dok_level("B7.1.1.1.1 Identify hardware components",
                     B7.value, 60) == 1
    assert dok_level("B9.1.1.1.1 Add two numbers", B9.value, 60) == 2


def test_dok_is_not_bloom():
    """DoK and the Bloom category stay distinct on the same action."""
    from src.curriculum.indicator_actions import bloom_category
    action = "compare"
    assert dok_level("compare fractions with the same denominator",
                     B7.value, 60, action) == 3      # strategic demand
    assert bloom_category(action) == "understand"    # Bloom maps differently


# ── 2. Items assess THIS indicator; marking is usable ──────────────────────


def test_items_are_indicator_aligned():
    items = build_assessment_items(
        indicator=B9_CREATIVE_INDICATOR, indicator_code="B9.1.3.1.1",
        focus="planning and executing design projects",
        subject="Creative Arts and Design", class_level="Basic 9",
        duration_minutes=60,
        context_hint="local craft, kente and adinkra examples",)
    assert 2 <= len(items) <= 3
    assert any("design project" in i["prompt"].lower() or
               "creative" in i["prompt"].lower() for i in items)
    assert all(i["marking_guidance"] and i["expected_evidence"] for i in items)
    assert all(i["misconception"] and i["follow_up"] for i in items)


def test_computed_math_answer_is_actually_correct():
    items = build_assessment_items(
        indicator="B6.1.1.1.1 Multiply two-digit numbers by two-digit numbers",
        indicator_code="B6.1.1.1.1", focus="multiplying two-digit numbers",
        subject="Mathematics", class_level="Basic 6", duration_minutes=60)
    comp = [i for i in items if i["prompt"].strip().startswith(("Work out",))]
    assert comp, "expected the computed check for a multiply indicator"
    prompt = comp[0]["prompt"]
    m = re.search(r"(\d+)\s*×\s*(\d+)", prompt)
    assert m, prompt
    product = int(m.group(1)) * int(m.group(2))
    expected = re.search(r"Correct answer: (\d+)",
                         comp[0]["marking_guidance"])
    assert expected and int(expected.group(1)) == product


def test_short_indicator_items_stay_usable():
    """Ordinary short checks carry marking guidance — never rubric theatre."""
    items = build_assessment_items(
        indicator="B7.1.1.1.1 Identify hardware components",
        indicator_code="B7.1.1.1.1", focus="hardware components",
        subject="ICT", class_level="Basic 7", duration_minutes=60)
    assert items and all(1 <= i["dok_level"] <= 4 for i in items)


# ── 3. Rubrics: only where a product/performance needs marking ─────────────


def test_rubric_for_extended_product_is_concrete():
    r = build_rubric(indicator=B9_CREATIVE_INDICATOR,
                     focus="a design project using the creative process",
                     class_level="Basic 9")
    assert r is not None
    criteria = [c["criterion"] for c in r["criteria"]]
    # Criteria observable on the learner's own design work — no generic text.
    assert any("criteria of the task" in c.lower() for c in criteria)
    assert any("justif" in c.lower() or "explanation" in c.lower()
               for c in criteria)
    weights = sum(c["points"] for c in r["criteria"])
    assert r["total_points"] == weights
    assert len(r["criteria"]) <= 4
    for c in r["criteria"]:
        assert len(c["descriptors"]) == 4
        assert [d["points"] for d in c["descriptors"]] == ["4", "3", "2", "1"]


def test_no_invented_official_grading_claims():
    r = build_rubric(indicator="B7.5.1.1.1 Create a simple pattern using "
                               "local materials",
                     focus="a pattern with local materials",
                     class_level="Basic 7")
    assert r is not None
    note = (r["authority_note"] or "").lower()
    assert "not an official na" in note and "ges" in note
    blob = " ".join(c["criterion"].lower() for c in r["criteria"]) + " " + \
        " ".join(d.get("descriptor", "").lower()
                 for c in r["criteria"] for d in c["descriptors"])
    assert "official" not in blob


def test_ordinary_check_lesson_gets_no_rubric():
    assert build_rubric(
        indicator="B7.1.1.1.1 Identify hardware components",
        focus="hardware components", class_level="Basic 7") is None
    assert build_rubric(indicator="B6.1.1.1.6 Skip count in 5,000s",
                        focus="skip counting", class_level="Basic 6") is None


def test_early_years_get_no_rubric():
    assert build_rubric(indicator="Sing some traditional songs",
                        focus="traditional songs",
                        class_level="Nursery 2") is None


# ── 4. Generation integration (AI OFF) ─────────────────────────────────────


def test_generation_attaches_structure_matching_demand():
    lp = _generate(*_b(level=B9, code="B9.1.3.1.1", subject=Subject.CREATIVE_ARTS,
                       indicator=B9_CREATIVE_INDICATOR))
    assert len(lp.assessment_items) >= 2
    assert lp.assessment_rubric is not None          # product/performance
    assert lp.assessment_rubric.total_points >= 3
    assert lp.assessment_rubric.criteria


def test_generation_no_rubric_for_simple_check():
    lp = _generate(*_b(level=B7, code="B7.1.1.1.1", subject=Subject.ICT,
                       indicator="B7.1.1.1.1 Identify hardware components"))
    assert lp.assessment_items, "items still required for a check lesson"
    assert lp.assessment_rubric is None              # short check → no rubric


def test_dok_spread_not_forced():
    """Recall-led lessons stay DoK 1/2 in the primary item — honest spread."""
    lp = _generate(*_b(level=B7, code="B7.1.1.1.1", subject=Subject.ICT,
                       indicator="B7.1.1.1.1 Identify hardware components"))
    top = lp.assessment_items[1]                     # the indicator-action check
    assert top.dok_level == 1


# ── 5. Persistence + reload (legacy-safe) ──────────────────────────────────


def test_save_and_reload_assessment_structure(db):
    from src.database import LessonPlanDB
    from src.service import DataService
    from tests.conftest import make_user
    from tests.test_export_download import make_job_with_lessons
    from src.routers.generation import parse_assessment_items, \
        parse_assessment_rubric

    user = make_user(db, email="p3-struct@t.test", role="teacher")
    _scheme, job = make_job_with_lessons(db, user, lessons=1)
    lp_db = db.query(LessonPlanDB).filter_by(job_id=job.id).first()
    model = parse_assessment_items(lp_db.assessment_items)
    assert isinstance(model, list)

    data_service = DataService()
    lp_model = _model_with_structure()
    data_service.create_lesson_plan(
        db, owner_id=user.id, job_id=job.id, scheme_id=job.scheme_id,
        lp=lp_model,
    )
    fresh = db.query(LessonPlanDB).filter_by(job_id=job.id).order_by(
        LessonPlanDB.created_at.desc()).first()
    assert fresh.assessment_rubric is not None
    assert fresh.assessment_items
    reloaded = parse_assessment_rubric(fresh.assessment_rubric)
    assert isinstance(reloaded, AnalyticRubric)
    # "create" category → 4 criteria × 4 points (4-band rubric) = 16.
    assert reloaded.total_points == 16
    items = parse_assessment_items(fresh.assessment_items)
    assert isinstance(items[0], AssessmentItem)
    assert items[0].dok_level >= 1


def _model_with_structure():
    """A generated Creative Arts lesson carrying items + a rubric."""
    lp = _generate(*_b(level=B9, code="B9.1.3.1.1",
                       subject=Subject.CREATIVE_ARTS,
                       indicator=B9_CREATIVE_INDICATOR))
    assert lp.assessment_rubric is not None
    return lp


# ── 6. WAPEF + DOCX/PDF fidelity ───────────────────────────────────────────


def test_wapef_unaffected():
    """The WAPEF teacher-selection fields are untouched by the new columns."""
    from src.engines.wapef_fields import (
        WAPEF_DEEP_HOPES, WAPEF_STORYLINES, WAPEF_THROUGH_LINES,
        WAPEF_GODS_STORY,
    )
    assert WAPEF_DEEP_HOPES and WAPEF_STORYLINES and WAPEF_THROUGH_LINES
    assert WAPEF_GODS_STORY
    lp = _generate(*_b(level=B7, code="B7.1.1.1.1", subject=Subject.ICT,
                       indicator="B7.1.1.1.1 Identify hardware components"))
    assert lp.wapef_deep_hope == "" and lp.wapef_through_lines == []


def test_assessment_appendix_uses_break_before_without_a_blank_page(tmp_path):
    """Starting the appendix must not insert a page-break paragraph.

    A break paragraph at the top of a page produces a blank page when the
    preceding GES form filled the previous page exactly. The heading carries
    page_break_before instead, which is blank-page-safe.
    """
    from docx import Document
    from src.engines.docx_export import (
        DOCXExportEngine, default_template_for_lessons,
    )

    lp = _model_with_structure()
    out = DOCXExportEngine().export_combined(
        [lp], default_template_for_lessons([lp]), tmp_path / "p3-combined.docx")
    paragraphs = Document(str(out)).paragraphs
    index = next(
        i for i, p in enumerate(paragraphs)
        if p.text.startswith("Assessment Structure")
    )
    assert paragraphs[index].paragraph_format.page_break_before is True
    if index:
        previous = paragraphs[index - 1]
        assert 'w:type="page"' not in previous._p.xml or previous.text.strip(), (
            "an empty page-break paragraph sits before the assessment appendix"
        )


def test_docx_export_contains_assessment_structure(tmp_path):
    lp = _model_with_structure()
    from src.engines.docx_export import DOCXExportEngine
    out = DOCXExportEngine().export_single(lp, output_path=tmp_path / "p3.docx")
    assert out.exists()
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "Assessment Questions / Checks" in text
    assert "Marking Rubric for the Learner Task" in text
    assert f"Total marks: {lp.assessment_rubric.total_points}" in text
    assert "not an official" in text.lower()
    # All four bands must keep their names in the export: the descriptor
    # points are 4/3/2/1, so a legacy 3-band label map silently printed
    # "4 points" for the top band.
    for band in ("4 — exceeds the standard", "3 — meets the standard",
                 "2 — developing", "1 — needs support"):
        assert band in text, band
    # the prompts themselves survive export
    first_prompt = lp.assessment_items[0].prompt
    assert first_prompt[:40] in text


def test_structured_pdf_contains_assessment_structure(tmp_path):
    lp = _model_with_structure()
    import src.engines.structured_pdf as spdf
    out = spdf.render_structured_pdf([lp], output_path=tmp_path / "p3.pdf")
    raw = open(out, "rb").read()
    assert raw[:5] == b"%PDF-"
    pypdf = pytest.importorskip("pypdf")
    reader = pypdf.PdfReader(str(out))
    text = "".join((p.extract_text() or "") for p in reader.pages)
    assert "Assessment Questions" in text
    assert "Marking Rubric" in text
    assert str(lp.assessment_rubric.total_points) in text
    assert "exceeds the standard" in text
    assert "needs support" in text


def test_legacy_pdf_export_path_still_works(tmp_path):
    """The plain PDF path must keep working (assessment prose unchanged)."""
    from src.engines.pdf_export import PDFExportEngine
    lp = _model_with_structure()
    result = PDFExportEngine().export_single(lp, output_path=tmp_path / "p3.pdf")
    raw = open(result, "rb").read()
    assert raw[:5] == b"%PDF-"


# ── 7. P4 quality gate stays green with the new structure ──────────────────


def test_p4_accepts_lessons_with_assessment_structure():
    lp = _model_with_structure()
    from src.curriculum.lesson_rubric import hard_failures as rubric_failures
    entry = dict(
        indicator="Plan and execute design projects using the creative process",
        code="B9.1.3.1.1", week=1, subject=Subject.CREATIVE_ARTS,
        class_level=ClassLevel.BASIC_9, strand="S", sub_strand="SS",
        content_standard="CS B9.1.3.1.1", content_standard_code="B9.1.3.1",
        resources=["chart"], provenance="", group="")
    assert rubric_failures(entry, lp) == []


# ── 8. API read path + partial-update + Zeli preservation ───────────────────


def _saved_structure_lesson(db, email: str):
    """Persist a generated lesson carrying items + rubric; return (user, lp_db)."""
    from tests.conftest import make_user
    from tests.test_export_download import make_job_with_lessons
    from src.service import DataService

    user = make_user(db, email=email, role="teacher")
    _scheme, job = make_job_with_lessons(db, user, lessons=1)
    lp_db = DataService().create_lesson_plan(
        db, owner_id=user.id, job_id=job.id, scheme_id=job.scheme_id,
        lp=_model_with_structure())
    return user, lp_db


def test_serialize_lesson_includes_assessment_structure(db):
    """The workspace read path (_serialize_lesson) carries the Phase 3 fields."""
    from src.routers.generation import _serialize_lesson

    _user, lp_db = _saved_structure_lesson(db, "p3-api@t.test")
    serialized = _serialize_lesson(lp_db)
    assert len(serialized["assessment_items"]) == len(lp_db.assessment_items or [])
    assert serialized["assessment_items"][0]["prompt"]
    assert serialized["assessment_items"][0]["dok_level"] >= 1
    assert serialized["assessment_rubric"] is not None
    assert serialized["assessment_rubric"]["total_points"] >= 3
    assert serialized["assessment_rubric"]["criteria"]
    # JSON-safe for the API response (no model objects leaking out).
    import json as _json
    _json.dumps(serialized["assessment_items"])
    _json.dumps(serialized["assessment_rubric"])


def test_serialize_lesson_tolerates_legacy_rows_without_structure(db):
    """Pre-Phase-3 lessons (NULL columns) stay readable on the same path."""
    from tests.conftest import make_user
    from tests.test_export_download import make_job_with_lessons
    from src.database import LessonPlanDB
    from src.routers.generation import _serialize_lesson

    user = make_user(db, email="p3-legacy-api@t.test", role="teacher")
    _scheme, job = make_job_with_lessons(db, user, lessons=1)
    lp_db = db.query(LessonPlanDB).filter_by(job_id=job.id).first()
    lp_db.assessment_items = None
    lp_db.assessment_rubric = None
    db.commit()
    serialized = _serialize_lesson(lp_db)
    assert serialized["assessment_items"] == []
    assert serialized["assessment_rubric"] is None


def test_partial_update_does_not_drop_assessment_structure(db):
    """A teacher's prose edit must never erase the structured assessment."""
    from src.routers.generation import parse_assessment_rubric
    from src.service import DataService

    user, lp_db = _saved_structure_lesson(db, "p3-partial@t.test")
    original_points = lp_db.assessment_rubric["total_points"]
    updated = DataService().update_lesson_plan(
        db, lp_db.id, user.id, {"conclusion": "Changed conclusion text."})
    assert updated is not None
    assert updated.conclusion == "Changed conclusion text."
    assert updated.assessment_items, "items dropped by a partial update"
    assert updated.assessment_rubric, "rubric dropped by a partial update"
    assert updated.assessment_rubric["total_points"] == original_points


@pytest.mark.asyncio
async def test_zeli_section_rewrite_preserves_assessment_structure(db):
    """An optional Zeli rewrite returns a suggestion and touches nothing else.

    Contract: the rewrite is a *suggestion* — the stored lesson (prose and
    every Phase 3 field) stays intact until the teacher saves; the structured
    assessment columns are never rewritten by Zeli at any point.
    """
    from unittest.mock import patch
    from tests.conftest import make_entitled_teacher
    from tests.test_ai_rewrite_wiring import _make_lesson, _StructuredProvider
    from src.routers import ai_regeneration as air

    user, _school, _lic = make_entitled_teacher(db, email="p3-zeli@t.test")
    lp = _make_lesson(db, user)
    lp.assessment_items = _model_with_structure().model_dump(
        mode="json")["assessment_items"]
    lp.assessment_rubric = _model_with_structure().model_dump(
        mode="json")["assessment_rubric"]
    db.commit()
    before_items = list(lp.assessment_items)
    before_rubric = dict(lp.assessment_rubric)
    before_activities = list(lp.main_activities or [])

    provider = _StructuredProvider()
    provider.is_available = lambda: True
    provider.get_name = lambda: "structured"
    provider.model = "test-model"
    with patch.object(air, "get_provider", return_value=provider):
        res = await air.regenerate_section(
            air.SectionRegenerateRequest(
                lesson_plan_id=lp.id, section="assessment", ai_mode="gemini"),
            user, db)
    assert res.source == "ai"
    assert res.new_content.startswith("New AI assessment")
    assert res.previous_content == "Old assessment text here."

    db.refresh(lp)
    # Suggestion only: the stored prose is untouched until the teacher saves.
    assert lp.assessment == "Old assessment text here."
    assert lp.assessment_items == before_items, "Zeli rewrote the items"
    assert lp.assessment_rubric == before_rubric, "Zeli rewrote the rubric"
    assert list(lp.main_activities or []) == before_activities
    # WAPEF + provenance-bearing columns untouched by a prose rewrite.
    assert lp.wapef_through_lines == []
