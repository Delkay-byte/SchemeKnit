"""Phase 11 regression tests — the Basic 9 Science fixture defects.

The attached ``Lesson_Plans_BASIC_9_SCIENCE_SCHEME_OF_LEARNING.docx``
regression fixture exposed four classes of defect:

1. Week sequence corruption: a five-week plan printed Week 1, 2, 2, 3, 4 —
   the second "WEEK 2" source section merged into the first, so a real week
   disappeared and another carried two lessons.
2. Header structure: every lesson page repeated "School / First Term. Week N /
   Teacher: …" instead of the required first-page order (school, teacher,
   Week 1) and week-only headings afterwards.
3. ``DATE / DAY:`` rows showing calendar dates squeezed beside day names.
4. AI-style drafting artefacts in teacher-facing prose ("What to look for:",
   "Quick check (unprompted):", raw "NS:" / "MA:" markers, mid-dot step
   labels, truncated "…" clauses, "Learners can Identify", "Charts And
   Pictures", concatenated indicators).

Each class gets a focused regression test here, built from the defect itself.
"""

import asyncio
import pytest
from pathlib import Path
from datetime import date

from docx import Document

from src.parsers.docx_parser import DOCXParser
from src.curriculum import text_hygiene
from src.engines.docx_export import DOCXExportEngine, render_custom_template
from src.engines.official_ges_template import (
    render_document as render_official_ges_document,
    render_header,
)
from src.engines.official_ges_levels import JHS_SPEC
from src.models import LessonPlan, ClassLevel, Subject, EducationalLevel


# ── Helpers ──────────────────────────────────────────────────────────────────

def _write_docx(path: Path, paragraphs, rows):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                table.rows[r].cells[c].text = val
    doc.save(str(path))
    return path


def _parse(path):
    return asyncio.run(DOCXParser().parse(path))


def _header_lines(doc) -> list:
    """Flatten a rendered document's paragraph text into header lines.

    A multi-line heading (school / teacher / week) is one paragraph with real
    line breaks, so paragraph text is split back into lines.
    """
    return [line.strip() for p in doc.paragraphs if p.text.strip()
            for line in p.text.split("\n") if line.strip()]


HEADER = ["WEEK ENDING", "WEEK", "STRAND", "SUB-STRAND", "CONTENT STANDARD",
          "INDICATORS", "RESOURCES"]


def _week_row(week, strand, substrand, cs, indicator, week_ending):
    return [week_ending, str(week), strand, substrand, cs, indicator, "charts"]


# ── 1. Week sequence: the five-week duplication defect ───────────────────────

class TestWeekSequenceRegression:
    def test_repeated_week_label_does_not_merge_two_curriculum_weeks(self, tmp_path):
        """A source whose sections repeat "WEEK 2" must yield FIVE weeks.

        The fixture printed Week 1, 2, 2, 3, 4 for a five-week plan because
        the second "WEEK 2" section (a different strand, the NEXT section's
        week-ending date) merged into the first. The parser must split the
        mislabeled section off and renumber so the extracted weeks are
        Week 1..Week 5 in order.
        """
        rows = [
            HEADER,
            _week_row("1", "Diversity of Matter", "Materials",
                      "B9.1.1.1 CS", "B9.1.1.1.1 Identify binary compounds",
                      "11 September 2026"),
            _week_row("2", "Diversity of Matter", "Materials",
                      "B9.1.1.1 CS", "B9.1.1.1.2 Discuss acids, bases and salts",
                      "18 September 2026"),
            # The duplicated section label: new strand content, same label,
            # but the NEXT section's week-ending date.
            _week_row("2", "Cycles", "Life Cycle of a Flowering Plant",
                      "B9.2.1.1 CS", "B9.2.1.1.1 Describe the life cycle",
                      "25 September 2026"),
            _week_row("3", "Systems", "The Solar System",
                      "B9.3.1.1 CS", "B9.3.1.1.1 Describe the solar system",
                      "02 October 2026"),
            _week_row("4", "Forces and Energy", "Electricity",
                      "B9.4.1.1 CS", "B9.4.1.1.1 Explain simple circuits",
                      "09 October 2026"),
        ]
        path = _write_docx(tmp_path / "dup_week2.docx", [], rows)
        parsed = _parse(path)
        numbers = [w.week_number for w in parsed.weeks]
        assert numbers == [1, 2, 3, 4, 5], numbers
        # Each curriculum week keeps its own strand — nothing was merged.
        strands = [w.strand for w in parsed.weeks]
        assert "Cycles" in strands and "Systems" in strands
        # The five distinct indicator requirements all survive.
        all_indicators = " ".join(
            " ".join(w.indicators or []) for w in parsed.weeks)
        assert "binary compounds" in all_indicators
        assert "life cycle" in all_indicators
        assert "solar system" in all_indicators
        assert "simple circuits" in all_indicators

    def test_repeating_same_strand_rows_stay_in_one_week(self, tmp_path):
        """Continuation rows legitimately repeat their week label + strand."""
        rows = [
            HEADER,
            _week_row("1", "Diversity of Matter", "Materials",
                      "B9.1.1.1 CS", "B9.1.1.1.1 Identify binary compounds",
                      "11 September 2026"),
            _week_row("1", "Diversity of Matter", "Materials",
                      "B9.1.1.1 CS", "B9.1.1.1.2 Discuss acids, bases and salts",
                      "11 September 2026"),
            _week_row("2", "Cycles", "Life Cycle",
                      "B9.2.1.1 CS", "B9.2.1.1.1 Describe the life cycle",
                      "18 September 2026"),
        ]
        path = _write_docx(tmp_path / "same_week_repeat.docx", [], rows)
        parsed = _parse(path)
        assert [w.week_number for w in parsed.weeks] == [1, 2]
        # Both week-1 indicators stay in week 1.
        assert len(parsed.weeks[0].indicators) == 2

    def test_multi_strand_week_sharing_one_date_stays_one_week(self, tmp_path):
        """A real document's multi-strand week is NOT split (Basic 7 guard)."""
        rows = [
            HEADER,
            _week_row("1", "Diversity of Matter", "Materials",
                      "B9.1.1.1 CS", "B9.1.1.1.1 Identify binary compounds",
                      "11 September 2026"),
            _week_row("1", "Cycles", "Life Cycle",
                      "B9.2.1.1 CS", "B9.2.1.1.1 Describe the life cycle",
                      "11 September 2026"),
            _week_row("2", "Systems", "Solar System",
                      "B9.3.1.1 CS", "B9.3.1.1.1 Describe the solar system",
                      "18 September 2026"),
        ]
        path = _write_docx(tmp_path / "multi_strand_week.docx", [], rows)
        parsed = _parse(path)
        assert [w.week_number for w in parsed.weeks] == [1, 2]
        assert len(parsed.weeks[0].indicators) == 2


# ── 2. Header structure (official GES renderer) ──────────────────────────────

def _lesson(week=1, school="Test School", teacher="Ama Mensah"):
    return LessonPlan(
        id=f"lp-w{week}", scheme_of_work_id="scheme-1",
        term_config_id="term-1",
        week_number=week, lesson_sequence=week,
        lesson_date=date(2026, 9, 7), week_ending=date(2026, 9, 11),
        class_level=ClassLevel.BASIC_9, subject=Subject.SCIENCE,
        educational_level=EducationalLevel.JHS,
        school_name=school, teacher_name=teacher,
        duration_minutes=60, class_size=31,
    )


class TestGesHeaderStructure:
    def test_first_page_shows_school_teacher_week_in_order(self):
        doc = Document()
        doc.add_paragraph("OFFICIAL GES / NaCCA WEEKLY LESSON PLAN TEMPLATE")
        replaced = render_header(doc, _lesson(1), first=True)
        assert replaced
        lines = _header_lines(doc)
        assert lines[0] == "Test School"
        assert lines[1] == "Ama Mensah"
        assert lines[2] == "Week 1"

    def test_later_pages_show_only_their_week_heading(self):
        doc = Document()
        doc.add_paragraph("OFFICIAL GES / NaCCA WEEKLY LESSON PLAN TEMPLATE")
        replaced = render_header(doc, _lesson(3), first=False)
        assert replaced
        lines = _header_lines(doc)
        assert lines == ["Week 3"]

    def test_no_identity_renders_nothing_invented(self):
        doc = Document()
        doc.add_paragraph("OFFICIAL GES / NaCCA WEEKLY LESSON PLAN TEMPLATE")
        replaced = render_header(doc, _lesson(2, school="", teacher=""), first=True)
        # The week heading still renders, but NO school/teacher is fabricated.
        assert replaced
        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "Test School" not in texts
        assert "Ama Mensah" not in texts

    def test_multi_page_document_renders_week_sequence_once(self):
        lessons = [_lesson(w) for w in (1, 2, 3, 4, 5)]
        doc = render_official_ges_document(JHS_SPEC, lessons, {})
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Header lines can be joined in one multi-line paragraph — flatten.
        lines = [line.strip() for t in texts for line in t.split("\n")]
        # The five pages carry Week 1..Week 5 exactly once each, in order.
        week_lines = [t for t in lines if t in
                      {f"Week {n}" for n in range(1, 6)}]
        assert week_lines == [f"Week {n}" for n in range(1, 6)]
        # School and teacher appear on the FIRST page only.
        joined = "\n".join(texts)
        assert joined.count("Test School") == 1
        assert joined.count("Ama Mensah") == 1


# ── 3. DATE / DAY row: selected day names, never calendar dates ──────────────

class TestDateDayRow:
    def test_teaching_day_names_are_ordered_and_deduped(self):
        from src.routers.generation import _teaching_day_names
        assert _teaching_day_names([4, 0, 2, 0]) == [
            "Monday", "Wednesday", "Friday"]
        assert _teaching_day_names(["wednesday", "Mon"]) == [
            "Monday", "Wednesday"]
        assert _teaching_day_names([]) == []

    def test_day_labelled_date_cell_renders_day_names(self):
        """A 'DATE / DAY:' sample cell shows the configured days, no dates."""
        doc = Document()
        structure = {"tables": [{
            "rows": 1, "cols": 1,
            "cells": [{"r": 0, "c": 0, "label": "DATE / DAY",
                       "field": "lesson_date"}],
        }], "meta": {}}
        lp = _lesson(1)
        ctx = {"teaching_days": ["Monday", "Wednesday", "Thursday"]}
        render_custom_template(doc, lp, structure, ctx, first_lesson=True)
        cell_text = doc.tables[0].rows[0].cells[0].text
        assert cell_text == "DATE / DAY: Monday, Wednesday, Thursday"
        assert "2026" not in cell_text
        # No calendar dates leak into the day row.
        assert "2026" not in cell_text


# ── 4. Grammar / drafting-artefact regressions (fixture strings) ─────────────

class TestArtefactRegressions:
    def test_concatenated_indicators_get_a_sentence_boundary(self):
        fixed = text_hygiene.hygiene_pass(
            "Discuss the formation of binary chemical compounds Describe the "
            "characteristics of common acids, bases and salts.")
        assert "compounds Describe" not in fixed
        assert "compounds. Describe" in fixed

    def test_modal_verb_capitalisation_is_not_split(self):
        fixed = text_hygiene.hygiene_pass(
            "Learners can Identify by name binary chemical compounds")
        assert fixed == "Learners can identify by name binary chemical compounds"

    def test_incomplete_trailing_clause_is_trimmed(self):
        fixed = text_hygiene.hygiene_pass(
            "Learners explain the process and explain it in.")
        assert not fixed.rstrip().endswith("it in.")
        assert fixed.endswith(".")

    def test_truncated_ellipsis_is_never_emitted(self):
        fixed = text_hygiene.complete_or_clean_clause(
            "Learners add plus two key words from the example of the topic "
            "of the day and share them with the class")
        assert "…" not in fixed
        assert fixed.endswith(".")

    def test_resource_capitalisation_fix_is_targeted(self):
        assert text_hygiene.fix_resource_capitalisation(
            "Charts And Pictures") == "Charts and pictures".replace(
            "pictures", "Pictures")
        assert text_hygiene.fix_resource_capitalisation("Holy Bible") == "Holy Bible"
        assert text_hygiene.fix_resource_capitalisation("charts") == "charts"

    def test_drafting_labels_are_replaced_with_teacher_prose(self):
        fixed = text_hygiene.strip_drafting_markers(
            "Starter task. What to look for: learners can name the compounds.")
        assert "What to look for:" not in fixed
        assert "As learners respond" in fixed

    def test_raw_band_markers_become_teacher_facing_prose(self):
        fixed = text_hygiene.strip_drafting_markers(
            "Task text. (NS: give word cards; MA: extend with naming five more).")
        assert "NS:" not in fixed and "MA:" not in fixed
        assert "For learners needing support: give word cards" in fixed
        assert "For learners ready for more" in fixed

    def test_mid_dot_step_label_is_removed(self):
        fixed = text_hygiene.strip_drafting_markers(
            "Step 2 · Classify: check each learner's sorting.")
        assert "·" not in fixed
        assert "Step 2:" in fixed

    def test_quality_issues_flag_what_cannot_be_autofixed(self):
        assert text_hygiene.quality_issues("An unfinished thought from the…")
        assert text_hygiene.quality_issues("Learners sort shapes (unclosed")
        assert text_hygiene.quality_issues("Ask learners to state what they learn about.")

    def test_gate_rejects_lesson_carrying_fixture_artefacts(self):
        """The quality gate fails a lesson whose activities carry the
        fixture's artefacts — the shared guarantee for BOTH generation paths."""
        from src.engines.lesson_quality_gate import hard_failures, build_entry
        lp = _lesson(1)
        lp.starter_activity = "Starter. What to look for: learners can name things."
        lp.assessment = "Quick check (unprompted): the compounds — each learner names two."
        entry = build_entry({"week": 1}, lp)
        failures = hard_failures(lp, entry)
        assert "drafting_label_present" in failures

# ── 5. Download-URL DOCX carries the teaching-day context ──────────────────

def test_issue_download_url_docx_passes_render_context():
    """The one-time download URL is the path the UI export button uses.
    Its DOCX branch must forward the render context, otherwise DATE / DAY
    silently falls back to calendar dates."""
    import inspect
    from src.routers import generation
    source = inspect.getsource(generation.issue_download_url)
    docx_branch = source.split('if fmt == "docx"')[1].split('elif fmt == "zip"')[0]
    assert "context=_render_context" in docx_branch


def test_official_ges_date_day_row_uses_configured_days():
    from src.engines import official_ges_template as ges
    spec = ges.spec_for_level("jhs")
    lp = _lesson(1)
    ctx = {"teaching_days": ["Monday", "Wednesday", "Thursday"]}
    values = ges.placeholder_values(spec, lp, ctx)
    assert values["INSERT_DATE_DAY"] == "Monday, Wednesday, Thursday"
    assert "2026" not in values["INSERT_DATE_DAY"]
