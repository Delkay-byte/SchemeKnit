"""
No synthetic default content (item 19).

Renderers must never print generic placeholder prose ("Charts, charts and
pictures…", "Learners participate in group discussions…") when a field has no
stored data. Real structured data renders; a genuinely empty field renders
blank. Curriculum content is never fabricated.
"""

import os
import sys
from datetime import date
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from src.engines.docx_export import DOCXExportEngine
from src.engines.template_engine import get_template_by_id
from src.models import ClassLevel, LessonPlan, Subject

#: Generic fallback strings that must never be printed by any renderer.
FORBIDDEN_FALLBACKS = [
    "Charts, charts and pictures, locally available materials",
    "Teacher presents the lesson content using board work and visual aids.",
    "Learners participate in group discussions and class activities.",
    "By the end of the lesson, learners should be able to:",
]


def _empty_lesson():
    return LessonPlan(
        scheme_of_work_id="s", term_config_id="j", week_number=1, lesson_sequence=1,
        lesson_date=date(2026, 9, 11), lesson_number=1,
        class_level=ClassLevel.BASIC_9, subject=Subject.SCIENCE,
        school_name="", teacher_name="",
        strand="Diversity of Matter", sub_strand="Materials",
        content_standard="B9.1.1.1 Show understanding",
        indicators=["B9.1.1.1.1 Identify"],
        learning_objectives=[], teaching_learning_resources=[],
        introduction="", main_activities=[], learner_activities=[],
        assessment="", conclusion="",
    )


def _document_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class TestLegacyRendererHasNoFabricatedFallbacks:
    """The legacy (non-form) templates must render blank, not filler prose."""

    @pytest.mark.parametrize("template_id", ["tpl-jhs-ges", "tpl-primary-standard"])
    def test_empty_fields_render_blank_not_filler(self, tmp_path, template_id):
        out = tmp_path / "legacy.docx"
        DOCXExportEngine().export_single(
            _empty_lesson(), get_template_by_id(template_id), out)
        blob = _document_text(out)
        for filler in FORBIDDEN_FALLBACKS:
            assert filler not in blob, filler

    def test_real_data_still_renders(self, tmp_path):
        """The point is never to drop real content — only fabricated filler."""
        lesson = _empty_lesson()
        lesson.teaching_learning_resources = ["Real chart", "Real objects"]
        lesson.introduction = "Real starter the teacher wrote."
        out = tmp_path / "legacy-real.docx"
        DOCXExportEngine().export_single(
            lesson, get_template_by_id("tpl-jhs-ges"), out)
        blob = _document_text(out)
        assert "Real chart" in blob
        assert "Real starter the teacher wrote." in blob


class TestApprovedFormHasNoFabricatedFallbacks:
    """The approved form fills in place: blanks stay blanks."""

    def test_empty_lesson_fields_render_blank(self, tmp_path):
        from src.validators.docx_structure_compare import compare_content

        out = tmp_path / "approved-empty.docx"
        DOCXExportEngine().export_single(
            _empty_lesson(),
            get_template_by_id("tpl-approved-org-headteacher"), out)
        rep = compare_content(out)
        assert rep["pass"], [c for c in rep["checks"] if not c["pass"]]
        blob = _document_text(out)
        for filler in FORBIDDEN_FALLBACKS:
            assert filler not in blob, filler
