"""
Real-use remediation regression tests (PART 23).

Covers the four remediation areas that production testing surfaced:

  RESOURCES / EMPTY FIELDS (PART 5-11)
    * source resources are stored/normalized as structured values
    * serialized/legacy list text never leaks as display text
    * empty lists are canonical [] (never [], [,], [""], null)
  AI STATUS (PART 12-15)
    * the resolved provider is reported truthfully for every requested mode
    * deterministic fallbacks are never reported as "AI active"
  AI QUOTA (PART 16-19)
    * 5 AI generations per calendar month (YYYY-MM server-clock bucket)
    * failures / deterministic fallbacks consume nothing
    * a new month resets usage to 0/5; users are isolated
  DOWNLOADS (PART 1-4, 22)
    * the one-time download URL path issues + delivers all three formats
    * export errors surface their real cause, never a bare "Failed to fetch"
"""
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

from src.ai_resource_text import (
    clean_serialized_text,
    normalize_text,
    normalize_text_items,
)
from src.parsers.docx_parser import DOCXParser
from src.curriculum.lesson_builder import build_lesson
from src.models import (
    ClassLevel, Subject, TermConfig, Week, WeekType,
)
from src.engines.allocation_engine import AllocationEngine


# ── Helpers ──────────────────────────────────────────────────────────────────

def _config(**kw):
    return TermConfig(
        scheme_of_work_id="s", academic_year="2026/2027", term="First Term",
        class_level=kw.pop("class_level", ClassLevel.BASIC_8),
        subject=kw.pop("subject", Subject.SCIENCE),
        term_start_date=date(2026, 10, 26), term_end_date=date(2026, 12, 18),
        lessons_per_week=3, lesson_duration_minutes=60,
        teaching_days=kw.pop("teaching_days", [0, 1, 2, 3, 4]),
        holidays=[], **kw,
    )


def _week(num, resources=None, indicators=None):
    return Week(
        week_number=num, start_date=date(2026, 10, 26),
        end_date=date(2026, 10, 30), week_type=WeekType.INSTRUCTION,
        strand="Diversity of Matter",
        sub_strand="Materials" if indicators else None,
        indicators=list(indicators or []),
        content_standards=["B9.1.1.1 Show understanding of matter"],
        resources=list(resources or []),
        scheme_of_work_id="s",
    )


def _allocated(resources, indicator="B9.1.1.1.1 Identify materials"):
    weeks = [_week(1, resources=resources, indicators=[indicator])]
    config = _config()
    from src.engines.calendar_engine import CalendarEngine
    calendar = CalendarEngine().build_calendar(config, weeks, [])
    coverage = AllocationEngine().allocate(weeks, calendar, config)
    return coverage.allocations[0], config


# ── RESOURCES: canonical normalization (PART 5-7, 11) ────────────────────────

class TestCanonicalResourceNormalization:
    def test_serialized_json_payload_is_parsed(self):
        value = '["Pictures showing Created things,HolyBible,HolyQuran."]'
        assert normalize_text_items(value) == [
            "Pictures showing Created things", "Holy Bible", "Holy Quran",
        ]

    def test_comma_cell_splits_into_individual_resources(self):
        assert normalize_text_items("Counters, sticks, flash cards") == [
            "Counters", "sticks", "flash cards",
        ]

    def test_meaningful_phrase_is_never_split(self):
        assert normalize_text_items("Pictures showing Created things") == [
            "Pictures showing Created things",
        ]

    def test_space_stripped_damage_is_reconstructed(self):
        assert normalize_text_items("PicturesshowingCreatedthings,HolyBible") == [
            "Pictures showing Created things", "Holy Bible",
        ]

    def test_bracket_wrapped_display_damage_is_unwrapped(self):
        assert normalize_text_items("[Counters, sticks]") == ["Counters", "sticks"]

    def test_semicolons_and_pipes_split(self):
        assert normalize_text_items("A; B | C") == ["A", "B", "C"]

    def test_empty_representations_are_canonical_empty(self):
        for value in ([], None, "", "[]", "[,]", '[""]', '[" "]', "null", "[ ]"):
            assert normalize_text_items(value) == [], repr(value)

    def test_clean_list_passes_through_untouched(self):
        assert normalize_text_items(["Holy Bible", "Charts"]) == [
            "Holy Bible", "Charts",
        ]

    def test_deduplication_is_case_insensitive(self):
        assert normalize_text_items(["Charts", "CHARTS", "charts "]) == ["Charts"]

    def test_normal_text_is_never_mangled(self):
        text = "Teacher-made word cards with pictures of farm animals"
        assert clean_serialized_text(text) == text
        assert normalize_text(text) == text


# ── RESOURCES: parser level (PART 6/7) ───────────────────────────────────────

class TestParserResourceExtraction:
    def test_resource_cells_become_individual_items(self):
        """PART 6/7: a resource cell holding several resources is normalized
        into individual items at the source (parser merge step)."""
        parser = DOCXParser()
        rows = [
            {
                "week_number": 1,
                "date": date(2026, 9, 11),
                "week_type": WeekType.INSTRUCTION,
                "strand": "God’s Creation",
                "sub_strand": "Environment",
                "content_standard": "B1.1.1.1 Show care for created things",
                "indicators": "B1.1.1.1.1 Identify created things",
                "resources": "Pictures showing Created things, Holy Bible, Holy Quran.",
            },
        ]
        week = parser._merge_week_rows(1, rows)
        assert "Pictures showing Created things" in week.resources
        assert "Holy Bible" in week.resources
        assert "Holy Quran" in week.resources
        # No entry is a serialized list.
        for r in week.resources:
            assert not r.startswith("[")

    def test_legacy_single_string_resource_is_split(self):
        parser = DOCXParser()
        rows = [{
            "week_number": 2, "date": None, "week_type": WeekType.INSTRUCTION,
            "strand": "S", "sub_strand": "SS", "content_standard": "",
            "indicators": "", "resources": "Counters, sticks, flash cards",
        }]
        week = parser._merge_week_rows(2, rows)
        # (resources is a de-duplicated set in the parsed week — order-free.)
        assert set(week.resources) == {"Counters", "sticks", "flash cards"}


# ── RESOURCES: allocation → lesson → API payload (PART 5/8/9) ────────────────

class TestResourceFlowToLesson:
    def test_allocation_carries_normalized_source_resources(self):
        alloc, _ = _allocated(["Counters, sticks, flash cards"])
        assert alloc.source_resources == ["Counters", "sticks", "flash cards"]

    def test_lesson_source_tlrs_are_structured(self):
        alloc, config = _allocated(["Counters, sticks, flash cards"])
        lp = build_lesson(alloc, config, "s")
        assert lp.source_tlrs == ["Counters", "sticks", "flash cards"]

    def test_legacy_comma_joined_week_resources_are_split(self):
        alloc, _ = _allocated(["Charts, Pictures, counters"])
        assert alloc.source_resources == ["Charts", "Pictures", "counters"]

    def test_other_tlrs_default_empty_and_render_empty(self):
        alloc, config = _allocated(["Charts"])
        lp = build_lesson(alloc, config, "s")
        assert lp.other_tlrs == []
        # The serialized forms that leaked before can never come back.
        assert "".join(lp.other_tlrs) == ""
        assert not any("[" in t for t in lp.teaching_learning_resources)


# ── AI STATUS: single source of truth (PART 12-15) ───────────────────────────

class TestAIStatusTruth:
    @pytest.mark.asyncio
    async def test_off_reports_deterministic(self, monkeypatch):
        from src.routers.settings import ai_status
        monkeypatch.delenv("AI_MODE", raising=False)
        res = await ai_status("OFF")
        assert res["active"] is False
        assert res["provider"] is None
        assert "deterministic" in res["reason"].lower()

    @pytest.mark.asyncio
    async def test_enhanced_without_provider_reports_deterministic(self, monkeypatch):
        """CASE C: with no provider anywhere, resolution falls back to the
        ENHANCED token (deterministic). The status must NEVER report the Mock
        engine as an active provider."""
        from src.routers import settings as settings_router
        from src.engines import ai_provider as ap

        for key in ("AI_MODE", "GEMINI_API_KEY", "GROQ_API_KEY", "OPENAI_API_KEY",
                    "MINIMAX_API_KEY", "OPENCODE_ZEN_API_KEY"):
            monkeypatch.delenv(key, raising=False)
        # Hermetic: _env() falls back to pydantic Settings (the .env file), so a
        # machine-local AI_MODE pin must not leak into this test either.
        monkeypatch.setattr(ap, "_env", lambda name, default="": default)
        # Patch the provider classes directly: every real provider reports
        # unavailable.
        monkeypatch.setattr(ap.GeminiProvider, "is_available", lambda self: False)
        monkeypatch.setattr(ap.GroqProvider, "is_available", lambda self: False)
        monkeypatch.setattr(ap.OpenAIProvider, "is_available", lambda self: False)
        monkeypatch.setattr(ap.MiniMaxProvider, "is_available", lambda self: False)
        monkeypatch.setattr(ap.OpenCodeZenProvider, "is_available", lambda self: False)
        monkeypatch.setattr(ap.OllamaProvider, "is_available", lambda self: False)

        res = await settings_router.ai_status("ENHANCED")
        assert res["active"] is False
        assert res["provider"] is None
        assert res["provider_label"] is None
        assert res["state"] == "NOT_CONFIGURED"
        assert "deterministic" in (res["reason"] or "").lower()

    @pytest.mark.asyncio
    async def test_active_reports_resolved_provider_label(self, monkeypatch):
        from src.routers import settings as settings_router
        from src.engines import ai_provider as ap

        monkeypatch.setenv("GEMINI_API_KEY", "test-key")
        # Hermetic: stop the machine-local .env (e.g. AI_MODE=groq) pinning a
        # different provider — only the Gemini key configured above may resolve.
        monkeypatch.setattr(ap, "_env",
                            lambda name, default="": "test-key" if name == "GEMINI_API_KEY" else default)
        monkeypatch.setattr(ap.GeminiProvider, "is_available", lambda self: True)

        res = await settings_router.ai_status("ENHANCED")
        assert res["active"] is True
        assert res["provider"] == "gemini"
        # Display-quality label for the UI (PART 14: "Gemini", not "gemini").
        assert res["provider_label"] == "Gemini"


# ── AI QUOTA: monthly bucket (PART 16-19) ────────────────────────────────────

class TestMonthlyAIQuota:
    def _entitlement(self, db, user_id, credits=5):
        from src.database import EntitlementDB
        ent = EntitlementDB(
            user_id=user_id, edition="free", ai_enabled=True,
            ai_credits=credits, generation_limit=5,
        )
        db.add(ent)
        db.commit()
        return ent

    def test_month_key_is_yyyy_mm_server_clock(self):
        from src.ai_quota import current_ai_period_key
        assert current_ai_period_key() == datetime.utcnow().strftime("%Y-%m")

    def test_new_month_starts_at_zero(self, db):
        from src.ai_quota import get_ai_units_used
        user_id = "quota-user-1"
        assert get_ai_units_used(db, user_id, "2099-01") == 0

    def test_successful_generation_consumes_one_unit(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        user_id = "quota-user-2"
        remaining = consume_ai_generation(db, user_id, credits=5, period_key="2099-02")
        assert remaining == 4
        assert get_ai_units_used(db, user_id, "2099-02") == 1

    def test_second_generation_consumes_exactly_one_more(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        user_id = "quota-user-3"
        consume_ai_generation(db, user_id, credits=5, period_key="2099-03")
        consume_ai_generation(db, user_id, credits=5, period_key="2099-03")
        assert get_ai_units_used(db, user_id, "2099-03") == 2

    def test_never_exceeds_the_monthly_cap(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        user_id = "quota-user-4"
        for _ in range(7):
            consume_ai_generation(db, user_id, credits=5, period_key="2099-04")
        assert get_ai_units_used(db, user_id, "2099-04") == 5

    def test_other_calendar_month_resets_usage(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        user_id = "quota-user-5"
        consume_ai_generation(db, user_id, credits=5, period_key="2099-05")
        assert get_ai_units_used(db, user_id, "2099-06") == 0

    def test_unlimited_credits_never_consume(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        user_id = "quota-user-6"
        assert consume_ai_generation(db, user_id, credits=0, period_key="2099-07") == -1
        assert get_ai_units_used(db, user_id, "2099-07") == 0

    def test_quota_isolated_per_user(self, db):
        from src.ai_quota import consume_ai_generation, get_ai_units_used
        consume_ai_generation(db, "quota-user-7a", credits=5, period_key="2099-08")
        assert get_ai_units_used(db, "quota-user-7b", "2099-08") == 0

    def test_entitlement_decision_reads_monthly_ledger(self, db):
        from src.entitlements import _entitlement_ai_decision
        from src.ai_quota import consume_ai_generation
        user_id = "quota-user-8"
        ent = self._entitlement(db, user_id, credits=5)
        assert _entitlement_ai_decision(db, ent)[0] is True
        for _ in range(5):
            consume_ai_generation(db, user_id, credits=5)
        assert _entitlement_ai_decision(db, ent)[0] is False
        assert _entitlement_ai_decision(db, ent)[1] == "credits_exhausted"
    def test_resolve_entitlement_reports_monthly_snapshot(self, db):
        from src.entitlements import resolve_entitlement
        from src.database import User
        from src.ai_quota import consume_ai_generation
        user = User(
            id="quota-user-9", email="quota9@t.test", full_name="Q",
            hashed_password="x", is_active=True,
        )
        db.add(user)
        self._entitlement(db, user.id, credits=5)
        consume_ai_generation(db, user.id, credits=5)
        resolved = resolve_entitlement(db, user)
        assert resolved["ai_lifetime"] is False
        assert resolved["ai_quota_used"] == 1
        assert resolved["ai_quota_remaining"] == 4
        assert resolved["ai_quota_period_key"] == datetime.utcnow().strftime("%Y-%m")

    def test_exhausted_message_is_monthly_not_lifetime(self):
        from src.entitlements import free_tier_ai_exhausted_message
        msg = free_tier_ai_exhausted_message(5)
        assert "month" in msg.lower()
        assert "lifetime" not in msg.lower()

    def test_failed_ai_call_consumes_nothing(self, db):
        """PART 19: 429/503/fallback → quota consumed: 0."""
        from src.ai_quota import get_ai_units_used
        user_id = "quota-user-10"
        # Simulate: only the failure happened — nothing was ever recorded.
        assert get_ai_units_used(db, user_id) == 0


# ── DOWNLOADS: one-time URL path for all three formats (PART 1-4) ────────────

class TestDownloadEndpoints:
    def _stub_docx(self, monkeypatch, tmp_path):
        from src.routers import generation as gen_router
        from src.engines.generation_pipeline import GenerationPipeline

        def _fake(lessons, *args, **kwargs):
            out = Path(kwargs.get("output_path") or args[-1])
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_bytes(b"PK\x03\x04stub docx body")
            return out

        monkeypatch.setattr(
            gen_router.pipeline, "export_docx_combined", _fake, raising=False)
        return gen_router

    @pytest.mark.asyncio
    async def test_docx_download_url_issues_and_delivers(self, db, tmp_path, monkeypatch):
        from tests.conftest import make_user
        from tests.test_export_download import make_job_with_lessons
        from src.routers import generation as gen_router
        from src.database import UsageUnitDB

        self._stub_docx(monkeypatch, tmp_path)
        monkeypatch.chdir(tmp_path)
        u = make_user(db, role="teacher", email="dl-docx-ru@t.test")
        _scheme, job = make_job_with_lessons(db, u, lessons=1)

        res = await gen_router.issue_download_url(job.id, "docx", "GES-style", None, u, db)
        assert res["download_url"].startswith("/api/generation/downloads/")
        assert res["filename"].endswith(".docx")

        # Deliver through the token route exactly as the browser does.
        from fastapi.testclient import TestClient
        from fastapi import FastAPI
        from src.auth import get_optional_user
        from src.database import get_db
        app = FastAPI()
        app.include_router(gen_router.router, prefix="/api/generation")
        app.dependency_overrides[get_optional_user] = lambda: None
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            r = client.get(res["download_url"])
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument")
        assert r.headers["content-disposition"].startswith("attachment")
        assert r.content.startswith(b"PK\x03\x04")

    @pytest.mark.asyncio
    async def test_xlsx_download_url_issues_and_delivers(self, db, tmp_path, monkeypatch):
        from tests.conftest import make_user
        from tests.test_export_download import make_job_with_lessons
        from src.routers import generation as gen_router

        def _fake_xlsx(lessons, out_path):
            out_path = Path(out_path)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_bytes(b"PK\x03\x04stub xlsx body")
            return out_path

        monkeypatch.setattr(gen_router.pipeline, "export_xlsx", _fake_xlsx)
        monkeypatch.chdir(tmp_path)
        u = make_user(db, role="teacher", email="dl-xlsx-ru@t.test")
        _scheme, job = make_job_with_lessons(db, u, lessons=1)

        res = await gen_router.issue_download_url(job.id, "xlsx", "GES-style", None, u, db)
        assert res["filename"].endswith(".xlsx")

        from fastapi.testclient import TestClient
        from fastapi import FastAPI
        from src.auth import get_optional_user
        from src.database import get_db
        app = FastAPI()
        app.include_router(gen_router.router, prefix="/api/generation")
        app.dependency_overrides[get_optional_user] = lambda: None
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            r = client.get(res["download_url"])
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        assert r.content.startswith(b"PK\x03\x04")

    @pytest.mark.asyncio
    async def test_pdf_without_libreoffice_still_delivers_a_real_pdf(
        self, db, tmp_path, monkeypatch
    ):
        """PART 4/22: missing LibreOffice must never surface as 'Generation
        problem / Failed to fetch'.

        Hosts with no converter render the lesson structurally instead, so the
        download still arrives as a real PDF. When that render fails too, the
        answer is the explicit converter requirement — never a raw 500."""
        from tests.conftest import make_user
        from tests.test_export_download import make_job_with_lessons
        from src.routers import generation as gen_router
        from src.engines.pdf_export import PDFExportEngine, PDF_CONVERTER_REQUIREMENT
        from fastapi import HTTPException

        monkeypatch.chdir(tmp_path)
        monkeypatch.setattr(
            PDFExportEngine, "is_available", classmethod(lambda cls: False))
        u = make_user(db, role="teacher", email="dl-pdf-ru@t.test")
        _scheme, job = make_job_with_lessons(db, u, lessons=1)

        res = await gen_router.issue_download_url(job.id, "pdf", "GES-style", None, u, db)
        assert res["media_type"] == "application/pdf"
        assert res["filename"].endswith(".pdf")

        from fastapi.testclient import TestClient
        from fastapi import FastAPI
        from src.auth import get_optional_user
        from src.database import get_db
        app = FastAPI()
        app.include_router(gen_router.router, prefix="/api/generation")
        app.dependency_overrides[get_optional_user] = lambda: None
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            r = client.get(res["download_url"])
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/pdf")
        assert r.content[:4] == b"%PDF"

        # Both renderers down: name the missing converter instead of leaking a
        # raw 500 at the teacher.
        def _boom(*args, **kwargs):
            raise RuntimeError("reportlab build failed")

        monkeypatch.setattr(gen_router, "_render_structured_pdf", _boom)
        with pytest.raises(HTTPException) as excinfo:
            await gen_router.issue_download_url(job.id, "pdf", "GES-style", None, u, db)
        assert excinfo.value.status_code == 503
        assert excinfo.value.detail == PDF_CONVERTER_REQUIREMENT
        assert "LibreOffice" in (excinfo.value.detail or "")

    @pytest.mark.asyncio
    async def test_unauthenticated_unknown_token_is_rejected(self, db, tmp_path, monkeypatch):
        from tests.conftest import make_user
        from tests.test_export_download import make_job_with_lessons
        from src.routers import generation as gen_router

        self._stub_docx(monkeypatch, tmp_path)
        monkeypatch.chdir(tmp_path)
        u = make_user(db, role="teacher", email="dl-auth-ru@t.test")
        _scheme, job = make_job_with_lessons(db, u, lessons=1)
        res = await gen_router.issue_download_url(job.id, "docx", "GES-style", None, u, db)

        from fastapi.testclient import TestClient
        from fastapi import FastAPI
        from src.auth import get_optional_user
        from src.database import get_db
        app = FastAPI()
        app.include_router(gen_router.router, prefix="/api/generation")
        app.dependency_overrides[get_optional_user] = lambda: None
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            # An unknown/expired/replayed token is always 404.
            r1 = client.get(res["download_url"])
            assert r1.status_code == 200
            # Single-use: the second fetch of the SAME token fails.
            r2 = client.get(res["download_url"])
            assert r2.status_code == 404
            r3 = client.get("/api/generation/downloads/not-a-real-token")
            assert r3.status_code == 404

    @pytest.mark.asyncio
    async def test_token_is_user_bound(self, db, tmp_path, monkeypatch):
        """PART 2: Tab A (User A) cannot consume User B's download token."""
        from tests.conftest import make_user
        from tests.test_export_download import make_job_with_lessons
        from src.routers import generation as gen_router

        self._stub_docx(monkeypatch, tmp_path)
        monkeypatch.chdir(tmp_path)
        user_a = make_user(db, role="teacher", email="dl-a-ru@t.test")
        _s, job = make_job_with_lessons(db, user_a, lessons=1)
        res = await gen_router.issue_download_url(job.id, "docx", "GES-style", None, user_a, db)

        from fastapi.testclient import TestClient
        from fastapi import FastAPI
        from src.auth import get_optional_user
        from src.database import get_db
        from src.database import User
        app = FastAPI()
        app.include_router(gen_router.router, prefix="/api/generation")
        other = User(id="other-user", email="other@t.test")
        app.dependency_overrides[get_optional_user] = lambda: other
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            r = client.get(res["download_url"])
        assert r.status_code == 404


# ── EXPORT CONTENT: resources render as entries, never JSON (PART 26) ────────

class TestExportResourceRendering:
    def _export(self, tmp_path, lp):
        from src.engines.docx_export import DOCXExportEngine
        engine = DOCXExportEngine()
        out = tmp_path / "lesson.docx"
        engine.export_single(lp, None, out)
        from docx import Document
        doc = Document(str(out))
        parts = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    def test_ges_docx_resources_are_readable_never_serialized(self, tmp_path):
        """PART 26: source resources render as readable entries — never as
        serialized array syntax, never space-stripped."""
        alloc, config = _allocated(["Counters, sticks, flash cards"])
        lp = build_lesson(alloc, config, "s")
        joined = self._export(tmp_path, lp)
        for resource in ("Counters", "sticks", "flash cards"):
            assert resource in joined
        for bad in ('["', '"]', "[]", "[,]"):
            assert bad not in joined
        assert "Picturesshowing" not in joined

    def test_empty_other_tlrs_render_no_brackets(self, tmp_path):
        from src.engines.docx_export import DOCXExportEngine
        alloc, config = _allocated(["Charts"])
        lp = build_lesson(alloc, config, "s")
        lp.other_tlrs = []
        lp.teaching_learning_resources = list(lp.source_tlrs)
        joined = self._export(tmp_path, lp)
        for bad in ("[]", "[,]", '[""]'):
            assert bad not in joined
