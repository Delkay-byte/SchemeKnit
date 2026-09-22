"""
Multi-subject document detection tests (§4).

A scheme of learning may be ONE document containing SEVERAL subjects for the
same level. SchemeKnit must identify the subject sections, let the teacher
confirm one, and never silently pick a subject.
"""

import asyncio
from pathlib import Path

import pytest
from docx import Document

from src.parsers.docx_parser import DOCXParser
from src.parsers.pdf_parser import PDFParser

HEADER = ["Week Ending", "Strand", "Sub-Strand", "Content Standard", "Indicators"]


def _week_row(n: int, code: str) -> list:
    return [
        f"{n}\n{10 + n}/09/2026",
        "Diversity of Matter",
        "Elements",
        f"B7.1.{n}.1 Content standard",
        f"{code} Indicator {n}",
    ]


def _add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(HEADER))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    return table


def build_docx(path: Path, sections):
    """sections: list of (heading, [(week_no, code), ...]). None heading ⇒ none."""
    doc = Document()
    doc.add_paragraph("BASIC 7 SCHEME OF LEARNING")
    for heading, weeks in sections:
        if heading:
            doc.add_paragraph(heading)
        rows = [HEADER] + [_week_row(n, code) for n, code in weeks]
        _add_table(doc, rows)
    doc.save(str(path))
    return path


def _pdf_lines(sections):
    lines = ["BASIC 7 SCHEME OF LEARNING"]
    for heading, weeks in sections:
        if heading:
            lines.append(heading)
        lines.append(" | ".join(HEADER))
        for n, code in weeks:
            lines.append(
                f"{n} {10 + n}/09/2026 | Diversity of Matter | Elements | "
                f"B7.1.{n}.1 Content standard | {code} Indicator {n}"
            )
    return lines


def build_pdf(path: Path, sections):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for line in _pdf_lines(sections):
        page.insert_text((50, y), line, fontsize=9)
        y += 13
    doc.save(str(path))
    doc.close()
    return path


MULTI = [
    ("ENGLISH LANGUAGE", [(1, "B7.1.1.1"), (2, "B7.1.2.1")]),
    ("MATHEMATICS", [(1, "B7.2.1.1"), (2, "B7.2.2.1")]),
    ("SCIENCE", [(1, "B7.3.1.1"), (2, "B7.3.2.1")]),
]

SINGLE = [("SCIENCE", [(1, "B7.3.1.1"), (2, "B7.3.2.1")])]


# ── DOCX ────────────────────────────────────────────────────────────────

class TestDocxDetection:
    def test_multi_subject_docx(self, tmp_path):
        path = build_docx(tmp_path / "multi.docx", MULTI)
        info = DOCXParser().analyze(path)
        assert info["detection_status"] == "multiple"
        assert set(info["detected_subjects"]) == {"English Language", "Mathematics", "Science"}

    def test_single_subject_docx(self, tmp_path):
        path = build_docx(tmp_path / "single.docx", SINGLE)
        info = DOCXParser().analyze(path)
        assert info["detection_status"] == "single"
        assert info["detected_subjects"] == ["Science"]

    def test_low_confidence_docx(self, tmp_path):
        # Weeks exist but no subject heading was recognised: the file is
        # extractable, yet the subject is unconfirmed. Part E/H report this
        # honestly as "needs_confirmation" (previously "low_confidence").
        path = build_docx(tmp_path / "plain.docx", [(None, [(1, "B7.3.1.1")])])
        info = DOCXParser().analyze(path)
        assert info["detection_status"] == "needs_confirmation"
        assert info["needs_confirmation"] is True

    def test_document_title_detected(self, tmp_path):
        path = build_docx(tmp_path / "titled.docx", MULTI)
        info = DOCXParser().analyze(path)
        assert "SCHEME OF LEARNING" in info["title"].upper()

    def test_confirm_subject_extracts_only_that_section(self, tmp_path):
        path = build_docx(tmp_path / "multi2.docx", MULTI)
        scheme = asyncio.run(DOCXParser().parse(path, target_subject="Science"))
        assert scheme.subject.value == "Science"
        codes = [c for w in scheme.weeks for c in w.indicators]
        assert codes, "science section should contain indicators"
        assert all("B7.3." in c for c in codes)
        # No maths/english content leaked in.
        assert not any("B7.2." in c or "B7.1.1.1 Indicator" in c for c in codes)

    def test_confirm_maths_section(self, tmp_path):
        path = build_docx(tmp_path / "multi3.docx", MULTI)
        scheme = asyncio.run(DOCXParser().parse(path, target_subject="Mathematics"))
        codes = [c for w in scheme.weeks for c in w.indicators]
        assert all("B7.2." in c for c in codes)


# ── PDF ─────────────────────────────────────────────────────────────────

class TestPdfDetection:
    def test_multi_subject_pdf(self, tmp_path):
        path = build_pdf(tmp_path / "multi.pdf", MULTI)
        info = PDFParser().analyze(path)
        assert info["detection_status"] == "multiple"
        assert set(info["detected_subjects"]) == {"English Language", "Mathematics", "Science"}

    def test_single_subject_pdf(self, tmp_path):
        path = build_pdf(tmp_path / "single.pdf", SINGLE)
        info = PDFParser().analyze(path)
        assert info["detection_status"] == "single"

    def test_parse_single_subject_pdf(self, tmp_path):
        path = build_pdf(tmp_path / "single2.pdf", SINGLE)
        scheme = asyncio.run(PDFParser().parse(path))
        assert len(scheme.weeks) == 2
        assert scheme.subject.value == "Science"

    def test_confirm_subject_pdf(self, tmp_path):
        path = build_pdf(tmp_path / "multi2.pdf", MULTI)
        scheme = asyncio.run(PDFParser().parse(path, target_subject="Science"))
        codes = [c for w in scheme.weeks for c in w.indicators]
        assert codes and all("B7.3." in c for c in codes)

    def test_low_confidence_pdf(self, tmp_path):
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 60), "Just some prose with no curriculum table.")
        out = tmp_path / "plain.pdf"
        doc.save(str(out))
        doc.close()
        info = PDFParser().analyze(out)
        # A prose-only PDF with no extractable curriculum table is an honest
        # extraction failure (Part J) — never presented as a usable scheme.
        assert info["detection_status"] == "extraction_failed"
