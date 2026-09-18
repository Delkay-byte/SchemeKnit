"""
Real Ghana scheme acceptance — BASIC 9 SCIENCE SCHEME OF LEARNING.docx (item 5).

No synthetic data: the scheme is parsed with the production DOCXParser, lessons
are generated with the production GenerationPipeline (AI OFF so nothing is
invented), rendered through the approved template renderer, and compared
structurally against the golden master of the approved source document.

Skipped cleanly when the uploaded scheme is not present (local-run repo), so the
suite stays green on machines without the acceptance upload.
"""

import os
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from tests.test_official_ges_template import (  # noqa: E402
    JHS_GUIDANCE,
    JHS_SPEC,
)

SCHEME_FILENAME = "1d3813d1-8fce-47b1-ac81-2ab40c5a8bb8_BASIC 9 SCIENCE SCHEME OF LEARNING.docx"
SCHEME_PATH = Path(__file__).parent.parent / "uploads" / SCHEME_FILENAME

pytestmark = pytest.mark.skipif(
    not SCHEME_PATH.exists(), reason="real Basic 9 Science scheme not uploaded")


@pytest.fixture(scope="module")
def real_scheme():
    import asyncio

    from src.parsers.docx_parser import DOCXParser

    parser = DOCXParser()
    return asyncio.run(parser.parse(SCHEME_PATH, "BASIC 9 SCIENCE SCHEME OF LEARNING.docx"))


@pytest.fixture(scope="module")
def generated(real_scheme):
    """Real generation run: pipeline, AI OFF, one week, approved template."""
    from src.models import AIMode, ClassLevel, Subject, TermConfig
    from src.engines.generation_pipeline import GenerationPipeline

    config = TermConfig(
        scheme_of_work_id=real_scheme.id,
        academic_year=real_scheme.academic_year,
        term=real_scheme.term,
        class_level=ClassLevel.BASIC_9,
        subject=Subject.SCIENCE,
        class_size=24,
        lesson_duration_minutes=60,
        lessons_per_week=2,
        term_start_date=real_scheme.weeks[0].start_date,
        term_end_date=real_scheme.weeks[-1].end_date,
        teaching_days=[0, 2],
        holidays=[],
        school_name="Acceptance Basic School",
        teacher_name="Acceptance Teacher",
        ai_mode=AIMode.OFF,
        include_special_weeks=False,
    )
    job = GenerationPipeline().generate_all(
        real_scheme, config, include_special_weeks=False, selected_weeks=[1])
    assert job.status == "completed", job.error_message
    lessons = job._lesson_plans
    assert lessons, "pipeline produced no lesson plans"
    return lessons


class TestRealSchemeExtraction:
    def test_real_curriculum_values(self, real_scheme):
        """The parser must extract the real scheme's curriculum data."""
        assert real_scheme.subject.value == "Science"
        assert real_scheme.class_level.value == "Basic 9"
        assert len(real_scheme.weeks) >= 12
        week1 = real_scheme.weeks[0]
        assert week1.strand == "Diversity of Matter"
        assert week1.sub_strand == "Materials"
        assert any("B9.1.1.1" in cs for cs in week1.content_standards)
        assert any("B9.1.1.1.1" in ind for ind in week1.indicators)
        assert week1.resources


class TestRealGeneration:
    def test_lessons_carry_the_real_curriculum(self, generated, real_scheme):
        week1 = real_scheme.weeks[0]
        for lp in generated:
            assert lp.strand == week1.strand
            assert lp.sub_strand == week1.sub_strand
            assert "B9.1.1.1" in (lp.content_standard or "")
            assert lp.class_level.value == "Basic 9"
            assert lp.subject.value == "Science"
            # AI OFF: no invented activities beyond deterministic seeding.
            assert lp.introduction or lp.main_activities

    def test_real_docx_structurally_matches_approved_source(self, generated, tmp_path):
        """Item 6: exact structural comparison APPROVED SOURCE vs GENERATED.

        Compared against the source .docx itself (topology, merges, label and
        phase coordinates) plus content hygiene — not merely "all fields filled".
        """
        from src.engines.docx_export import DOCXExportEngine
        from src.validators.docx_structure_compare import (
            golden_master_source_path, validate_generated,
        )

        DOCXExportEngine().export_batch(
            generated, None, tmp_path)  # default template resolves from class level
        outs = list(tmp_path.glob("*.docx"))
        assert outs
        out = outs[0]

        report = validate_generated(
            out,
            expected_values=["Basic 9", "Diversity of Matter", "Materials",
                             "B9.1.1.1"],
            forbidden_values=JHS_GUIDANCE + ["Responding to questions, brainstorming",
                                             "Forces & Energy", "B7.4.3.1"],
            source_path=golden_master_source_path(),
        )
        failures = [(g, c["name"], c["detail"])
                    for g in ("structure", "content")
                    for c in report[g]["checks"] if not c["pass"]]
        assert report["pass"], failures

    def test_no_invented_curriculum_codes(self, generated, tmp_path):
        """Values on the plan come from the scheme, not from the renderer."""
        from src.engines.docx_export import DOCXExportEngine
        from docx import Document

        out_dir = tmp_path / "codes"
        DOCXExportEngine().export_batch(generated, None, out_dir)
        out = next(out_dir.glob("*.docx"))
        blob = "\n".join(p.text for p in Document(str(out)).paragraphs)
        for t in Document(str(out)).tables:
            for row in t.rows:
                for cell in row.cells:
                    blob += "\n" + cell.text
        week1 = real_scheme_week1()
        assert "B9.1.1.1" in blob          # real code present
        assert "B99.9.9.9" not in blob     # no invented code pattern
        assert week1.strand in blob


def real_scheme_week1():
    import asyncio

    from src.parsers.docx_parser import DOCXParser

    parser = DOCXParser()
    scheme = asyncio.run(parser.parse(SCHEME_PATH, "BASIC 9 SCIENCE SCHEME OF LEARNING.docx"))
    return scheme.weeks[0]


class TestRealBatchAndZip:
    def test_zip_members_match_approved_structure(self, generated, tmp_path):
        """Item 7: at least one ZIP member is opened and structurally validated
        against the approved source."""
        from src.engines.zip_export import ZIPExportEngine
        from src.validators.docx_structure_compare import (
            golden_master_source_path, validate_generated,
        )

        out = tmp_path / "real_science_plans.zip"
        ZIPExportEngine().export_batch(generated, out)
        assert out.exists()

        with zipfile.ZipFile(out) as zf:
            members = [n for n in zf.namelist() if n.endswith(".docx")]
            assert members
            member = tmp_path / members[0]
            member.write_bytes(zf.read(members[0]))

        report = validate_generated(
            member,
            expected_values=["Basic 9", "Diversity of Matter", "B9.1.1.1"],
            forbidden_values=JHS_GUIDANCE + ["Forces & Energy", "B7.4.3.1"],
            source_path=golden_master_source_path(),
        )
        failures = [(g, c["name"], c["detail"])
                    for g in ("structure", "content")
                    for c in report[g]["checks"] if not c["pass"]]
        assert report["pass"], failures
