"""
Phase 16.6 Section A — subject-section detection (generic, no hardcoded lists).

Real Basic 6 / Basic 7 documents establish the census; a synthetic fixture
covers differently formatted headings (wrapped titles + in-table cell titles)
so the fix is not tied to one visual template.
"""
import asyncio
from pathlib import Path

import pytest
from docx import Document

from src.models import Subject
from src.parsers.docx_parser import DOCXParser
from src.parsers.pdf_parser import PDFParser, get_document_parser
from src.parsers.subject_keywords import canonical_subject_from_heading

REAL = Path(__file__).resolve().parent.parent / "real_documents"

B6_DOCX = REAL / "BASIC 6 TERM 1.docx"
B6_PDF = REAL / "BASIC 6 TERM 1.pdf"
B6_FRENCH = REAL / "BASIC 6 FRENCH SCHEME OF LEARNING.docx"
B7_DOCX = REAL / "BASIC 7 TERM 1.docx"
B7_PDF = REAL / "BASIC 7 TERM 1.pdf"

#: Subject section headings actually present in BASIC 6 TERM 1.docx
#: (derived from document structure: paragraph titles
#: "FIRST TERM SCHEME OF LEARNING FOR BASIC 6 - <SUBJECT>").
B6_EXPECTED_SUBJECTS = {
    "Science",
    "ICT",
    "English Language",
    "French",
    "Creative Arts and Design",
    "Mathematics",
    "Religious and Moral Education",
    "History",
    "Ghanaian Language",
}


pytestmark = pytest.mark.skipif(
    not B6_DOCX.exists(),
    reason="real Basic 6 document not present in backend/real_documents",
)


# ── Real document census ─────────────────────────────────────────────────

class TestRealBasic6Census:
    def test_b6_docx_census(self):
        info = DOCXParser().analyze(B6_DOCX, original_filename=B6_DOCX.name)
        detected = set(info["detected_subjects"])
        assert detected == B6_EXPECTED_SUBJECTS, (
            f"MISSING={sorted(B6_EXPECTED_SUBJECTS - detected)} "
            f"EXTRA={sorted(detected - B6_EXPECTED_SUBJECTS)}"
        )
        assert info["detection_status"] == "multiple"
        meta = info["metadata"]
        assert meta["class_level"] == "Basic 6"
        assert meta["class_signals"]["conflict"] is False

    def test_b6_docx_sections_have_weeks(self):
        info = DOCXParser().analyze(B6_DOCX, original_filename=B6_DOCX.name)
        by_subject = {s["subject"]: s for s in info["sections"]}
        assert set(by_subject) == B6_EXPECTED_SUBJECTS
        for name, section in by_subject.items():
            assert section["week_count"] >= 12, (
                f"{name} only has {section['week_count']} weeks"
            )

    @pytest.mark.skipif(not B6_PDF.exists(), reason="B6 pdf missing")
    def test_b6_pdf_mirrors_docx_census(self):
        info = PDFParser().analyze(B6_PDF, original_filename=B6_PDF.name)
        assert set(info["detected_subjects"]) == B6_EXPECTED_SUBJECTS
        assert info["metadata"]["class_level"] == "Basic 6"

    def test_b6_french_single_subject_not_basic9(self):
        info = DOCXParser().analyze(
            B6_FRENCH, original_filename=B6_FRENCH.name
        )
        assert info["detection_status"] == "single"
        assert info["detected_subjects"] == ["French"]
        assert info["metadata"]["class_level"] == "Basic 6"
        assert info["metadata"]["class_signals"]["conflict"] is False
        sections = info["sections"]
        assert sections and sections[0]["week_count"] == 15

    @pytest.mark.skipif(not B7_DOCX.exists(), reason="B7 docx missing")
    def test_b7_docx_keeps_prior_subjects_plus_wrapped_titles(self):
        info = DOCXParser().analyze(B7_DOCX, original_filename=B7_DOCX.name)
        detected = set(info["detected_subjects"])
        # Previously detected subset must remain; wrapped English/French titles
        # are now recognised too.
        assert {
            "Science", "ICT", "Creative Arts and Design", "Ghanaian Language",
            "Mathematics", "Religious and Moral Education",
            "Career Technology", "Social Studies",
        } <= detected
        assert "English Language" in detected
        assert info["metadata"]["class_level"] == "Basic 7"

    @pytest.mark.skipif(not B7_PDF.exists(), reason="B7 pdf missing")
    def test_b7_pdf_class_preserved(self):
        info = PDFParser().analyze(B7_PDF, original_filename=B7_PDF.name)
        assert info["metadata"]["class_level"] == "Basic 7"
        assert "English Language" in info["detected_subjects"]


# ── Heading normalisation (shared primitive) ─────────────────────────────

class TestHeadingNormalisation:
    def test_wrapped_district_title_english(self):
        assert canonical_subject_from_heading(
            "FIRST TERM SCHEME OF LEARNING FOR BASIC 6 - ENGLISH LANGUAGE"
        ) is Subject.ENGLISH

    def test_wrapped_district_title_french(self):
        assert canonical_subject_from_heading(
            "FIRST TERM SCHEME OF LEARNING FOR BASIC 6 - FRENCH LANGUAGE"
        ) is Subject.FRENCH

    def test_bare_subject_still_matches(self):
        assert canonical_subject_from_heading("MATHEMATICS") is Subject.MATHEMATICS

    def test_prose_is_not_a_heading(self):
        assert canonical_subject_from_heading(
            "Learners explain the science of measurement in groups"
        ) is None

    def test_long_text_rejected(self):
        assert canonical_subject_from_heading("Science " + "detail " * 40) is None


# ── Synthetic differently-formatted multi-subject fixture ────────────────

HEADER = ["Week Ending", "Strand", "Sub-Strand", "Content Standard", "Indicators"]


def _row(n: int, code: str):
    return [
        f"{n}\n{10 + n}/09/2026",
        "Diversity of Matter",
        "Elements",
        f"Content standard {n}",
        f"{code} Indicator {n}",
    ]


def _add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(HEADER))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value


def build_synthetic(path: Path) -> Path:
    """Multi-subject DOCX with headings the old matcher/template missed.

    - Paragraph title uses the district wrapper pattern (FOR … - SUBJECT).
    - One subject title lives ONLY inside a table row (cell heading).
    - Ordinary curriculum prose sits between sections (must not split).
    """
    doc = Document()
    doc.add_paragraph("TERM 1 SCHEME OF LEARNING 2026/2027")
    doc.add_paragraph(
        "SECOND TERM SCHEME OF LEARNING FOR BASIC 8 - SOCIAL STUDIES"
    )
    _add_table(doc, [HEADER] + [_row(1, "B8.1.1.1"), _row(2, "B8.1.2.1")])
    doc.add_paragraph(
        "This week learners review the science of measurement in groups."
    )
    # Cell-only subject heading: single short cell row, then data.
    table = doc.add_table(rows=3, cols=len(HEADER))
    for c, value in enumerate(["CAREER TECHNOLOGY"] + [""] * (len(HEADER) - 1)):
        table.cell(0, c).text = value
    for c, value in enumerate(HEADER):
        table.cell(1, c).text = value
    for c, value in enumerate(_row(1, "B8.2.1.1")):
        table.cell(2, c).text = value
    doc.add_paragraph("THE SCIENCE OF MEASUREMENT IS FUN")  # prose-ish, not a subject
    _add_table(doc, [HEADER] + [_row(1, "B8.3.1.1")])
    doc.add_paragraph("HISTORY")
    _add_table(doc, [HEADER] + [_row(1, "B8.4.1.1")])
    doc.save(str(path))
    return path


class TestSyntheticFormatVariant:
    @pytest.fixture
    def synthetic(self, tmp_path):
        return build_synthetic(tmp_path / "variant_multi.docx")

    def test_detects_wrapped_cell_and_bare_headings(self, synthetic):
        info = DOCXParser().analyze(synthetic, original_filename=synthetic.name)
        detected = set(info["detected_subjects"])
        assert "Social Studies" in detected
        assert "Career Technology" in detected, "table-cell heading not detected"
        assert "History" in detected
        assert info["detection_status"] == "multiple"

    def test_prose_does_not_create_false_section(self, synthetic):
        info = DOCXParser().analyze(synthetic, original_filename=synthetic.name)
        # "science of measurement" prose must not open a Science section.
        assert "Science" not in set(info["detected_subjects"])

    def test_target_subject_isolation_on_synthetic(self, synthetic):
        scheme = asyncio.run(
            DOCXParser().parse(synthetic, target_subject="Social Studies")
        )
        codes = [c for w in scheme.weeks for c in w.indicators]
        assert codes and all("B8.1." in c for c in codes)
        assert not any("B8.2." in c or "B8.4." in c for c in codes)


class TestGetDocumentParserRouting:
    def test_docx_routes_to_docx_parser(self):
        if not B6_DOCX.exists():
            pytest.skip("b6 missing")
        parser = get_document_parser(B6_DOCX)
        assert isinstance(parser, DOCXParser)
