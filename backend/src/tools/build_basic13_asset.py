"""
Build the bundled Approved WAPEF Basic 1-3 weekly class-plan template asset.

SOURCE OF TRUTH
---------------
``WAPEF BASIC 1 PLAN.docx`` — the weekly class-teacher plan supplied by WAPEF
(bundled verbatim at ``backend/tests/fixtures/wapef/WAPEF BASIC 1 PLAN.docx``).
One weekly class plan holds one table per subject section:

    metadata rows (Week Ending … Key words)
    DAYS | PHASE 1: STARTER | PHASE 2: MAIN | PHASE 3: REFLECTION
    one row per teaching day (MONDAY … FRIDAY, or a grouped "MONDAY & THURSDAY")

This script derives the *token* asset from that document — the same approach the
Approved WAPEF Plan and the official GES forms use: the source's own table
topology (5 columns, merged metadata value cells, the printed phase headings) is
kept and only the VALUES are replaced with bracketed tokens. The weekly renderer
then fills the tokens in place, cloning the block once per subject, so a
generated weekly plan is topologically identical to the supplied source by
construction.

The asset is committed; regenerate with:

    python -m src.tools.build_basic13_asset

Nothing in the fixture's lesson prose survives into the asset (every value cell
becomes a token) and the school identity is removed, so no sample text can leak
into a generated plan.
"""

from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn

BACKEND_DIR = Path(__file__).resolve().parent.parent.parent
REPO_FIXTURE = (BACKEND_DIR / "tests" / "fixtures" / "wapef"
                / "WAPEF BASIC 1 PLAN.docx")
ASSET_PATH = (BACKEND_DIR / "src" / "engines" / "assets"
              / "wapef_basic13_weekly_template.docx")

#: The exported source document this asset is derived from (repo-relative),
#: quoted by the provenance registry and the acceptance report.
SOURCE_DOCUMENT = "backend/tests/fixtures/wapef/WAPEF BASIC 1 PLAN.docx"

#: Row index -> token for the plain "label | value" metadata rows of the source
#: section table. The label cell is kept verbatim; only the value changes.
VALUE_ROW_TOKENS = {
    0: "WEEK_ENDING",
    1: "CLASS",
    2: "SUBJECT",
    3: "REFERENCE",
    4: "CONTENT_STANDARD",
    5: "LEARNING_INDICATOR",
    6: "PERFORMANCE_INDICATOR",
    7: "STRAND",
    8: "SUB_STRAND",
    9: "RESOURCES",
}

#: Self-labelled rows: (row, cell) -> token. The cell's own label text
#: ("Throughlines:", "God's Story:", "Deep Hope :", "Storyline :",
#: "Core Competencies:", "Key words:") is preserved and the token appended, so
#: the labels stay exactly as the supplied plan prints them.
LABELLED_CELL_TOKENS = {
    (10, 0): "THROUGH_LINES",
    (10, 1): "GODS_STORY",
    (11, 0): "DEEP_HOPE",
    (11, 1): "STORYLINE",
    (12, 0): "CORE_COMPETENCIES",
    (13, 0): "KEY_WORDS",
}

#: The source's day rows: 5 teaching-day rows live after the DAYS header row.
FIRST_DAY_ROW = 16
DAY_ROW_COUNT = 5



def _unique_cells(row):
    """A row's cells with horizontal-merge duplicates collapsed (grid order).

    The source's metadata rows merge the label across one grid column group and
    the value across the rest, so ``row.cells[1]`` can be the LABEL's own tc.
    Addressing by the row's *distinct* cells (index 0 = label, 1 = value, and
    for day rows 0 = day label, 1..3 = starter/main/reflection) is the only
    addressing that follows the supplied document's own layout.
    """
    out, seen = [], set()
    for cell in row.cells:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        out.append(cell)
    return out


def _fill(paragraph, text: str) -> None:
    from src.engines.official_ges_template import _fill_paragraph

    _fill_paragraph(paragraph, text)


def _set_text(cell, text: str) -> None:
    """Replace a cell's whole content with ``text`` (first run keeps its format)."""
    from src.engines.official_ges_template import _fill_paragraph

    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        cell.text = text
        return
    _fill_paragraph(paragraphs[0], text)
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def _label_prefix(text: str) -> str:
    """Label part of a self-labelled cell ('Throughlines: Order…' -> 'Throughlines:')."""
    raw = " ".join((text or "").split())
    head, sep, _ = raw.partition(":")
    return f"{head}:" if sep else raw


def _find_section_table(document):
    """The source's own subject-section table (the one carrying Strand/Sub strand)."""
    for table in document.tables:
        has_strand = any((row.cells[0].text or "").strip().lower() == "strand"
                         for row in table.rows)
        if has_strand:
            return table
    raise SystemExit("No subject-section table found in the source document")


def build(source: Path = REPO_FIXTURE, dest: Path = ASSET_PATH) -> Path:
    if not source.is_file():
        raise SystemExit(f"Source document missing: {source}")

    document = Document(str(source))
    table = _find_section_table(document)

    # Keep ONLY the canonical subject-section table; the weekly renderer clones
    # it once per subject (the subject count is data, never hard-coded).
    keep = table._tbl
    for other in list(document.tables):
        if other._tbl is not keep:
            other._tbl.getparent().remove(other._tbl)

    # Header: school identity and the week number become tokens; the source's
    # own "LESSON PLAN" title line is preserved verbatim.
    paragraphs = list(document.paragraphs)
    if len(paragraphs) < 3:
        raise SystemExit("Source header paragraphs missing")
    _fill(paragraphs[0], "[SCHOOL_NAME]")
    _fill(paragraphs[1], "LESSON PLAN")
    _fill(paragraphs[2], "WEEK [WEEK]")

    # Stray page-break / blank paragraphs from the source are dropped: the
    # weekly renderer owns section separation.
    for paragraph in list(document.paragraphs):
        if paragraph._element.findall(".//" + qn("w:br")):
            paragraph._element.getparent().remove(paragraph._element)
        elif not (paragraph.text or "").strip():
            paragraph._element.getparent().remove(paragraph._element)

    rows = table.rows
    if len(rows) < FIRST_DAY_ROW + DAY_ROW_COUNT:
        raise SystemExit("Source section table has fewer day rows than expected")

    # The sample's week-2 module prints its own row label as "2nd Week Ending";
    # that row is the week-ending row, so the label is normalized to the form's
    # canonical wording (the "2nd" belongs to the sample's week number, not to
    # the template).
    week_label = " ".join((_unique_cells(rows[0])[0].text or "").split())
    if week_label.lower().endswith("week ending"):
        _set_text(_unique_cells(rows[0])[0], "Week Ending")

    for index, token in VALUE_ROW_TOKENS.items():
        cells = _unique_cells(rows[index])
        _set_text(cells[1], f"[{token}]")

    for (row_index, cell_index), token in LABELLED_CELL_TOKENS.items():
        cell = _unique_cells(rows[row_index])[cell_index]
        prefix = _label_prefix(cell.text)
        _set_text(cell, f"{prefix} [{token}]" if prefix else f"[{token}]")

    for offset in range(DAY_ROW_COUNT):
        cells = _unique_cells(rows[FIRST_DAY_ROW + offset])
        number = offset + 1
        if len(cells) < 4:
            raise SystemExit("Day row does not carry four cells")
        _set_text(cells[0], f"[DAY_LABEL_{number}]")
        for position, token in enumerate(("PHASE1", "PHASE2", "PHASE3"), start=1):
            _set_text(cells[position], f"[{token}_{number}]")

    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dest))
    return dest


def main() -> int:
    path = build()
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
