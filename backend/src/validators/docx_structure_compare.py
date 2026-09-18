"""Exact structural comparison between the approved source DOCX and a generated
SchemeKnit DOCX.

Acceptance rule (see docs/APPROVED_TEMPLATE_ACCEPTANCE.md): a generated lesson
plan is not acceptable merely because every field was filled. It must be
*structurally* the approved source's form — same table count and order, same row
and column counts, same merge map, same label and phase placement — with the
source's own guidance and example prose gone and real lesson data in its place.

Two independent checks run:

``compare_structure(source_fp, generated_fp)``
    Topology only: table count/order, per-table rows, per-table grid columns,
    the horizontal- and vertical-merge multisets, the coordinate of every key
    label, and the coordinate and text of every phase marker. Content is
    ignored, so this is the check that survives different lessons.

``compare_content(generated_fp, expected_values, forbidden_values)``
    Content hygiene: every expected lesson value present; no leftover bracketed
    token, no bare ``[``, no ``]``, no ``INSERT_``, none of the source's
    instruction-sheet text, none of the source's example bullets, and none of
    the forbidden foreign sample content.

Both are pure functions over python-docx; they never touch the network or AI.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from docx import Document  # type: ignore
from docx.oxml.ns import qn  # type: ignore

from ..engines.approved_template import APPROVED_PHASE_NAMES
from ..engines.official_ges_levels import JHS_SPEC
from ..engines.official_ges_template import GUIDANCE_PREFIX, load_document, strip_guidance

#: Labels whose grid coordinates must match the source exactly. Taken from the
#: golden master's own label cells.
KEY_LABELS: Sequence[str] = (
    "WEEK ENDING", "CLASS / LEVEL", "SUBJECT", "CLASS SIZE", "DATE / DAY",
    "PERIOD", "DURATION", "REFERENCE", "STRAND", "SUB-STRAND",
    "CONTENT STANDARD", "INDICATOR", "CORE COMPETENCIES",
    "PERFORMANCE INDICATOR", "KEYWORDS / VOCABULARY",
    "TEACHING & LEARNING RESOURCES (TLRs)",
)

#: The source's authoring-instruction phrase. Its presence in a generated plan
#: means the instruction sheet leaked onto a teacher's page.
GUIDANCE_PHRASES = (
    GUIDANCE_PREFIX,
    "Instructions for AI Engine",
    "THE CORE TRIDELIVERY LESSON TIMELINE TABLE",
    "Parse this table to extract",
    "Map these variables exactly",
    "Populate the lesson prerequisites",
    "Break down instructional actions",
)

#: Example bullets the source prints beside its phase tokens. A generated plan
#: must carry the teacher's real activities, never these.
SOURCE_EXAMPLE_PROSE = (
    "Review prior knowledge.",
    "Brain preparation hook.",
    "Step-by-step teacher instructions.",
    "Active individual or group work.",
    "Differentiated learning engagement tasks.",
    "Recapping key milestones.",
    "Post-lesson homework or exit ticket prompts.",
    "Teacher reflection space.",
)


def _grid_span(tc) -> int:
    tc_pr = tc.tcPr
    if tc_pr is None:
        return 1
    grid_span = tc_pr.find(qn("w:gridSpan"))
    try:
        return int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
    except (TypeError, ValueError):
        return 1


def _v_merge(tc) -> Optional[str]:
    tc_pr = tc.tcPr
    if tc_pr is None:
        return None
    v_merge = tc_pr.find(qn("w:vMerge"))
    if v_merge is None:
        return None
    val = v_merge.get(qn("w:val"))
    return "continue" if (val is None or val == "continue") else "restart"


def _table_topology(table) -> Dict[str, Any]:
    """Row count, grid column count, and the merge multiset of one table."""
    tbl = table._tbl
    rows_xml = tbl.findall(qn("w:tr"))
    grid = tbl.find(qn("w:tblGrid"))
    grid_cols = len(grid.findall(qn("w:gridCol"))) if grid is not None else len(table.columns)

    h_spans: List[int] = []
    v_events: List[str] = []
    cell_grid: List[List[str]] = []
    for tr in rows_xml:
        row_texts: List[str] = []
        for tc in tr.findall(qn("w:tc")):
            span = _grid_span(tc)
            vm = _v_merge(tc)
            if span > 1:
                h_spans.append(span)
            if vm:
                v_events.append(vm)
            text = " ".join("".join(p.itertext()) for p in tc.findall(qn("w:p")))
            row_texts.append(" ".join(text.split()))
            for _ in range(span - 1):
                row_texts.append("")  # covered slot
        cell_grid.append(row_texts)
    return {
        "rows": len(rows_xml),
        "grid_cols": grid_cols,
        "h_merges": sorted(h_spans),
        "v_merges": sorted(v_events),
        "cell_grid": cell_grid,
    }


def _document_topology(path) -> Dict[str, Any]:
    document = Document(str(path))
    return {
        "table_count": len(document.tables),
        "tables": [_table_topology(t) for t in document.tables],
    }


def _find_label(grid: List[List[str]], label: str) -> Optional[tuple]:
    """Grid coordinate (row, col) of a cell whose text starts with ``label``."""
    for r, row in enumerate(grid):
        for c, text in enumerate(row):
            if text and text.rstrip(":").strip().upper().startswith(label.upper()):
                return (r, c)
    return None


def _find_label_in_tables(tables: List[Dict[str, Any]], label: str) -> Optional[tuple]:
    """(table index, row, col) of the cell whose text starts with ``label``."""
    for t_index, table in enumerate(tables):
        at = _find_label(table["cell_grid"], label)
        if at is not None:
            return (t_index, at[0], at[1])
    return None


def _norm(text: str) -> str:
    return " ".join((text or "").split())


def compare_structure(source_path, generated_path) -> Dict[str, Any]:
    """Topology comparison of the approved source against a generated document.

    A multi-lesson document (one form per lesson, page-separated) is allowed:
    its table count must be an exact multiple of the source's and every block
    must match. Returns ``{"pass": bool, "checks": [...]}``; never raises.
    """
    checks: List[Dict[str, Any]] = []

    def check(name: str, ok: bool, detail: str = ""):
        checks.append({"name": name, "pass": bool(ok), "detail": detail})

    try:
        source = _document_topology(source_path)
        generated = _document_topology(generated_path)
    except Exception as exc:  # pragma: no cover - defensive
        check("documents_parseable", False, str(exc))
        return {"pass": False, "checks": checks}

    src_tables = source["tables"]
    out_tables = generated["tables"]
    n = len(src_tables)

    if len(out_tables) == n:
        blocks = [out_tables]
        check("table_count", True, f"source={n} generated={len(out_tables)}")
    elif n and len(out_tables) % n == 0:
        blocks = [out_tables[i:i + n] for i in range(0, len(out_tables), n)]
        check("table_count", True,
              f"source={n} generated={len(out_tables)} lessons={len(blocks)}")
    else:
        check("table_count", False, f"source={n} generated={len(out_tables)}")
        return {"pass": False, "checks": checks}

    for block_index, block in enumerate(blocks):
        prefix = f"L{block_index + 1}:" if len(blocks) > 1 else ""

        dims_ok = ([ (t["rows"], t["grid_cols"]) for t in block ]
                   == [ (t["rows"], t["grid_cols"]) for t in src_tables ])
        check(f"{prefix}table_dimensions", dims_ok,
              f"source={[(t['rows'], t['grid_cols']) for t in src_tables]} "
              f"output={[(t['rows'], t['grid_cols']) for t in block]}")

        merges_ok = ([ (t["h_merges"], t["v_merges"]) for t in block ]
                     == [ (t["h_merges"], t["v_merges"]) for t in src_tables ])
        check(f"{prefix}merge_map", merges_ok,
              "horizontal-span and vertical-merge multisets per table")

        # Key labels sit at the same (table, row, column) as the source.
        for label in KEY_LABELS:
            src_at = _find_label_in_tables(src_tables, label)
            out_at = _find_label_in_tables(block, label)
            check(f"{prefix}label:{label}",
                  src_at is not None and src_at == out_at,
                  f"source={src_at} output={out_at}")

        # Phase markers: same table, same row, same printed name.
        delivery = block[-1]["cell_grid"]
        for phase_index, phase_name in enumerate(APPROVED_PHASE_NAMES, start=1):
            found = None
            for r, row in enumerate(delivery):
                for c, text in enumerate(row):
                    if _norm(text).upper().startswith(phase_name.upper()):
                        found = (r, c)
            check(f"{prefix}phase:{phase_name}",
                  found is not None and found[0] == phase_index,
                  f"expected row {phase_index} in delivery table, found {found}")

    overall = all(c["pass"] for c in checks)
    return {"pass": overall, "checks": checks, "lessons": len(blocks)}


def compare_content(generated_path,
                    expected_values: Sequence[str] = (),
                    forbidden_values: Sequence[str] = ()) -> Dict[str, Any]:
    """Content hygiene of a generated document. Never raises."""
    checks: List[Dict[str, Any]] = []

    def check(name: str, ok: bool, detail: str = ""):
        checks.append({"name": name, "pass": bool(ok), "detail": detail})

    try:
        document = Document(str(generated_path))
    except Exception as exc:  # pragma: no cover - defensive
        check("document_parseable", False, str(exc))
        return {"pass": False, "checks": checks}

    parts: List[str] = []
    for paragraph in document.paragraphs:
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    blob = "\n".join(parts)
    flat = " ".join(blob.split())

    for value in expected_values:
        check(f"value_present:{_norm(value)[:40]}", _norm(value) in flat)

    for value in forbidden_values:
        check(f"foreign_absent:{_norm(value)[:40]}", _norm(value) not in flat)

    # No placeholder machinery may survive into a teacher's plan.
    check("no_placeholder_tokens", "[INSERT" not in blob and "[PHASE_" not in blob,
          "bracketed [INSERT_*] / [PHASE_*] tokens must be filled")
    check("no_brackets", "[" not in blob and "]" not in blob,
          "no bare [ or ] anywhere in the document")
    check("no_insert_marker", "INSERT_" not in blob)
    for phrase in GUIDANCE_PHRASES:
        check(f"no_guidance:{phrase[:30]}", phrase not in blob)
    for prose in SOURCE_EXAMPLE_PROSE:
        check(f"no_example_prose:{prose[:30]}", prose not in blob)

    overall = all(c["pass"] for c in checks)
    return {"pass": overall, "checks": checks}


def golden_master_source_path() -> Path:
    """Absolute path of the approved source document (the golden master)."""
    from ..engines.official_ges_template import template_path

    return template_path(JHS_SPEC)


def golden_master_fingerprint() -> Dict[str, Any]:
    """Normalized fingerprint of the approved source, guidance stripped.

    This is the structural contract: table count, order, dimensions and merge
    map of the form a generated plan must reproduce.
    """
    document = load_document(JHS_SPEC)
    strip_guidance(document)
    fingerprint = {
        "source": str(golden_master_source_path()),
        "table_count": len(document.tables),
        "tables": [_table_topology(t) for t in document.tables],
        "phase_names": list(APPROVED_PHASE_NAMES),
        "key_labels": list(KEY_LABELS),
        "tokens": list(JHS_SPEC.tokens),
    }
    return fingerprint


def validate_generated(generated_path,
                       expected_values: Sequence[str] = (),
                       forbidden_values: Sequence[str] = (),
                       source_path=None) -> Dict[str, Any]:
    """Full acceptance: structural comparison AND content hygiene."""
    source = source_path or golden_master_source_path()
    structure = compare_structure(source, generated_path)
    content = compare_content(generated_path, expected_values, forbidden_values)
    return {
        "pass": bool(structure["pass"] and content["pass"]),
        "structure": structure,
        "content": content,
    }
