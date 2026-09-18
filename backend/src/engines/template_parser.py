"""
SchemeKnit DOCX Template Parser

Analyzes an uploaded .docx file to extract lesson plan template structure.
Identifies sections, fields, tables, and layout properties.
"""

import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..models import Template, TemplateSection, TemplateField, TemplateFieldType, TemplateLayout, TemplateFamily, EducationalLevel


KNOWN_SECTIONS = {
    "header": ["school name", "subject", "class", "term", "week", "date", "teacher"],
    "strand_info": ["strand", "sub-strand", "sub strand"],
    "content_standard": ["content standard", "curriculum standard", "standard code"],
    "indicators": ["indicator", "curriculum indicator", "learning indicator"],
    "objectives": ["objective", "learning objective", "learning outcomes", "by the end"],
    "resources": ["resource", "teaching learning resource", "materials", "instructional material"],
    "previous_knowledge": ["previous knowledge", "prior knowledge", "prerequisite"],
    "introduction": ["introduction", "introductory activity", "starter", "warm-up", "warm up", "set induction"],
    "main_activities": ["main activity", "teacher activity", "teacher's activity", "presentation", "development"],
    "learner_activities": ["learner activity", "student activity", "pupil activity", "practice", "class activity"],
    "teacher_activities": ["teacher activity", "facilitation", "monitoring"],
    "assessment": ["assessment", "evaluation", "assessment criteria", "closer"],
    "conclusion": ["conclusion", "summary", "wrap-up", "closure", "reflection"],
    "homework": ["homework", "assignment", "take-home"],
    "differentiation": ["differentiation", "extension", "remediation", "enrichment"],
    "core_competencies": ["core competency", "competency", "21st century"],
    "keywords": ["keyword", "key word", "vocabulary", "key concepts"],
    "essential_questions": ["essential question", "big question", "driving question"],
}


def parse_docx_template(file_path: str) -> Dict:
    """
    Parse a DOCX file and extract template structure.
    Returns a dict with sections, fields, layout, and raw content.
    """
    doc = Document(file_path)
    
    sections = []
    tables_found = []
    headings_found = []
    paragraphs_sample = []
    
    # Extract layout info
    layout = _extract_layout(doc)
    
    # Analyze paragraphs for section headings
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        
        # Check if it's a heading
        if para.style.name.startswith("Heading") or _looks_like_heading(para):
            headings_found.append({
                "level": _heading_level(para),
                "text": text,
                "index": i,
            })
        
        # Sample paragraphs for content analysis
        if len(paragraphs_sample) < 100:
            paragraphs_sample.append(text)
    
    # Analyze tables
    for t_idx, table in enumerate(doc.tables):
        table_data = _parse_table(table)
        tables_found.append({
            "index": t_idx,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "data": table_data,
        })
    
    # Match headings to known sections
    sections = _match_sections(headings_found, tables_found, paragraphs_sample)
    
    # If no sections matched, create a generic structure
    if not sections:
        sections = _infer_sections_from_content(paragraphs_sample, tables_found)
    
    # Detect educational level and family
    family, level = _detect_level(paragraphs_sample)
    
    return {
        "sections": sections,
        "tables": tables_found,
        "headings": headings_found,
        "layout": layout,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "detected_family": family,
        "detected_level": level,
        "sample_content": paragraphs_sample[:20],
    }


def _extract_layout(doc: Document) -> Dict:
    """Extract layout properties from document."""
    layout = {
        "page_size": "A4",
        "orientation": "portrait",
        "font_family": "Arial",
        "font_size": 11,
    }
    
    try:
        section = doc.sections[0]
        width_inches = section.page_width / 914400  # EMU to inches
        height_inches = section.page_height / 914400
        
        if width_inches > height_inches:
            layout["orientation"] = "landscape"
        
        if abs(width_inches - 8.5) < 0.5 and abs(height_inches - 11) < 0.5:
            layout["page_size"] = "Letter"
    except Exception:
        pass
    
    # Try to get font from first paragraph
    try:
        for para in doc.paragraphs:
            if para.runs:
                run = para.runs[0]
                if run.font.name:
                    layout["font_family"] = run.font.name
                if run.font.size:
                    layout["font_size"] = int(run.font.size.pt)
                break
    except Exception:
        pass
    
    return layout


def _looks_like_heading(para) -> bool:
    """Check if a paragraph looks like a heading based on formatting."""
    if not para.runs:
        return False
    run = para.runs[0]
    if run.bold and run.font.size and run.font.size.pt >= 12:
        return True
    if para.style.name == "Normal" and run.bold:
        text = para.text.strip()
        if text.endswith(":") or re.match(r"^[A-Z][\w\s]+:$", text):
            return True
    return False


def _heading_level(para) -> int:
    """Get heading level from style name."""
    match = re.match(r"Heading (\d)", para.style.name)
    if match:
        return int(match.group(1))
    if para.runs and para.runs[0].bold:
        return 2
    return 3


def _parse_table(table) -> List[List[str]]:
    """Parse a docx table into a 2D list of strings."""
    data = []
    for row in table.rows:
        row_data = []
        for cell in row.cells:
            row_data.append(cell.text.strip())
        data.append(row_data)
    return data


def _match_sections(headings: List[Dict], tables: List[Dict], paragraphs: List[str]) -> List[Dict]:
    """Match detected headings to known template sections."""
    sections = []
    matched_indices = set()
    
    for h in headings:
        text_lower = h["text"].lower().rstrip(":")
        for section_name, keywords in KNOWN_SECTIONS.items():
            if any(kw in text_lower for kw in keywords):
                if section_name not in [s["name"] for s in sections]:
                    sections.append({
                        "name": section_name,
                        "label": h["text"].rstrip(":"),
                        "order": len(sections),
                        "found_at": h["index"],
                        "has_table": any(
                            abs(t["index"] - h["index"]) <= 3
                            for t in tables
                        ),
                    })
                    matched_indices.add(h["index"])
                break
    
    return sections


def _infer_sections_from_content(paragraphs: List[str], tables: List[Dict]) -> List[Dict]:
    """Infer sections from content when headings don't match."""
    sections = []
    
    # Always add header
    sections.append({
        "name": "header",
        "label": "Lesson Header",
        "order": 0,
        "found_at": 0,
        "has_table": len(tables) > 0,
    })
    
    # Check for common patterns in paragraphs
    text_blob = " ".join(paragraphs).lower()
    
    section_checks = [
        ("introduction", ["introduction", "intro", "starter", "warm"]),
        ("main_activities", ["main activity", "presentation", "development", "teacher activity"]),
        ("learner_activities", ["learner", "student", "pupil", "class activity", "practice"]),
        ("assessment", ["assessment", "evaluation", "quiz", "test"]),
        ("conclusion", ["conclusion", "summary", "wrap", "closure"]),
        ("objectives", ["objective", "learning outcome", "by the end"]),
        ("resources", ["resource", "material", "instructional"]),
    ]
    
    order = 1
    for name, keywords in section_checks:
        if any(kw in text_blob for kw in keywords):
            sections.append({
                "name": name,
                "label": name.replace("_", " ").title(),
                "order": order,
                "found_at": order,
                "has_table": False,
            })
            order += 1
    
    return sections


def _detect_level(paragraphs: List[str]) -> Tuple[str, str]:
    """Detect educational level from content."""
    text_blob = " ".join(paragraphs).lower()
    
    if any(kw in text_blob for kw in ["nursery", "kg 1", "kg 2", "kindergarten"]):
        return "early_childhood", "Nursery/KG"
    if any(kw in text_blob for kw in ["basic 1", "basic 2", "basic 3", "basic 4", "basic 5", "basic 6", "primary"]):
        return "primary", "Primary"
    if any(kw in text_blob for kw in ["basic 7", "basic 8", "basic 9", "jhs", "junior high"]):
        return "jhs", "Junior High School"
    if any(kw in text_blob for kw in ["shs", "senior high", "basic 10", "basic 11", "basic 12"]):
        return "shs", "Senior High School"
    
    return "jhs", "Junior High School"  # default


def parsed_to_template_sections(parsed: Dict) -> List[Dict]:
    """Convert parsed result to template sections format."""
    sections = []
    for i, s in enumerate(parsed.get("sections", [])):
        sections.append({
            "name": s["name"],
            "label": s.get("label", s["name"].replace("_", " ").title()),
            "visible": True,
            "required": s["name"] in ("header", "introduction", "main_activities", "assessment"),
            "order": i,
            "fields": _default_fields_for_section(s["name"]),
        })
    return sections


def _default_fields_for_section(section_name: str) -> List[Dict]:
    """Get default fields for a section based on its name."""
    field_map = {
        "header": [
            {"name": "school_name", "label": "School Name", "field_type": "text", "required": False},
            {"name": "subject", "label": "Subject", "field_type": "text", "required": True},
            {"name": "class_level", "label": "Class", "field_type": "text", "required": True},
            {"name": "term", "label": "Term", "field_type": "text", "required": True},
            {"name": "week_number", "label": "Week", "field_type": "number", "required": True},
            {"name": "lesson_date", "label": "Date", "field_type": "date", "required": False},
            {"name": "teacher_name", "label": "Teacher", "field_type": "text", "required": False},
        ],
        "strand_info": [
            {"name": "strand", "label": "Strand", "field_type": "text", "required": True},
            {"name": "sub_strand", "label": "Sub-strand", "field_type": "text", "required": False},
        ],
        "content_standard": [
            {"name": "content_standard", "label": "Content Standard", "field_type": "textarea", "required": True},
            {"name": "content_standard_code", "label": "Standard Code", "field_type": "text", "required": True},
        ],
        "indicators": [
            {"name": "indicators", "label": "Indicators", "field_type": "textarea", "required": True},
        ],
        "objectives": [
            {"name": "learning_objectives", "label": "Learning Objectives", "field_type": "textarea", "required": True},
        ],
        "resources": [
            {"name": "teaching_learning_resources", "label": "Resources", "field_type": "textarea", "required": False},
        ],
        "previous_knowledge": [
            {"name": "previous_knowledge", "label": "Previous Knowledge", "field_type": "textarea", "required": False},
        ],
        "introduction": [
            {"name": "introduction", "label": "Introduction", "field_type": "textarea", "required": True},
        ],
        "main_activities": [
            {"name": "main_activities", "label": "Main Activities", "field_type": "textarea", "required": True},
        ],
        "learner_activities": [
            {"name": "learner_activities", "label": "Learner Activities", "field_type": "textarea", "required": True},
        ],
        "teacher_activities": [
            {"name": "teacher_activities", "label": "Teacher Activities", "field_type": "textarea", "required": False},
        ],
        "assessment": [
            {"name": "assessment", "label": "Assessment", "field_type": "textarea", "required": True},
        ],
        "conclusion": [
            {"name": "conclusion", "label": "Conclusion", "field_type": "textarea", "required": False},
        ],
        "homework": [
            {"name": "homework", "label": "Homework", "field_type": "textarea", "required": False},
        ],
        "differentiation": [
            {"name": "differentiation", "label": "Differentiation", "field_type": "textarea", "required": False},
        ],
        "core_competencies": [
            {"name": "core_competencies", "label": "Core Competencies", "field_type": "textarea", "required": False},
        ],
        "keywords": [
            {"name": "keywords", "label": "Keywords", "field_type": "text", "required": False},
        ],
        "essential_questions": [
            {"name": "essential_questions", "label": "Essential Questions", "field_type": "textarea", "required": False},
        ],
    }
    return field_map.get(section_name, [
        {"name": section_name, "label": section_name.replace("_", " ").title(), "field_type": "textarea", "required": False},
    ])
