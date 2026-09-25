"""
Approved WAPEF Basic 1-3 Plan — weekly CLASS-TEACHER plan (in-place renderer).

SOURCE OF TRUTH: ``WAPEF BASIC 1 PLAN.docx``, the weekly class plan supplied by
WAPEF (bundled verbatim at ``backend/tests/fixtures/wapef/WAPEF BASIC 1
PLAN.docx``). It is ONE weekly document — school + "LESSON PLAN" + "WEEK n"
followed by one table per SUBJECT SECTION:

    Week Ending / Class / Subject / Reference / Content Standard /
    Learning Indicator(s) / Performance Indicator / [Strand] / [Sub strand] /
    Teaching-Learning Resources / Throughlines / God's Story / Deep Hope /
    Storyline / Core Competencies / Key words
    DAYS | PHASE 1: STARTER | PHASE 2: MAIN | PHASE 3: REFLECTION
    one row per teaching day

This module NEVER hand-writes that topology. The bundled, tokenized asset
(``assets/wapef_basic13_weekly_template.docx``, derived from the supplied
document by ``tools/build_basic13_asset.py``) IS the subject-section block: the
renderer fills its tokens in place and clones it once per subject, so a
generated weekly plan is topologically identical to the supplied plan by
construction — the same approach the Approved WAPEF Plan and the official
GES/NaCCA forms use.

Contract
--------
- Administrative and curriculum fields are DETERMINISTIC-ONLY: they resolve from
  the stored lesson rows (source scheme data) and never from AI.
- The four WAPEF fields (Deep Hope, Storyline, Through lines, God's Story) are
  TEACHER-SELECTED structured values; blank stays blank — never invented.
- DAYS is a first-class structured column: each teaching day (or the teacher's
  grouped "MONDAY & THURSDAY" entry) has its own STARTER / MAIN / REFLECTION.
- Subject sections are isolated: nothing (metadata, indicators, resources,
  WAPEF selections) leaks from one subject into another.
- No subject list, day pattern or weekday schedule is hard-coded: the subjects
  are the teacher's, the days are the teacher's, and an optional metadata row
  (Strand / Sub strand) is dropped when the source has nothing for it.
"""

from copy import deepcopy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from ..models import ClassLevel
from .official_ges_template import (
    _fill_paragraph,
    _iter_paragraphs,
    _sectpr,
    PLACEHOLDER_RE,
)
from .wapef_template import format_wapef_date
from .wapef_fields import format_through_lines, normalize_through_lines

TEMPLATE_VERSION = "1.0"

#: Stable internal template id. The subject-teacher template
#: (``tpl-wapef-approved-plan``) is untouched: this is an ADDITIONAL template.
TEMPLATE_ID = "tpl-wapef-basic13-weekly-plan"

TEMPLATE_NAME = "Approved WAPEF Basic 1\u20133 Plan"

#: The supplied weekly class plan this asset is derived from (repo-relative).
SOURCE_DOCUMENT = "backend/tests/fixtures/wapef/WAPEF BASIC 1 PLAN.docx"

ASSET_FILENAME = "wapef_basic13_weekly_template.docx"

ASSETS_DIR = Path(__file__).resolve().parent / "assets"

#: The class-teacher band. Basic 4-JHS stays on the subject-teacher Approved
#: WAPEF Plan; this template must never be routed there (or the reverse).
BASIC13_CLASS_LEVELS = (
    ClassLevel.BASIC_1,
    ClassLevel.BASIC_2,
    ClassLevel.BASIC_3,
)

#: Canonical teaching days, in week order (Python ``date.weekday()`` numbering:
#: Monday = 0). The names are the ones the supplied plan prints.
WEEK_DAYS = ("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY")

#: ``weekday()`` -> canonical day name (Saturday/Sunday are never teaching days
#: and map to an empty string rather than being invented as one).
WEEKDAY_TO_NAME = {
    0: "MONDAY",
    1: "TUESDAY",
    2: "WEDNESDAY",
    3: "THURSDAY",
    4: "FRIDAY",
}

#: Teacher-facing day options (value = weekday int, label = printed name).
DAY_OPTIONS = [
    {"value": index, "label": name.title()} for index, name in enumerate(WEEK_DAYS)
]

#: The three delivery phases, exactly as the supplied plan prints them.
WAPEF_BASIC13_PHASES = (
    "PHASE 1: STARTER",
    "PHASE 2: MAIN",
    "PHASE 3: REFLECTION",
)

#: Metadata rows the supplied plan only prints when the source has a value for
#: them (English Language prints no Strand/Sub strand row at all). A blank value
#: drops the row instead of printing an empty label.
OPTIONAL_METADATA_ROWS = ("Strand", "Sub strand")



# ── Class routing (Basic 1-3 boundary) ────────────────────────────────────────

def _class_level_value(value) -> str:
    if isinstance(value, ClassLevel):
        return value.value
    return str(getattr(value, "value", value) or "")


def is_basic13_class(class_level) -> bool:
    """True for Basic 1, Basic 2 and Basic 3 (and nothing else)."""
    return _class_level_value(class_level) in [c.value for c in BASIC13_CLASS_LEVELS]


def is_basic13_template(template) -> bool:
    """True when ``template`` (a Template, a template id, or None) is this plan."""
    if template is None:
        return False
    ident = template if isinstance(template, str) else getattr(template, "id", None)
    return ident == TEMPLATE_ID


def wapef_template_for_class(class_level) -> str:
    """The WAPEF template id for one class level.

    Basic 1-3 uses the class-teacher weekly plan; Basic 4-JHS keeps the
    existing subject-teacher Approved WAPEF Plan. This single rule is the only
    place the boundary is decided, so it cannot drift between the API, the UI
    and the tests: Basic 3 must never resolve to the subject-teacher form, and
    Basic 4 must never resolve to the class-teacher form.
    """
    from .wapef_template import TEMPLATE_ID as WAPEF_APPROVED_ID

    return TEMPLATE_ID if is_basic13_class(class_level) else WAPEF_APPROVED_ID


def wapef_routing(class_level) -> Dict[str, Any]:
    """Routing description for the UI/API (which WAPEF plan applies here)."""
    from .wapef_template import TEMPLATE_ID as WAPEF_APPROVED_ID

    template_id = wapef_template_for_class(class_level)
    return {
        "class_level": _class_level_value(class_level),
        "template_id": template_id,
        "template_name": (TEMPLATE_NAME if template_id == TEMPLATE_ID
                          else "Approved WAPEF Plan"),
        "planning_model": ("class_teacher" if template_id == TEMPLATE_ID
                           else "subject_teacher"),
        "basic13_template_id": TEMPLATE_ID,
        "subject_teacher_template_id": WAPEF_APPROVED_ID,
        "class_levels": [c.value for c in BASIC13_CLASS_LEVELS],
    }


# ── Day model ─────────────────────────────────────────────────────────────────

def normalize_teaching_days(values) -> List[str]:
    """Canonical teaching-day names, week order, deduplicated.

    Accepts ``"Monday"``, ``"MONDAY"``, ``"Mon"``, ``0`` (Python
    ``date.weekday()``) or a mix. Anything that is not a Monday-Friday teaching
    day is dropped — a weekday is never invented and a Saturday/Sunday entry is
    not silently relabelled.
    """
    if values is None:
        return []
    if isinstance(values, (str, int)) and not isinstance(values, bool):
        values = [values]
    names = set()
    for value in values:
        if value is None or isinstance(value, bool):
            continue
        if isinstance(value, int) or (isinstance(value, str) and value.strip().isdigit()):
            index = int(value)
            if 0 <= index < len(WEEK_DAYS):
                names.add(WEEK_DAYS[index])
            continue
        text = " ".join(str(value).split()).strip().upper()
        for name in WEEK_DAYS:
            if text == name or (text and name.startswith(text) and len(text) >= 3):
                names.add(name)
                break
    return [name for name in WEEK_DAYS if name in names]


def normalize_day_groups(groups) -> List[List[str]]:
    """Canonical teaching-day GROUPS, week order, each deduplicated.

    One inner list is one shared teaching entry — the supplied plan's
    "MONDAY & THURSDAY" RME row is ``["MONDAY", "THURSDAY"]``. Groups are
    ordered by their earliest day and a day never repeats across groups.
    """
    cleaned: List[List[str]] = []
    seen: set = set()
    for group in groups or []:
        days = [d for d in normalize_teaching_days(group) if d not in seen]
        if not days:
            continue
        seen.update(days)
        cleaned.append(days)
    cleaned.sort(key=lambda g: WEEK_DAYS.index(g[0]))
    return cleaned


def flatten_day_groups(groups) -> List[str]:
    """Every weekday declared by the groups, week order (no duplicates)."""
    return normalize_teaching_days([d for group in groups or [] for d in group])


def format_day_label(days) -> str:
    """The DAYS cell text: "MONDAY", "MONDAY & THURSDAY", "MONDAY, TUESDAY & FRIDAY"."""
    names = normalize_teaching_days(days)
    if not names:
        return ""
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + f" & {names[-1]}"


def day_name_for_weekday(weekday: int) -> str:
    """Canonical day name for a ``date.weekday()`` value ("" outside Mon-Fri)."""
    return WEEKDAY_TO_NAME.get(int(weekday), "")


def weekday_index(day: str) -> int:
    """``date.weekday()`` index for a canonical day name (-1 when not a day)."""
    names = normalize_teaching_days([day])
    return WEEK_DAYS.index(names[0]) if names else -1


def template_path() -> Path:
    return ASSETS_DIR / ASSET_FILENAME



# ── Deterministic value resolution ────────────────────────────────────────────

def _get(obj, name: str, default=None):
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def _text(value) -> str:
    """Plain string for a stored value, unwrapping enums."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(getattr(value, "value", value))


def _text_items(value) -> List[str]:
    """Stored JSON column (list | JSON string | text | None) -> non-empty strings."""
    import json

    if value is None or value == "":
        return []
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return []
        if text[0] in "[{":
            try:
                value = json.loads(text)
            except Exception:
                return [" ".join(line.split()) for line in text.splitlines()
                        if line.strip()]
        else:
            return [" ".join(line.split()) for line in text.splitlines()
                    if line.strip()]
    if isinstance(value, dict):
        value = [value]
    if not isinstance(value, (list, tuple, set)):
        return [str(value)]
    out: List[str] = []
    for item in value:
        if item is None:
            continue
        if isinstance(item, str):
            out.append(" ".join(item.split()))
        elif isinstance(item, dict):
            desc = (item.get("description") or item.get("title")
                    or item.get("text") or "")
            out.append(" ".join(str(desc).split()))
        else:
            out.append(" ".join(str(getattr(item, "description", item)).split()))
    return [x for x in out if x]


def subject_metadata_values(section) -> Dict[str, str]:
    """``token -> text`` for one subject section's metadata block.

    Every value comes from the stored lesson rows — the source scheme's own
    curriculum data — plus the teacher's WAPEF selections. A value nothing
    supplies stays blank, and its optional row is dropped on render.
    """
    meta = _metadata(section)
    section_subject = _text(_get(meta, "subject") or _get(section, "subject"))

    standards = _unique(_text_items(_get(meta, "content_standards")))
    standard_codes = _unique(_text_items(_get(meta, "content_standard_codes")))
    indicator_texts = _unique(_text_items(_get(meta, "indicators")))
    indicator_codes = _unique(_text_items(_get(meta, "indicator_codes")))
    performance = _unique(_text_items(_get(meta, "performance_indicators")))

    standard_lines: List[str] = []
    for index, text in enumerate(standards):
        code = standard_codes[index] if index < len(standard_codes) else ""
        standard_lines.append(f"{code} {text}".strip() if code else text)

    return {
        "WEEK_ENDING": format_wapef_date(_get(meta, "week_ending")
                                         or _get(section, "week_ending")),
        "CLASS": _text(_get(meta, "class_level")
                       or _get(section, "class_level")),
        "SUBJECT": section_subject,
        "REFERENCE": _text(_get(meta, "reference")),
        "CONTENT_STANDARD": "\n".join(_unique(standard_lines)),
        "LEARNING_INDICATOR": "\n".join(
            _numbered_codes(indicator_codes, indicator_texts)),
        "PERFORMANCE_INDICATOR": "\n".join(performance),
        "STRAND": _text(_get(meta, "strand")),
        "SUB_STRAND": _text(_get(meta, "sub_strand")),
        "RESOURCES": ", ".join(_unique(
            _text_items(_get(meta, "teaching_learning_resources")))),
        "THROUGH_LINES": format_through_lines(normalize_through_lines(
            _text_items(_get(meta, "wapef_through_lines")))),
        "GODS_STORY": _text(_get(meta, "wapef_gods_story")),
        "DEEP_HOPE": _text(_get(meta, "wapef_deep_hope")),
        "STORYLINE": _text(_get(meta, "wapef_storyline")),
        "CORE_COMPETENCIES": ", ".join(_unique(
            _text_items(_get(meta, "core_competencies")))),
        "KEY_WORDS": ", ".join(_unique(_text_items(_get(meta, "keywords")))),
    }


def day_plan_values(day_plan) -> Dict[str, str]:
    """``token suffix -> text`` for one teaching-day row."""
    main_lines = [_text(_get(a, "description"))
                  for a in (_get(day_plan, "main_activities") or [])]
    return {
        "DAY_LABEL": (format_day_label(_get(day_plan, "days"))
                      or _text(_get(day_plan, "day_label"))),
        "PHASE1": _text(_get(day_plan, "starter")),
        "PHASE2": "\n".join(_unique(main_lines)),
        "PHASE3": _text(_get(day_plan, "reflection")),
    }


def subject_plan_values(section) -> Dict[str, str]:
    """Full ``token -> text`` map for one subject section (metadata + days)."""
    values = subject_metadata_values(section)
    for index, day_plan in enumerate(_get(section, "day_plans") or [], start=1):
        for token, text in day_plan_values(day_plan).items():
            values[f"{token}_{index}"] = text
    return values

# ── Rendering ─────────────────────────────────────────────────────────────────

def _unique_cells(row) -> List[Any]:
    """A row's cells with horizontal-merge duplicates collapsed (grid order).

    The source's metadata rows merge their label and value cells, so
    ``row.cells[1]`` can be the label's own tc; addressing the row's *distinct*
    cells is what follows the supplied document's layout.
    """
    out, seen = [], set()
    for cell in row.cells:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        out.append(cell)
    return out


def _row_label(row) -> str:
    cells = _unique_cells(row)
    return " ".join((cells[0].text or "").split()) if cells else ""


def _row_is_empty(row) -> bool:
    return not _row_label(row) and all(
        not (cell.text or "").strip() for cell in _unique_cells(row))


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in _unique_cells(row):
            yield from _iter_paragraphs(cell)


def _drop_structural_rows(table) -> int:
    """Remove the rows the source itself only prints when it has data.

    * an OPTIONAL metadata row (Strand / Sub strand) whose value is blank is
      dropped — the supplied plan's English section prints no Strand row at all,
      while its Maths/Science/RME/History sections do; and
    * teaching-day rows beyond the days the teacher actually teaches are
      dropped, so a subject taught on two days has two day rows and never an
      empty MONDAY-FRIDAY skeleton.
    """
    rows = list(table.rows)
    header_index = next(
        (i for i, row in enumerate(rows)
         if _row_label(row).upper().startswith("DAYS")), None)
    removed = 0
    for index, row in enumerate(rows):
        if header_index is None:
            break
        if index <= header_index:
            continue
        if _row_is_empty(row):
            row._tr.getparent().remove(row._tr)
            removed += 1
    for row in list(table.rows):
        label = _row_label(row).rstrip(":").strip()
        if label in OPTIONAL_METADATA_ROWS:
            cells = _unique_cells(row)
            if len(cells) >= 2 and not (cells[1].text or "").strip():
                row._tr.getparent().remove(row._tr)
                removed += 1
    return removed


def render_section_into(table, section) -> int:
    """Fill one subject section (metadata + day rows) of the weekly plan.

    Tokens are substituted *inside* each paragraph's text, so the supplied
    plan's own inline labels ("Throughlines: [THROUGH_LINES]", "God's Story:
    [GODS_STORY]", "Deep Hope : [DEEP_HOPE]", "Core Competencies: …") keep
    printing exactly as the source prints them, with only the value changing.
    Returns the number of tokens filled. Safe to call on a cloned block.
    """
    values = subject_plan_values(section)
    filled = 0
    for paragraph in _iter_table_paragraphs(table):
        text = paragraph.text or ""
        tokens = PLACEHOLDER_RE.findall(text)
        if not tokens:
            continue
        resolved = PLACEHOLDER_RE.sub(lambda m: values.get(m.group(1), ""), text)
        _fill_paragraph(paragraph, resolved)
        filled += len(tokens)
    _drop_structural_rows(table)
    return filled


def load_document():
    """Open the bundled weekly class-plan form as a python-docx Document."""
    from docx import Document

    path = template_path()
    if not path.is_file():
        raise FileNotFoundError(f"Basic 1-3 weekly template missing: {path}")
    return Document(str(path))


def render_header(document, weekly_plan, context: Optional[Dict[str, Any]] = None):
    """Fill the school line and the WEEK line above the subject sections.

    The source prints exactly: [0] school name, [1] LESSON PLAN, [2] WEEK n.
    Identity is server-derived; a blank school renders blank (never a
    placeholder).
    """
    ctx = context or {}
    school = _text(ctx.get("school_name") or _get(weekly_plan, "school_name"))
    week = _get(weekly_plan, "week_number")
    paragraphs = list(document.paragraphs)
    if len(paragraphs) >= 1:
        _fill_paragraph(paragraphs[0], school)
    if len(paragraphs) >= 3:
        _fill_paragraph(paragraphs[2],
                        f"WEEK {_text(week)}" if week not in (None, "", 0) else "WEEK")


def render_weekly_document(weekly_plan, context: Optional[Dict[str, Any]] = None):
    """One Document: ONE weekly class plan holding every subject's section.

    The subject-section block is cloned per subject (a token can only be filled
    once), and each section after the first starts on its own page so the
    subjects are clearly separated — never merged into one metadata block and
    never emitted as independent lesson documents.
    """
    from docx.enum.text import WD_BREAK
    from docx.table import Table

    subjects = list(_get(weekly_plan, "subjects") or [])
    document = load_document()
    render_header(document, weekly_plan, context)
    if not subjects:
        return document

    body = document.element.body
    block = deepcopy(document.tables[0]._tbl)
    render_section_into(document.tables[0], subjects[0])

    for section in subjects[1:]:
        page_break = document.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)
        cloned = deepcopy(block)
        _sectpr(body).addprevious(cloned)
        render_section_into(Table(cloned, document), section)

    return document


def validate_rendered_document(document) -> List[str]:
    """Any placeholder left unfilled (empty list means a clean weekly plan)."""
    leftovers: List[str] = []
    for paragraph in _iter_paragraphs(document):
        leftovers.extend(PLACEHOLDER_RE.findall(paragraph.text or ""))
    for table in document.tables:
        for row in table.rows:
            for cell in _unique_cells(row):
                for paragraph in _iter_paragraphs(cell):
                    leftovers.extend(PLACEHOLDER_RE.findall(paragraph.text or ""))
    return leftovers


# ── Template registry sections ────────────────────────────────────────────────

def template_sections() -> List[Any]:
    """Sections/fields for the Basic 1-3 weekly plan, derived from the asset's
    own token contract.

    The four WAPEF fields stay SELECT fields bound to the approved option lists
    (``wapef_fields.py``) and the teaching days are a teacher multi-select, so
    the review UI gives the class teacher exactly the controls this plan needs —
    no second copy of the WAPEF field definitions.
    """
    from ..models import TemplateField, TemplateFieldType, TemplateSection
    from .wapef_fields import (
        WAPEF_DEEP_HOPES, WAPEF_GODS_STORY, WAPEF_STORYLINES,
        WAPEF_THROUGH_LINES,
    )

    def field(name: str, label: str, *, options=None,
              field_type=TemplateFieldType.TEXT, order: int = 0,
              source=None) -> Any:
        from ..models import ContentSource
        return TemplateField(
            name=name, label=label, field_type=field_type, order=order,
            options=list(options or []),
            source=source or ContentSource.DETERMINISTIC,
        )

    metadata = TemplateSection(
        name="weekly_plan_metadata", label="Weekly Plan", order=0,
        description="Week, class and subject identity of this section.",
        fields=[
            field("week_number", "Week", field_type=TemplateFieldType.NUMBER, order=0),
            field("week_ending", "Week Ending", field_type=TemplateFieldType.DATE, order=1),
            field("class_level", "Class", order=2),
            field("subject", "Subject", order=3),
            field("reference", "Reference", order=4),
        ],
    )

    curriculum = TemplateSection(
        name="weekly_plan_curriculum", label="Curriculum Alignment", order=1,
        description="Authoritative curriculum data taken from the scheme.",
        fields=[
            field("content_standard", "Content Standard",
                  field_type=TemplateFieldType.TEXTAREA, order=0),
            field("indicators", "Learning Indicator(s)",
                  field_type=TemplateFieldType.TEXTAREA, order=1),
            field("performance_indicators", "Performance Indicator",
                  field_type=TemplateFieldType.TEXTAREA, order=2),
            field("strand", "Strand", order=3),
            field("sub_strand", "Sub strand", order=4),
            field("teaching_learning_resources", "Teaching/ Learning Resources",
                  field_type=TemplateFieldType.TEXTAREA, order=5),
            field("core_competencies", "Core Competencies",
                  field_type=TemplateFieldType.TEXTAREA, order=6),
            field("keywords", "Key words", order=7),
        ],
    )

    special = TemplateSection(
        name="weekly_plan_special", label="WAPEF Special Fields", order=2,
        description="Teacher-selected WAPEF values (never AI-chosen).",
        fields=[
            field("wapef_through_lines", "Through lines",
                  field_type=TemplateFieldType.SELECT,
                  options=WAPEF_THROUGH_LINES, order=0),
            field("wapef_gods_story", "God's Story",
                  field_type=TemplateFieldType.SELECT,
                  options=WAPEF_GODS_STORY, order=1),
            field("wapef_deep_hope", "Deep Hope",
                  field_type=TemplateFieldType.SELECT,
                  options=WAPEF_DEEP_HOPES, order=2),
            field("wapef_storyline", "Storyline",
                  field_type=TemplateFieldType.SELECT,
                  options=WAPEF_STORYLINES, order=3),
        ],
    )

    days = TemplateSection(
        name="weekly_plan_days", label="Teaching Days", order=3,
        description=("The days this subject is taught this week. One entry per "
                     "shared teaching day group (a grouped entry such as "
                     "'Monday & Thursday' is supported)."),
        fields=[
            field("teaching_day_groups", "Teaching Days",
                  field_type=TemplateFieldType.SELECT,
                  options=[d["label"] for d in DAY_OPTIONS], order=0),
        ],
    )

    for index, section in enumerate((metadata, curriculum, special, days)):
        section.order = index
        section.fields = [
            f.model_copy(update={"order": i}) for i, f in enumerate(section.fields)
        ]
    return [metadata, curriculum, special, days]


def _unique(items: Iterable[str]) -> List[str]:
    """Whitespace-insensitive de-duplication preserving first-seen order."""
    out: List[str] = []
    seen = set()
    for item in items:
        text = " ".join((item or "").split())
        if not text:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(text)
    return out


def _code_and_text(code, text) -> str:
    code = _text(code).strip()
    text = _text(text).strip()
    if code and text:
        return f"{code} {text}"
    return code or text


def _numbered_codes(codes: List[str], texts: List[str]) -> List[str]:
    """Pair indicator codes with their descriptions (WAPEF convention)."""
    out: List[str] = []
    for index, text in enumerate(texts):
        code = codes[index] if index < len(codes) else ""
        if code and text.lower().startswith(code.lower()):
            out.append(text)
        else:
            out.append(_code_and_text(code, text))
    return _unique(out)


def _metadata(plan_like) -> Any:
    """The SubjectMetadata of a SubjectPlan (or the object itself)."""
    metadata = _get(plan_like, "metadata")
    return metadata if metadata is not None else plan_like
