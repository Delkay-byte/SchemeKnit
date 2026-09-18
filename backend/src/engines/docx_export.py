"""
SchemeKnit DOCX Export Engine

Production-quality DOCX generation for lesson plans.
Uses python-docx for clean, professional document creation.
"""

from pathlib import Path
from typing import List, Optional
from datetime import date
import re

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from ..models import (
    LessonPlan, Template, TemplateType, Week, Subject, ClassLevel, EducationalLevel,
    CLASS_LEVEL_TO_EDUCATIONAL_LEVEL,
)
from .template_engine import (
    get_visible_sections, get_template_by_type, get_default_template_for_class_level,
)
from .official_ges_levels import (
    JHS_SPEC, LEVEL_JHS, LEVEL_KG, LEVEL_PRIMARY, LEVEL_SHS, GESLevelSpec,
    spec_for_level, spec_for_template_id,
)
from .official_ges_template import (
    is_ges_form_template,
    render_document as render_official_ges_document,
)


class DOCXExportEngine:
    """Production DOCX generation for lesson plans."""

    def export_single(
        self,
        lesson_plan: LessonPlan,
        template: Optional[Template] = None,
        output_path: Optional[Path] = None,
    ) -> Path:
        """Generate a single lesson plan DOCX."""
        if template is None:
            template = default_template_for_lessons([lesson_plan], TemplateType.GES_STYLE)

        if is_ges_form_template(template):
            if output_path is None:
                output_path = Path(f"lesson_plan_{lesson_plan.id}.docx")
            return self.export_official_ges([lesson_plan], output_path,
                                            spec=spec_for_template_id(template.id))

        doc = Document()
        self._setup_styles(doc)
        self._render_lesson_plan(doc, lesson_plan, template)

        if output_path is None:
            output_path = Path(f"lesson_plan_{lesson_plan.id}.docx")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def export_batch(
        self,
        lesson_plans: List[LessonPlan],
        template: Optional[Template] = None,
        output_dir: Path = Path("exports"),
        naming_pattern: str = "{subject}_{class}_W{week:02d}_L{lesson:02d}_{date}",
    ) -> List[Path]:
        """Generate DOCX files for a batch of lesson plans."""
        if template is None:
            template = default_template_for_lessons(lesson_plans, TemplateType.GES_STYLE)

        output_dir.mkdir(parents=True, exist_ok=True)
        generated: List[Path] = []

        for lp in lesson_plans:
            filename = self._make_filename(lp, naming_pattern)
            filepath = output_dir / filename
            self.export_single(lp, template, filepath)
            generated.append(filepath)

        return generated

    def export_single_custom(
        self,
        lesson_plan: LessonPlan,
        structure: dict,
        output_path: Optional[Path] = None,
        context: dict = None,
    ) -> Path:
        """Generate a single lesson plan DOCX through a custom sample structure."""
        doc = Document()
        self._setup_styles(doc)
        self._apply_custom_layout(doc, structure.get("layout", {}))
        render_custom_template(doc, lesson_plan, structure, context)
        if output_path is None:
            output_path = Path(f"lesson_plan_{lesson_plan.id}.docx")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def export_combined_custom(
        self,
        lesson_plans: List[LessonPlan],
        structure: dict,
        output_path: Path = Path("exports/lesson_plans.docx"),
        context: dict = None,
    ) -> Path:
        """Generate one DOCX with all lesson plans rendered via a custom structure."""
        doc = Document()
        self._setup_styles(doc)
        self._apply_custom_layout(doc, structure.get("layout", {}))
        for index, lp in enumerate(lesson_plans):
            if index > 0:
                doc.add_page_break()
            render_custom_template(doc, lp, structure, context)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def export_batch_custom(
        self,
        lesson_plans: List[LessonPlan],
        structure: dict,
        output_dir: Path = Path("exports"),
        naming_pattern: str = "{subject}_{class}_W{week:02d}_L{lesson:02d}_{date}",
        context: dict = None,
    ) -> List[Path]:
        """Generate one DOCX per lesson plan via a custom structure (for ZIP)."""
        output_dir.mkdir(parents=True, exist_ok=True)
        generated: List[Path] = []
        for lp in lesson_plans:
            filename = self._make_filename(lp, naming_pattern)
            generated.append(self.export_single_custom(lp, structure, output_dir / filename, context))
        return generated

    def _apply_custom_layout(self, doc: Document, layout: dict):
        """Apply sampled page orientation to the document."""
        try:
            if (layout or {}).get("orientation") == "landscape":
                section = doc.sections[0]
                w, h = section.page_width, section.page_height
                section.page_width, section.page_height = h, w
        except Exception:
            pass

    def export_combined(
        self,
        lesson_plans: List[LessonPlan],
        template: Optional[Template] = None,
        output_path: Path = Path("exports/lesson_plans.docx"),
    ) -> Path:
        """Generate one DOCX containing all lesson plans, each on its own page."""
        if template is None:
            template = default_template_for_lessons(lesson_plans, TemplateType.GES_STYLE)

        if is_ges_form_template(template):
            return self.export_official_ges(lesson_plans, output_path,
                                            spec=spec_for_template_id(template.id))

        doc = Document()
        self._setup_styles(doc)

        for index, lp in enumerate(lesson_plans):
            if index > 0:
                doc.add_page_break()
            self._render_lesson_plan(doc, lp, template)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def export_official_ges(
        self,
        lesson_plans: List[LessonPlan],
        output_path: Optional[Path] = None,
        spec: Optional[GESLevelSpec] = None,
        context: dict = None,
    ) -> Path:
        """Render lessons onto one official GES/NaCCA placeholder form.

        The bundled form *is* the document (its bracketed placeholders are
        filled in place, one page per lesson) rather than a layout to imitate,
        so this path bypasses the section renderers entirely. ``spec`` picks the
        level's form; when the caller names none, the lessons' own level does.
        """
        if output_path is None:
            output_path = Path("exports/lesson_plans.docx")
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        lessons = list(lesson_plans)
        document = render_official_ges_document(spec or ges_spec_for_lessons(lessons),
                                               lessons, context)
        document.save(str(output_path))
        return output_path

    def _make_filename(self, lp: LessonPlan, pattern: str) -> str:
        subject = lp.subject.value.replace(" ", "_").upper() if isinstance(lp.subject, Subject) else str(lp.subject).upper()
        cls = lp.class_level.value.replace(" ", "") if isinstance(lp.class_level, ClassLevel) else str(lp.class_level)
        date_str = lp.lesson_date.strftime("%Y-%m-%d") if lp.lesson_date else "nodate"
        kwargs = {
            "subject": subject,
            "class": cls,
            "week": lp.week_number,
            "lesson": lp.lesson_sequence,
            "date": date_str,
        }
        safe = pattern.format(**kwargs)
        safe = re.sub(r'[<>:"/\\|?*]', '_', safe)
        if not safe.endswith('.docx'):
            safe += '.docx'
        return safe

    def _setup_styles(self, doc: Document):
        """Set up document styles."""
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)

        for size in range(1, 4):
            heading_style = doc.styles[f'Heading {size}']
            heading_style.font.color.rgb = RGBColor(0, 51, 102)

    def _render_lesson_plan(
        self,
        doc: Document,
        lp: LessonPlan,
        template: Template,
    ):
        """Render a lesson plan into a document."""
        visible = get_visible_sections(template)

        for section in visible:
            handler = SECTION_RENDERERS.get(section.name)
            if handler:
                handler(self, doc, lp)

        doc.add_page_break()

    def _render_header(self, doc: Document, lp: LessonPlan):
        title = doc.add_heading('LESSON PLAN', level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        info_table = doc.add_table(rows=6, cols=2)
        info_table.style = 'Table Grid'
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        subject_val = lp.subject.value if isinstance(lp.subject, Subject) else str(lp.subject)
        class_val = lp.class_level.value if isinstance(lp.class_level, ClassLevel) else str(lp.class_level)

        info_data = [
            ("School", lp.school_name or "________________"),
            ("Teacher", lp.teacher_name or "________________"),
            ("Subject", subject_val),
            ("Class", f"{class_val}  |  Size: {lp.class_size}"),
            ("Date", lp.lesson_date.strftime("%d %B %Y") if lp.lesson_date else ""),
            ("Duration", f"{lp.duration_minutes} minutes"),
        ]

        for i, (label, value) in enumerate(info_data):
            cell_label = info_table.cell(i, 0)
            cell_label.text = label
            for p in cell_label.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(10)
            cell_value = info_table.cell(i, 1)
            cell_value.text = str(value)
            for p in cell_value.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)

    def _render_strand_info(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Strand / Sub-Strand', level=2)
        strand = lp.strand or ""
        sub = lp.sub_strand or ""
        if strand or sub:
            p = doc.add_paragraph()
            if strand:
                run = p.add_run(f"Strand: {strand}")
                run.bold = True
            if sub:
                p.add_run(f"\nSub-Strand: {sub}")

    def _render_content_standard(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Content Standard', level=2)
        cs = lp.content_standard or ""
        if cs:
            p = doc.add_paragraph()
            if lp.content_standard_code:
                run = p.add_run(f"{lp.content_standard_code} ")
                run.bold = True
            p.add_run(cs)

    def _render_indicators(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Indicators', level=2)
        if lp.indicators:
            for ind in lp.indicators:
                doc.add_paragraph(ind, style='List Bullet')
        elif lp.indicator_codes:
            for code in lp.indicator_codes:
                doc.add_paragraph(code, style='List Bullet')

    def _render_objectives(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Learning Objectives', level=2)
        # Real data only: never print a invented objective stem. A lesson with
        # no stored objectives renders an empty section rather than filler.
        for obj in lp.learning_objectives:
            doc.add_paragraph(obj.description, style='List Bullet')

    def _render_resources(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Teaching & Learning Resources', level=2)
        # No stored resources renders blank — never fabricated TLR suggestions.
        for r in lp.teaching_learning_resources or []:
            doc.add_paragraph(r, style='List Bullet')

    def _render_introduction(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Introduction / Starter', level=2)
        if lp.introduction:
            doc.add_paragraph(lp.introduction)

    def _render_main_activities(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Main Teaching Activities', level=2)
        for act in lp.main_activities:
            p = doc.add_paragraph()
            run = p.add_run(f"{act.phase}: ")
            run.bold = True
            p.add_run(act.description)
            if act.duration_minutes:
                p.add_run(f" ({act.duration_minutes} min)")

    def _render_learner_activities(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Learner Activities', level=2)
        for act in lp.learner_activities:
            p = doc.add_paragraph()
            p.add_run(act.description)
            if act.duration_minutes:
                p.add_run(f" ({act.duration_minutes} min)")

    def _render_teacher_activities(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Teacher Activities', level=2)
        if lp.teacher_activities:
            for act in lp.teacher_activities:
                doc.add_paragraph(act.description)

    def _render_assessment(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Assessment', level=2)
        if lp.assessment:
            doc.add_paragraph(lp.assessment)

    def _render_conclusion(self, doc: Document, lp: LessonPlan):
        doc.add_heading('Conclusion / Reflection', level=2)
        if lp.conclusion:
            doc.add_paragraph(lp.conclusion)

    def _render_previous_knowledge(self, doc: Document, lp: LessonPlan):
        if lp.previous_knowledge:
            doc.add_heading('Previous Knowledge', level=2)
            doc.add_paragraph(lp.previous_knowledge)

    def _render_core_competencies(self, doc: Document, lp: LessonPlan):
        if lp.core_competencies:
            doc.add_heading('Core Competencies', level=2)
            for c in lp.core_competencies:
                doc.add_paragraph(c, style='List Bullet')

    def _render_references(self, doc: Document, lp: LessonPlan):
        if lp.references:
            doc.add_heading('References', level=2)
            for r in lp.references:
                doc.add_paragraph(r, style='List Bullet')


#: Official GES/NaCCA form to use per educational level.
_GES_LEVEL_KEY_BY_LEVEL = {
    EducationalLevel.EARLY_CHILDHOOD: LEVEL_KG,
    EducationalLevel.PRIMARY: LEVEL_PRIMARY,
    EducationalLevel.JHS: LEVEL_JHS,
    EducationalLevel.SHS: LEVEL_SHS,
}


def _educational_level_of(lesson) -> Optional[EducationalLevel]:
    """Educational level of a lesson plan model or dict, or None when unknown."""
    raw = (lesson.get("educational_level") if isinstance(lesson, dict)
           else getattr(lesson, "educational_level", None))
    if isinstance(raw, EducationalLevel):
        return raw
    if isinstance(raw, str):
        for level in EducationalLevel:
            if raw in (level.value, level.name):
                return level
    raw_class = (lesson.get("class_level") if isinstance(lesson, dict)
                 else getattr(lesson, "class_level", None))
    if isinstance(raw_class, ClassLevel):
        return CLASS_LEVEL_TO_EDUCATIONAL_LEVEL.get(raw_class)
    if isinstance(raw_class, str):
        for class_level in ClassLevel:
            if raw_class in (class_level.value, class_level.name):
                return CLASS_LEVEL_TO_EDUCATIONAL_LEVEL.get(class_level)
    return None


def default_template_for_lessons(lesson_plans, template_type: TemplateType = None) -> Template:
    """Template for an export that named none.

    Every official GES/NaCCA form belongs to one level, so the batch's own class
    level decides: a KG export must not fall back to the head of
    DEFAULT_TEMPLATES (the JHS form). Classes with no official form — Basic 4-6,
    because the GES primary form is a LOWER PRIMARY form — fall back to their
    family's default template.
    """
    for lesson in lesson_plans or []:
        class_level = (lesson.get("class_level") if isinstance(lesson, dict)
                       else getattr(lesson, "class_level", None))
        if isinstance(class_level, str):
            try:
                class_level = ClassLevel(class_level)
            except ValueError:
                continue
        if isinstance(class_level, ClassLevel):
            return get_default_template_for_class_level(class_level)
    return get_template_by_type(template_type or TemplateType.GES_STYLE)


def ges_spec_for_lessons(lesson_plans) -> GESLevelSpec:
    """Level contract to render a batch with when no template id was given.

    A KG batch must not render a JHS form, so the lessons' own educational level
    decides. JHS is the last resort, matching the head of DEFAULT_TEMPLATES.
    """
    for lesson in lesson_plans or []:
        level = _educational_level_of(lesson)
        if level is not None:
            return spec_for_level(_GES_LEVEL_KEY_BY_LEVEL.get(level, LEVEL_JHS)) or JHS_SPEC
    return JHS_SPEC


def _lesson_field_value(lp, field: Optional[str]) -> str:
    """Render one canonical SchemeKnit field from a lesson plan (model or dict)."""
    if not field:
        return ""
    if isinstance(lp, dict):
        raw = lp.get(field)
    else:
        raw = getattr(lp, field, None)
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return str(raw)
    if isinstance(raw, (int, float)):
        return str(raw)
    # date / datetime
    if hasattr(raw, "strftime"):
        try:
            return raw.strftime("%d/%m/%Y")
        except Exception:
            return str(raw)
    if isinstance(raw, str):
        return raw
    if isinstance(raw, (list, tuple)):
        parts = []
        for item in raw:
            if item is None:
                continue
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                parts.append(item.get("description") or item.get("text") or str(item))
            else:
                desc = getattr(item, "description", None)
                if desc:
                    phase = getattr(item, "phase", None)
                    dur = getattr(item, "duration_minutes", None)
                    line = f"{phase}: {desc}" if phase else str(desc)
                    if dur:
                        line += f" ({dur} min)"
                    parts.append(line)
                else:
                    parts.append(str(item))
        return "\n".join(p for p in parts if p)
    # Enum-like
    val = getattr(raw, "value", raw)
    return str(val)


def _activity_lines(lp, key: str) -> list:
    """Extract description lines from an activity list (model or dict items)."""
    raw = lp.get(key) if isinstance(lp, dict) else getattr(lp, key, None)
    lines = []
    for item in raw or []:
        if isinstance(item, str):
            lines.append(item)
        elif isinstance(item, dict):
            desc = item.get("description") or item.get("text")
            if desc:
                dur = item.get("duration_minutes")
                lines.append(f"{desc} ({dur} min)" if dur else desc)
        else:
            desc = getattr(item, "description", None)
            if desc:
                dur = getattr(item, "duration_minutes", None)
                lines.append(f"{desc} ({dur} min)" if dur else str(desc))
    return [l for l in lines if l]


def _activity_block(lp, kind: str) -> str:
    """Positional delivery-grid content by row phase kind.

    starter → introduction; main/unknown → learner + main activities;
    reflection → assessment + conclusion.
    """
    if kind == "starter":
        return _lesson_field_value(lp, "introduction")
    if kind == "reflection":
        parts = [_lesson_field_value(lp, "assessment"), _lesson_field_value(lp, "conclusion")]
        return "\n".join(p for p in parts if p)
    lines = _activity_lines(lp, "learner_activities") + _activity_lines(lp, "main_activities")
    return "\n".join(lines)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)


def render_custom_template(doc: Document, lp, structure: dict, render_context: dict = None):
    """Render a lesson plan through a saved sample TemplateStructure.

    Rules (presentation vs content separation):
    - table count/order/dimensions/merges reproduce the sample grid
    - label cells keep the sample's label text (bold)
    - value cells are filled from mapped SchemeKnit fields; unknown/custom
      fields and sample prose render blank (sample content is never copied)
    - phase-marker cells keep their structural marker text
    """
    tables = (structure or {}).get("tables", [])
    if not tables:
        # Fallback: render key fields as a simple grid rather than failing
        doc.add_heading("LESSON PLAN", level=1)
        for field in ["school_name", "teacher_name", "subject", "class_level",
                      "lesson_date", "strand", "content_standard", "indicators",
                      "introduction", "assessment", "conclusion"]:
            val = _lesson_field_value(lp, field)
            if val:
                doc.add_heading(field.replace("_", " ").title(), level=2)
                doc.add_paragraph(val)
        doc.add_page_break()
        return

    # Optional title from sample meta; approved templates use title_lines
    # with {term} / {week} placeholders. Scheme context (term) is supplied by
    # the caller when known; otherwise the line is skipped rather than invented.
    meta = (structure or {}).get("meta") or {}
    title_lines = meta.get("title_lines") or ([meta.get("title")] if meta.get("title") else [])
    if title_lines:
        context = render_context or {}
        def _field(name, default=""):
            v = _lesson_field_value(lp, name)
            return v if v else default
        term = str(context.get("term") or _field("term") or _field("academic_year") or "").upper()
        mapping = {
            "term": term,
            "week": str(_field("week_number") or ""),
            "subject": str(_field("subject") or ""),
            "class": str(_field("class_level") or ""),
        }
        for line in title_lines:
            try:
                text = str(line).format(**mapping)
            except Exception:
                text = str(line)
            if text.strip():
                h = doc.add_heading(text.strip(), level=1)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for table in tables:
        rows, cols = int(table.get("rows", 0)), int(table.get("cols", 0))
        if rows <= 0 or cols <= 0:
            continue
        grid = doc.add_table(rows=rows, cols=cols)
        grid.style = "Table Grid"
        grid.alignment = WD_TABLE_ALIGNMENT.CENTER
        # write master cells (skip covered slots)
        written = {}
        for c in table.get("cells", []):
            if c.get("covered"):
                continue
            r, cc = int(c.get("r", 0)), int(c.get("c", 0))
            if r >= rows or cc >= cols:
                continue
            cell = grid.cell(r, cc)
            label = c.get("label")
            field = c.get("field")
            role = c.get("role")
            if c.get("structural_header"):
                _set_cell_text(cell, c.get("text", ""), bold=c.get("bold", True))
            elif c.get("is_phase_marker"):
                _set_cell_text(cell, c.get("text", ""), bold=True)
            elif role == "activity_content":
                _set_cell_text(cell, _activity_block(lp, c.get("row_phase_kind", "unknown")))
            elif role == "resource_content":
                _set_cell_text(cell, _lesson_field_value(lp, "teaching_learning_resources"))
            elif label and field:
                val = _lesson_field_value(lp, field)
                _set_cell_text(cell, f"{label}: {val}" if val else f"{label}:", bold=False)
                # bold the label run only
                try:
                    if cell.paragraphs and cell.paragraphs[0].runs:
                        full = cell.paragraphs[0].text
                        cell.text = ""
                        p = cell.paragraphs[0]
                        lab_run = p.add_run(f"{label}: ")
                        lab_run.bold = True
                        lab_run.font.size = Pt(10)
                        v_run = p.add_run(val)
                        v_run.font.size = Pt(10)
                except Exception:
                    pass
            elif label:
                # custom/unknown label: keep label, blank value (never copy sample data)
                _set_cell_text(cell, f"{label}:", bold=False)
                try:
                    if cell.paragraphs and cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].bold = True
                except Exception:
                    pass
            else:
                # unlabeled prose cell from the sample: leave blank
                _set_cell_text(cell, "")
            written[(r, cc)] = cell
        # apply merges (horizontal then vertical)
        for c in table.get("cells", []):
            r, cc = int(c.get("r", 0)), int(c.get("c", 0))
            span = int(c.get("colspan", 1) or 1)
            if span > 1 and (r, cc) in written and cc + span - 1 < cols:
                try:
                    written[(r, cc)].merge(grid.cell(r, cc + span - 1))
                except Exception:
                    pass
        # vertical merges: restart cell merges down through consecutive covered slots
        covered = {(int(c.get("r", 0)), int(c.get("c", 0)))
                   for c in table.get("cells", []) if c.get("covered")}
        for c in table.get("cells", []):
            if c.get("covered") or not c.get("v_merge_restart"):
                continue
            r, cc = int(c.get("r", 0)), int(c.get("c", 0))
            end_r = r
            while (end_r + 1, cc) in covered and end_r + 1 < rows:
                end_r += 1
            if end_r > r and (r, cc) in written:
                try:
                    written[(r, cc)].merge(grid.cell(end_r, cc))
                except Exception:
                    pass
        doc.add_paragraph("")
    doc.add_page_break()


SECTION_RENDERERS = {
    "header": DOCXExportEngine._render_header,
    "strand_info": DOCXExportEngine._render_strand_info,
    "content_standard": DOCXExportEngine._render_content_standard,
    "indicators": DOCXExportEngine._render_indicators,
    "objectives": DOCXExportEngine._render_objectives,
    "resources": DOCXExportEngine._render_resources,
    "introduction": DOCXExportEngine._render_introduction,
    "main_activities": DOCXExportEngine._render_main_activities,
    "learner_activities": DOCXExportEngine._render_learner_activities,
    "teacher_activities": DOCXExportEngine._render_teacher_activities,
    "assessment": DOCXExportEngine._render_assessment,
    "conclusion": DOCXExportEngine._render_conclusion,
    "previous_knowledge": DOCXExportEngine._render_previous_knowledge,
    "core_competencies": DOCXExportEngine._render_core_competencies,
    "references": DOCXExportEngine._render_references,
}
