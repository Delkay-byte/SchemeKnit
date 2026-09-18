"""
SchemeKnit XLSX Export Engine

Generates curriculum/lesson register in Excel format.
"""

from pathlib import Path
from typing import List, Optional
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from ..models import LessonPlan, Subject, ClassLevel


class XLSXExportEngine:
    """Excel register export."""

    HEADERS = [
        "Week", "Week Ending", "Lesson Date", "Lesson #",
        "Strand", "Sub-Strand", "Content Standard",
        "Indicator", "Topic", "Status",
    ]

    def export_register(
        self,
        lesson_plans: List[LessonPlan],
        output_path: Path,
        title: str = "Lesson Plan Register",
    ) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Register"

        self._write_title(ws, title)
        self._write_headers(ws)
        self._write_data(ws, lesson_plans)
        self._auto_width(ws)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        return output_path

    def _write_title(self, ws, title: str):
        ws.merge_cells('A1:J1')
        cell = ws['A1']
        cell.value = title
        cell.font = Font(name='Arial', size=14, bold=True, color='003366')
        cell.alignment = Alignment(horizontal='center')
        ws.row_dimensions[1].height = 30

    def _write_headers(self, ws):
        header_fill = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
        header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        for col, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        ws.row_dimensions[3].height = 25

    def _write_data(self, ws, lesson_plans: List[LessonPlan]):
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        alt_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

        for idx, lp in enumerate(lesson_plans, 4):
            subject_val = lp.subject.value if isinstance(lp.subject, Subject) else str(lp.subject)
            class_val = lp.class_level.value if isinstance(lp.class_level, ClassLevel) else str(lp.class_level)

            strand = lp.strand or ""
            sub = lp.sub_strand or ""
            cs = lp.content_standard or ""
            indicators = "; ".join(lp.indicators) if lp.indicators else ""
            topic = lp.lesson_topic or f"{strand} - {sub}"
            status = lp.status.value if hasattr(lp.status, 'value') else str(lp.status)
            date_val = lp.lesson_date.strftime("%d/%m/%Y") if lp.lesson_date else ""

            row_data = [
                lp.week_number,
                date_val,
                date_val,
                lp.lesson_sequence,
                strand,
                sub,
                cs,
                indicators,
                topic,
                status,
            ]

            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=idx, column=col, value=val)
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = thin_border
                if idx % 2 == 0:
                    cell.fill = alt_fill

    def _auto_width(self, ws):
        col_widths = [8, 12, 12, 8, 20, 20, 30, 30, 25, 10]
        for i, width in enumerate(col_widths, 1):
            col_letter = chr(64 + i)
            ws.column_dimensions[col_letter].width = width
