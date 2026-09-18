"""
SchemeKnit Sample Template Analyzer (deterministic, no LLM).

Analyzes an uploaded sample lesson-plan DOCX and produces a normalized
TemplateStructure intermediate representation:

- document metadata (paragraphs, tables, title, layout)
- per-table grids with merged-cell maps (gridSpan + vMerge)
- per-cell label/value split ("Label: value", "Label:\\nvalue")
- field detection against the canonical SchemeKnit field vocabulary,
  with confidence (high / medium / manual) and match provenance
- unknown labels preserved as custom-keep fields (never silently dropped)

The structure is persisted with the custom template and later drives
DOCX rendering (see docx_export.render_custom_template).
"""

import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn


# ── Canonical SchemeKnit lesson-plan field vocabulary ──────────────────────────
# Every key is a renderable LessonPlan attribute (or accepted custom slot).

TEACHFLOW_FIELDS = [
    "school_name",
    "teacher_name",
    "lesson_date",
    "class_level",
    "subject",
    "class_size",
    "duration_minutes",
    "week_number",
    "lesson_number",
    "term",
    "academic_year",
    "strand",
    "sub_strand",
    "content_standard",
    "content_standard_code",
    "indicators",
    "indicator_codes",
    "learning_objectives",
    "core_competencies",
    "teaching_learning_resources",
    "previous_knowledge",
    "introduction",
    "main_activities",
    "learner_activities",
    "teacher_activities",
    "assessment",
    "conclusion",
    "references",
    "lesson_topic",
    "homework",
]

# Normalized label (lowercase, no trailing colon, collapsed whitespace)
# -> (SchemeKnit field, confidence, match kind)
LABEL_ALIASES: Dict[str, Tuple[str, str, str]] = {
    # Exact canonical labels
    "school": ("school_name", "high", "exact"),
    "school name": ("school_name", "high", "exact"),
    "teacher": ("teacher_name", "high", "exact"),
    "teacher name": ("teacher_name", "high", "exact"),
    "name of teacher": ("teacher_name", "high", "alias"),
    "date": ("lesson_date", "high", "exact"),
    "lesson date": ("lesson_date", "high", "alias"),
    "class": ("class_level", "high", "exact"),
    "subject": ("subject", "high", "exact"),
    "class size": ("class_size", "high", "exact"),
    "duration": ("duration_minutes", "high", "exact"),
    "week": ("week_number", "high", "exact"),
    "lesson": ("lesson_number", "high", "exact"),
    "lesson number": ("lesson_number", "high", "alias"),
    "term": ("term", "high", "exact"),
    "strand": ("strand", "high", "exact"),
    "sub strand": ("sub_strand", "high", "exact"),
    "sub-strand": ("sub_strand", "high", "exact"),
    "content standard": ("content_standard", "high", "exact"),
    "content standards": ("content_standard", "high", "alias"),
    "standard code": ("content_standard_code", "high", "alias"),
    "indicator": ("indicators", "high", "exact"),
    "indicators": ("indicators", "high", "exact"),
    "performance indicator": ("learning_objectives", "medium", "alias"),
    "performance indicators": ("learning_objectives", "medium", "alias"),
    "objectives": ("learning_objectives", "high", "alias"),
    "learning objectives": ("learning_objectives", "high", "exact"),
    "core competencies": ("core_competencies", "high", "exact"),
    "core competency": ("core_competencies", "high", "alias"),
    "resources": ("teaching_learning_resources", "high", "alias"),
    "teaching learning resources": ("teaching_learning_resources", "high", "exact"),
    "teaching/learning resources": ("teaching_learning_resources", "high", "alias"),
    "tlm": ("teaching_learning_resources", "medium", "alias"),
    "instructional materials": ("teaching_learning_resources", "medium", "alias"),
    "previous knowledge": ("previous_knowledge", "high", "exact"),
    "prior knowledge": ("previous_knowledge", "high", "alias"),
    "introduction": ("introduction", "high", "exact"),
    "starter": ("introduction", "medium", "alias"),
    "main activities": ("main_activities", "high", "exact"),
    "teacher activities": ("teacher_activities", "high", "exact"),
    "learner activities": ("learner_activities", "high", "exact"),
    "learner activity": ("learner_activities", "high", "alias"),
    "pupil activities": ("learner_activities", "medium", "alias"),
    "assessment": ("assessment", "high", "exact"),
    "evaluation": ("assessment", "medium", "alias"),
    "conclusion": ("conclusion", "high", "exact"),
    "reflection": ("conclusion", "medium", "alias"),
    "summary": ("conclusion", "medium", "alias"),
    "references": ("references", "high", "exact"),
    "reference": ("references", "high", "alias"),
    "topic": ("lesson_topic", "high", "alias"),
    "lesson topic": ("lesson_topic", "high", "exact"),
    "homework": ("homework", "high", "exact"),
    "assignment": ("homework", "medium", "alias"),
}

# Cells matching these are structural section/phase markers, not fields.
# A digit is required so delivery-grid headers ("Phase/Duration") are never
# misclassified as markers.
PHASE_MARKER_RE = re.compile(r"^\s*phase\s+\d+", re.IGNORECASE)

# Header-row keyword sets for the lesson-delivery (activity) grid:
# e.g. "Phase/Duration | Learners Activities | Resources".
PHASE_COL_KEYWORDS = ("phase", "duration", "stage", "step", "time")
ACTIVITY_COL_KEYWORDS = ("activit", "learner", "pupil", "student", "task", "group work", "class work")
RESOURCE_COL_KEYWORDS = ("resource", "material", "tlm", "tlrs", "teaching aid", "apparatus")


def _col_kind(label: Optional[str]) -> Optional[str]:
    """Classify a header label into a delivery-grid column role."""
    if not label:
        return None
    norm = normalize_label(label)
    if any(k in norm for k in PHASE_COL_KEYWORDS) and not any(
            k in norm for k in ACTIVITY_COL_KEYWORDS + RESOURCE_COL_KEYWORDS):
        return "phase"
    if any(k in norm for k in ACTIVITY_COL_KEYWORDS):
        return "activity"
    if any(k in norm for k in RESOURCE_COL_KEYWORDS):
        return "resource"
    return None


def phase_kind(marker_text: str) -> str:
    """Classify a phase marker into starter | main | reflection."""
    t = (marker_text or "").lower()
    if re.search(r"starter|intro|warm|begin|previous|recap|review", t):
        return "starter"
    if re.search(r"reflect|assess|evaluat|clos|summar|plenary|conclusion", t):
        return "reflection"
    return "main"


def normalize_label(text: str) -> str:
    """Normalize a cell label for alias lookup."""
    text = text.strip().rstrip(":").strip()
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def detect_field(label: str) -> Dict:
    """Map a label to a SchemeKnit field with confidence.

    Returns {field|None, confidence: high|medium|manual, match_kind, custom}.
    Unknown labels are preserved as custom-keep (field=None) — never dropped.
    """
    norm = normalize_label(label)
    if not norm:
        return {"field": None, "confidence": "manual", "match_kind": "empty", "custom": False}
    if norm in LABEL_ALIASES:
        field, confidence, kind = LABEL_ALIASES[norm]
        return {"field": field, "confidence": confidence, "match_kind": kind, "custom": False}
    # Keyword-contains fallback (medium/low confidence, teacher must confirm)
    for alias, (field, _, _) in LABEL_ALIASES.items():
        if len(alias) >= 5 and alias in norm:
            return {"field": field, "confidence": "medium", "match_kind": "contains", "custom": False}
    return {"field": None, "confidence": "manual", "match_kind": "unknown", "custom": True}


def split_label_value(text: str) -> Tuple[Optional[str], str]:
    """Split 'Label: value' / 'Label:\\nvalue' / 'Label: / value' cells.

    Returns (label|None, value). A bare 'Label:' yields (label, '').
    """
    if not text or not text.strip():
        return None, ""
    # First line up to ':' or newline is the candidate label
    m = re.match(r"^\s*([^:\n]{1,60}?)\s*:\s*(.*)$", text, re.DOTALL)
    if m:
        label, rest = m.group(1).strip(), m.group(2).strip()
        # Strip sample separators like leading '/' used in some teacher docs
        rest = re.sub(r"^/+\s*", "", rest)
        if label and len(label) <= 60:
            return label, rest
    return None, text.strip()


def _grid_span(tc) -> int:
    tcPr = tc.find(qn("w:tcPr"))
    gs = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
    try:
        return int(gs.get(qn("w:val"))) if gs is not None else 1
    except (ValueError, TypeError):
        return 1


def _v_merge(tc) -> Optional[str]:
    """Return 'restart' for merge start, 'continue' for covered cells, None otherwise."""
    tcPr = tc.find(qn("w:tcPr"))
    vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
    if vm is None:
        return None
    val = vm.get(qn("w:val"))
    return "continue" if (val is None or val == "continue") else "restart"


def _cell_bold(cell) -> bool:
    try:
        for p in cell.paragraphs:
            for r in p.runs:
                if r.bold:
                    return True
    except Exception:
        pass
    return False


def _analyze_table(table, index: int) -> Dict:
    """Build a normalized grid for one table with merge maps and field detection."""
    tc_rows = [row._tr.findall(qn("w:tc")) for row in table.rows]
    n_grid_cols = 0
    for tcs in tc_rows:
        n_grid_cols = max(n_grid_cols, sum(_grid_span(tc) for tc in tcs))

    cells: List[Dict] = []
    covered = set()  # (r, c) positions covered by merges
    for ri, (row, tcs) in enumerate(zip(table.rows, tc_rows)):
        cursor = 0
        tc_iter = list(zip(tcs, row.cells)) if len(tcs) == len(row.cells) else None
        for tc in tcs:
            # advance cursor past covered slots
            while (ri, cursor) in covered:
                cursor += 1
            span = _grid_span(tc)
            vm = _v_merge(tc)
            # find display cell for text (best effort alignment with python-docx cells)
            text = ""
            bold = False
            try:
                # locate matching _Cell by tc identity
                for c in row.cells:
                    if c._tc is tc:
                        text = c.text.strip()
                        bold = _cell_bold(c)
                        break
            except Exception:
                pass
            if vm == "continue":
                covered.add((ri, cursor))
                # mark the whole span covered
                for k in range(1, span):
                    covered.add((ri, cursor + k))
                # persist covered slots so the renderer can rebuild vertical merges
                cells.append({
                    "r": ri,
                    "c": cursor,
                    "rowspan": 1,
                    "colspan": span,
                    "v_merge_restart": False,
                    "covered": True,
                    "text": "",
                    "label": None,
                    "sample_value": "",
                    "is_phase_marker": False,
                    "field": None,
                    "confidence": "manual",
                    "match_kind": "covered",
                    "custom": False,
                    "bold": False,
                })
                cursor += span
                continue
            # mark horizontal span coverage
            for k in range(1, span):
                covered.add((ri, cursor + k))
            # mark vertical restart coverage for rows below is implicit via continue cells
            label, value = split_label_value(text)
            is_phase = bool(label and PHASE_MARKER_RE.match(label)) or bool(
                not label and PHASE_MARKER_RE.match(text)
            )
            det = {"field": None, "confidence": "manual", "match_kind": "none", "custom": False}
            if label and not is_phase:
                det = detect_field(label)
            cells.append({
                "r": ri,
                "c": cursor,
                "rowspan": 1,
                "colspan": span,
                "v_merge_restart": vm == "restart",
                "text": text,
                "label": label,
                "sample_value": value,
                "is_phase_marker": is_phase,
                "field": det["field"],
                "confidence": det["confidence"],
                "match_kind": det["match_kind"],
                "custom": det["custom"],
                "bold": bold,
            })
            cursor += span

    grid = {
        "index": index,
        "rows": len(table.rows),
        "cols": n_grid_cols,
        "cells": cells,
    }
    _detect_activity_grid(grid)
    return grid


def _detect_activity_grid(grid: Dict):
    """Detect a lesson-delivery header row and assign column roles + row phase kinds.

    Header example: Phase/Duration | Learners Activities | Resources.
    Body rows then carry activity_content / resource_content roles so the
    renderer can place generated activities and resources positionally
    (these cells have no labels in the sample).
    Tables without a header but with phase markers get per-row inference.
    """
    by_row: Dict[int, list] = {}
    for c in grid["cells"]:
        if c.get("covered"):
            continue
        by_row.setdefault(c["r"], []).append(c)

    grid["activity_grid"] = None

    # 1) Header-row detection (labels, falling back to raw cell text —
    #    delivery-grid headers like "Phase/Duration" carry no colon).
    #    Rows containing phase markers are never headers (they are body rows,
    #    including in continuation tables) — this prevents filled body content
    #    from re-classifying as a header on output re-analysis.
    for ri in sorted(by_row):
        if any(c.get("is_phase_marker") for c in by_row[ri]):
            continue
        kinds = {}
        for c in by_row[ri]:
            k = _col_kind(c.get("label") or c.get("text"))
            if k:
                kinds.setdefault(k, []).append(c)
        if "activity" in kinds and ("phase" in kinds or "resource" in kinds) \
                and sum(len(v) for v in kinds.values()) >= 2:
            phase_cols, act_cols, res_cols = set(), set(), set()
            for c in kinds.get("phase", []):
                phase_cols.update(range(c["c"], c["c"] + c["colspan"]))
            for c in kinds.get("activity", []):
                act_cols.update(range(c["c"], c["c"] + c["colspan"]))
            for c in kinds.get("resource", []):
                res_cols.update(range(c["c"], c["c"] + c["colspan"]))
            grid["activity_grid"] = {
                "header_row": ri,
                "phase_cols": sorted(phase_cols),
                "activity_cols": sorted(act_cols),
                "resource_cols": sorted(res_cols),
            }
            for c in by_row[ri]:
                c["structural_header"] = True
            break

    ag = grid["activity_grid"]
    if ag:
        for ri in sorted(by_row):
            if ri == ag["header_row"]:
                continue
            # row phase kind from its phase-column marker text
            marker = ""
            for c in by_row[ri]:
                if c["c"] in ag["phase_cols"] and c.get("text"):
                    marker = c["text"]
                    break
            kind = phase_kind(marker) if marker else "unknown"
            for c in by_row[ri]:
                c["row_phase_kind"] = kind
                if c.get("label") or c.get("is_phase_marker") or c.get("structural_header"):
                    continue
                if c["c"] in ag["activity_cols"]:
                    c["role"] = "activity_content"
                elif c["c"] in ag["resource_cols"]:
                    c["role"] = "resource_content"
        return

    # 2) No header: infer from phase markers (continuation tables)
    for ri in sorted(by_row):
        markers = [c for c in by_row[ri] if c.get("is_phase_marker")]
        if not markers:
            continue
        kind = phase_kind(markers[0].get("text", ""))
        candidates = [c for c in by_row[ri]
                      if not c.get("label") and not c.get("is_phase_marker")
                      and (c.get("text") is not None)]
        for c in by_row[ri]:
            c["row_phase_kind"] = kind
        if candidates:
            widest = max(candidates, key=lambda c: c.get("colspan", 1))
            widest["role"] = "activity_content"
            widest["role_inferred"] = True


def _extract_layout(doc) -> Dict:
    layout = {"page_size": "A4", "orientation": "portrait", "font_family": "Arial", "font_size": 11}
    try:
        section = doc.sections[0]
        w = section.page_width / 914400
        h = section.page_height / 914400
        if w > h:
            layout["orientation"] = "landscape"
        if abs(w - 8.5) < 0.5 and abs(h - 11) < 0.5:
            layout["page_size"] = "Letter"
    except Exception:
        pass
    try:
        for para in doc.paragraphs:
            if para.runs:
                run = para.runs[0]
                if run.font.name:
                    layout["font_family"] = run.font.name
                if run.font.size:
                    layout["font_size"] = int(run.font.size.pt)
                break
    except Exception:
        pass
    return layout


def validate_custom_docx(sample_structure: Dict, output_path: str,
                         expected_values: Optional[List[str]] = None,
                         forbidden_values: Optional[List[str]] = None) -> Dict:
    """Deterministic structural validation of a generated DOCX against the sample IR.

    Inspects table topology (count, dimensions, grid spans), label presence,
    expected lesson-value placement, and absence of foreign sample content.
    Compares structure, not just plain text.
    """
    checks: List[Dict] = []

    def check(name: str, ok: bool, detail: str = ""):
        checks.append({"name": name, "pass": bool(ok), "detail": detail})

    try:
        out = analyze_docx_sample(output_path)
    except Exception as e:
        check("output_parseable", False, str(e))
        return {"pass": False, "checks": checks}

    s_tables = sample_structure.get("tables", [])
    o_tables = out.get("tables", [])
    if s_tables and len(o_tables) != len(s_tables) and len(o_tables) % len(s_tables) == 0:
        # Multi-lesson document: each lesson repeats the sample table block.
        # Validate every lesson block and aggregate.
        n = len(s_tables)
        chunks = [o_tables[i:i + n] for i in range(0, len(o_tables), n)]
        check("table_count", True,
              f"sample={len(s_tables)} output={len(o_tables)} lessons={len(chunks)}")
        block_results = [_validate_block(s_tables, chunk) for chunk in chunks]
        for name in block_results[0]:
            bad = [(i, b[name][1]) for i, b in enumerate(block_results) if not b[name][0]]
            detail = "; ".join(f"lesson{i + 1}:{d}" for i, d in bad) or f"{len(chunks)} lessons ok"
            check(name, not bad, detail)
    else:
        check("table_count", len(o_tables) == len(s_tables),
              f"sample={len(s_tables)} output={len(o_tables)}")
        for name, (ok, detail) in _validate_block(s_tables, o_tables).items():
            check(name, ok, detail)

    sample_labels = {m["label"] for m in sample_structure.get("mappings", []) if m.get("label")}
    out_text = " ".join(
        c.get("text", "") for t in out.get("tables", []) for c in t.get("cells", [])
    )
    try:
        _d = Document(output_path)
        out_text += " " + " ".join(p.text for p in _d.paragraphs if p.text.strip())
    except Exception:
        pass
    missing = sorted(l for l in sample_labels if l not in out_text)
    check("labels_present", not missing, f"missing={missing[:5]}")

    for val in expected_values or []:
        check(f"value_present:{val[:30]}", val in out_text)
    for val in forbidden_values or []:
        check(f"foreign_absent:{val[:30]}", val not in out_text)

    overall = all(c["pass"] for c in checks)
    return {"pass": overall, "checks": checks}


def _validate_block(s_tables, o_tables) -> Dict[str, tuple]:
    """Compare one lesson-block of output tables against the sample block.

    Topology only for merges (span + vMerge multisets): filled lesson content
    can legitimately mimic a label on re-analysis and must not fail topology.
    Returns {check_name: (pass, detail)}.
    """
    def dims(t):
        return (t.get("rows"), t.get("cols"))

    def spans(t):
        return sorted(
            (c.get("colspan", 1), c.get("v_merge_restart", False))
            for c in t.get("cells", []) if not c.get("covered")
        )

    def grids(t):
        return sum(1 for _ in [t] if t.get("activity_grid"))

    return {
        "table_dimensions": (
            [dims(t) for t in o_tables] == [dims(t) for t in s_tables],
            f"sample={[dims(t) for t in s_tables]} output={[dims(t) for t in o_tables]}",
        ),
        "merge_pattern": (
            [spans(t) for t in o_tables] == [spans(t) for t in s_tables],
            "gridSpan/vMerge multisets per table",
        ),
        "activity_grids_preserved": (
            sum(grids(t) for t in o_tables) == sum(grids(t) for t in s_tables),
            f"sample={sum(grids(t) for t in s_tables)} "
            f"output={sum(grids(t) for t in o_tables)}",
        ),
    }


def _detect_level(text_blob: str) -> Tuple[str, str]:
    blob = text_blob.lower()
    if any(k in blob for k in ["nursery", "kg 1", "kg 2", "kindergarten", "early childhood"]):
        return "early_childhood", "Early Childhood"
    if any(k in blob for k in ["basic 1", "basic 2", "basic 3", "basic 4", "basic 5", "basic 6", "primary"]):
        return "primary", "Primary"
    if any(k in blob for k in ["basic 7", "basic 8", "basic 9", "jhs", "junior high"]):
        return "jhs", "Junior High School"
    if any(k in blob for k in ["shs", "senior high", "basic 10", "basic 11", "basic 12"]):
        return "shs", "Senior High School"
    return "jhs", "Junior High School"


def analyze_docx_document(doc, file_path: Optional[str] = None) -> Dict:
    """Analyze an already-open lesson-plan DOCX into a TemplateStructure IR.

    ``analyze_docx_sample`` opens a file; this entry point accepts an in-memory
    ``Document`` so the approved-template golden master can derive its IR from
    the real source form (after stripping its authoring guidance) instead of
    hand-transcribing a structure that can drift from the source.
    """
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs and not doc.tables:
        raise ValueError("Document has no readable content")

    title = ""
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            title = t
            break

    tables = [_analyze_table(t, i) for i, t in enumerate(doc.tables)]
    text_blob = " ".join(paragraphs)
    for t in tables:
        for c in t["cells"]:
            text_blob += " " + (c["text"] or "")
    family, level = _detect_level(text_blob)

    # Consolidated mapping list (dedup by label, first occurrence wins)
    mappings: List[Dict] = []
    seen = set()
    for t in tables:
        for c in t["cells"]:
            if c["label"] and c["label"] not in seen and not c["is_phase_marker"]:
                seen.add(c["label"])
                mappings.append({
                    "label": c["label"],
                    "field": c["field"],
                    "confidence": c["confidence"],
                    "match_kind": c["match_kind"],
                    "custom": c["custom"],
                    "teacher_confirmed": False,
                })

    return {
        "meta": {
            "title": title,
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        },
        "layout": _extract_layout(doc),
        "tables": tables,
        "mappings": mappings,
        "detected_family": family,
        "detected_level": level,
        "sample_paragraphs": paragraphs[:20],
    }


def analyze_docx_sample(file_path: str) -> Dict:
    """Analyze a sample lesson-plan DOCX file into a TemplateStructure IR.

    Raises ValueError on malformed/empty documents.
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Malformed document: {e}")
    return analyze_docx_document(doc, file_path)
