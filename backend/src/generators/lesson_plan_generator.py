"""
SchemeKnit Lesson Plan Generator
"""

from typing import List, Optional, Dict, Any
from datetime import date, timedelta
import uuid

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from ..models import (
    SchemeOfWork, Week, Strand, SubStrand, ContentStandard, Indicator,
    Lesson, Activity, Resource, Assessment, GenerationConfig, Template,
    TemplateField, CurriculumAllocation
)


class LessonPlanGenerator:
    """Generates lesson plans from structured curriculum data."""
    
    def __init__(self):
        self.activity_templates = {
            "STARTER": [
                "Revise with learners on the previous lesson",
                "Call volunteer learners to the board to solve sample questions",
                "Introduce the lesson by sharing performance indicators",
                "Use questions and answers to review previous concepts"
            ],
            "NEW_LEARNING": [
                "Guide learners to explore the concept through demonstration",
                "Use local materials to illustrate the concept",
                "Engage learners in group discussion",
                "Provide worked examples step by step"
            ],
            "REFLECTION": [
                "Use peer discussion to find out what learners have learned",
                "Take feedback from learners and summarize the lesson",
                "Ask effective questions to check understanding",
                "Review key points and address misconceptions"
            ]
        }
    
    async def generate_lesson_plans(
        self,
        scheme: SchemeOfWork,
        config: GenerationConfig,
        template: Template,
        allocation: CurriculumAllocation
    ) -> List[Lesson]:
        """Generate lesson plans for the entire term."""
        lessons = []
        
        # Generate lessons for each week
        current_date = config.term_start_date
        
        for week in scheme.weeks:
            week_lessons = await self._generate_week_lessons(
                week, config, template, current_date
            )
            lessons.extend(week_lessons)
            
            # Move to next week
            current_date += timedelta(days=7)
        
        return lessons
    
    async def _generate_week_lessons(
        self,
        week: Week,
        config: GenerationConfig,
        template: Template,
        week_start_date: date
    ) -> List[Lesson]:
        """Generate lessons for a single week."""
        lessons = []
        
        # Calculate teaching days for this week
        teaching_days = self._get_teaching_days(week_start_date, config)
        
        # Allocate curriculum content to lessons
        curriculum_for_week = self._allocate_curriculum_to_week(week, config)
        
        # Generate lessons for each teaching day
        for i, teaching_day in enumerate(teaching_days[:config.lessons_per_week]):
            lesson_number = i + 1
            
            # Get curriculum for this lesson
            lesson_curriculum = self._get_curriculum_for_lesson(
                curriculum_for_week, i, config.lessons_per_week
            )
            
            # Create lesson object
            lesson = Lesson(
                id=str(uuid.uuid4()),
                week_id=week.id,
                lesson_number=lesson_number,
                date=teaching_day,
                period=lesson_number,
                duration_minutes=config.lesson_duration_minutes,
                strand=lesson_curriculum.get('strand'),
                sub_strand=lesson_curriculum.get('sub_strand'),
                content_standard=lesson_curriculum.get('content_standard'),
                indicator=lesson_curriculum.get('indicator'),
                objectives=lesson_curriculum.get('objectives', []),
                activities=self._generate_activities(lesson_curriculum, config),
                resources=lesson_curriculum.get('resources', []),
                assessments=self._generate_assessments(lesson_curriculum, config),
                status="generated"
            )
            
            lessons.append(lesson)
        
        return lessons
    
    def _get_teaching_days(self, week_start_date: date, config: GenerationConfig) -> List[date]:
        """Get teaching days for a week."""
        teaching_days = []
        
        # Check each day of the week
        for i in range(7):
            current_date = week_start_date + timedelta(days=i)
            
            # Check if it's a teaching day
            if current_date.weekday() in config.teaching_days:
                # Check if it's not a holiday
                is_holiday = False
                for holiday in config.holidays:
                    if holiday.date == current_date:
                        is_holiday = True
                        break
                
                if not is_holiday:
                    teaching_days.append(current_date)
        
        return teaching_days
    
    def _allocate_curriculum_to_week(self, week: Week, config: GenerationConfig) -> List[Dict[str, Any]]:
        """Allocate curriculum content to a week."""
        curriculum_items = []
        
        for strand in week.strands:
            for sub_strand in strand.sub_strands:
                for cs in sub_strand.content_standards:
                    for indicator in cs.indicators:
                        curriculum_items.append({
                            'strand': strand,
                            'sub_strand': sub_strand,
                            'content_standard': cs,
                            'indicator': indicator,
                            'objectives': indicator.objectives,
                            'resources': self._generate_resources(indicator, config)
                        })
        
        return curriculum_items
    
    def _get_curriculum_for_lesson(
        self, 
        curriculum_items: List[Dict[str, Any]], 
        lesson_index: int, 
        lessons_per_week: int
    ) -> Dict[str, Any]:
        """Get curriculum content for a specific lesson."""
        if not curriculum_items:
            return {}
        
        # Simple distribution: divide items equally among lessons
        items_per_lesson = max(1, len(curriculum_items) // lessons_per_week)
        start_index = lesson_index * items_per_lesson
        end_index = min(start_index + items_per_lesson, len(curriculum_items))
        
        if start_index >= len(curriculum_items):
            return curriculum_items[-1] if curriculum_items else {}
        
        # Return the first item for this lesson (can be enhanced later)
        return curriculum_items[start_index]
    
    def _generate_activities(self, curriculum: Dict[str, Any], config: GenerationConfig) -> List[Activity]:
        """Generate activities for a lesson."""
        activities = []
        
        # Generate starter activity
        starter_desc = self._select_activity("STARTER", curriculum)
        activities.append(Activity(
            phase="STARTER",
            description=starter_desc,
            duration_minutes=10,
            resources=[]
        ))
        
        # Generate new learning activity
        new_learning_desc = self._select_activity("NEW_LEARNING", curriculum)
        activities.append(Activity(
            phase="NEW_LEARNING",
            description=new_learning_desc,
            duration_minutes=40,
            resources=curriculum.get('resources', [])
        ))
        
        # Generate reflection activity
        reflection_desc = self._select_activity("REFLECTION", curriculum)
        activities.append(Activity(
            phase="REFLECTION",
            description=reflection_desc,
            duration_minutes=10,
            resources=[]
        ))
        
        return activities
    
    def _select_activity(self, phase: str, curriculum: Dict[str, Any]) -> str:
        """Select an appropriate activity for the phase."""
        import random
        
        templates = self.activity_templates.get(phase, [])
        if not templates:
            return f"Conduct {phase.lower()} activity"
        
        # Select a random template
        base_activity = random.choice(templates)
        
        # Customize based on curriculum
        indicator = curriculum.get('indicator')
        if indicator:
            # Add specific content to the activity
            return f"{base_activity} focusing on {indicator.description[:50]}..."
        
        return base_activity
    
    def _generate_resources(self, indicator: Indicator, config: GenerationConfig) -> List[Resource]:
        """Generate appropriate resources for a lesson."""
        resources = []
        
        # Basic resources based on subject
        if config.subject.value == "Mathematics":
            resources.extend([
                Resource(name="Ruler", type="MATERIAL"),
                Resource(name="Pencil", type="MATERIAL"),
                Resource(name="Exercise book", type="MATERIAL")
            ])
        elif config.subject.value == "Science":
            resources.extend([
                Resource(name="Charts", type="LOCAL"),
                Resource(name="Pictures", type="LOCAL"),
                Resource(name="Real objects", type="LOCAL")
            ])
        
        return resources
    
    def _generate_assessments(self, curriculum: Dict[str, Any], config: GenerationConfig) -> List[Assessment]:
        """Generate assessment items for a lesson."""
        assessments = []
        
        indicator = curriculum.get('indicator')
        if indicator:
            # Create a formative assessment based on the indicator
            assessments.append(Assessment(
                type="FORMATIVE",
                question=f"What is {indicator.description.lower().split('.')[0]}?",
                marks=5
            ))
        
        return assessments
    
    async def render_to_docx(
        self, 
        lessons: List[Lesson], 
        config: GenerationConfig,
        template: Template
    ) -> Document:
        """Render lesson plans to a DOCX document."""
        doc = Document()
        
        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)
        
        # Add title
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(f"TERM THREE WEEKLY LESSON NOTES – {config.class_level.value}")
        run.bold = True
        run.font.size = Pt(16)
        
        # Group lessons by week
        weeks_lessons = {}
        for lesson in lessons:
            week_num = (lesson.date - config.term_start_date).days // 7 + 1
            if week_num not in weeks_lessons:
                weeks_lessons[week_num] = []
            weeks_lessons[week_num].append(lesson)
        
        # Generate each week
        for week_num in sorted(weeks_lessons.keys()):
            week_lessons = weeks_lessons[week_num]
            
            # Add week header
            week_header = doc.add_paragraph()
            week_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = week_header.add_run(f"WEEK {week_num}")
            run.bold = True
            run.font.size = Pt(14)
            
            # Generate each lesson in the week
            for lesson in week_lessons:
                self._add_lesson_to_docx(doc, lesson, config, template)
                
                # Add page break after each lesson (except the last one)
                if lesson != week_lessons[-1]:
                    doc.add_page_break()
        
        return doc
    
    def _add_lesson_to_docx(
        self, 
        doc: Document, 
        lesson: Lesson, 
        config: GenerationConfig,
        template: Template
    ):
        """Add a single lesson to the DOCX document."""
        # Create metadata table
        meta_table = doc.add_table(rows=12, cols=2)
        meta_table.style = 'Table Grid'
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Add metadata rows
        meta_data = [
            ("Date:", lesson.date.strftime("%d/%m/%Y")),
            ("Period:", f"{lesson.period}{'st' if lesson.period == 1 else 'nd' if lesson.period == 2 else 'rd'} period"),
            ("Subject:", config.subject.value),
            ("Duration:", f"{lesson.duration_minutes}mins"),
            ("Strand:", lesson.strand.name if lesson.strand else ""),
            ("Class:", config.class_level.value),
            ("Class Size:", str(config.class_size)),
            ("Sub Strand:", lesson.sub_strand.name if lesson.sub_strand else ""),
            ("Content Standard:", lesson.content_standard.description if lesson.content_standard else ""),
            ("Indicator:", lesson.indicator.description if lesson.indicator else ""),
            ("Lesson:", f"{lesson.lesson_number} of {config.lessons_per_week}"),
            ("Performance Indicator:", f"Learners can {lesson.indicator.description.lower()}" if lesson.indicator else "")
        ]
        
        for i, (label, value) in enumerate(meta_data):
            row = meta_table.rows[i]
            
            # Label cell
            cell_label = row.cells[0]
            cell_label.text = label
            for paragraph in cell_label.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
            
            # Value cell
            cell_value = row.cells[1]
            cell_value.text = value
            
            # Set column widths
            row.cells[0].width = Inches(2.0)
            row.cells[1].width = Inches(4.5)
        
        # Add spacing
        doc.add_paragraph()
        
        # Add activities table
        activities_table = doc.add_table(rows=len(lesson.activities) + 1, cols=3)
        activities_table.style = 'Table Grid'
        activities_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header row
        header_row = activities_table.rows[0]
        headers = ["Phase/Duration", "Learners Activities", "Resources"]
        for i, header_text in enumerate(headers):
            cell = header_row.cells[i]
            cell.text = header_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        # Activity rows
        for i, activity in enumerate(lesson.activities):
            row = activities_table.rows[i + 1]
            row.cells[0].text = f"{activity.phase} ({activity.duration_minutes} mins)"
            row.cells[1].text = activity.description
            row.cells[2].text = ", ".join([r.name for r in activity.resources]) if activity.resources else ""
            
            # Set column widths
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(4.0)
            row.cells[2].width = Inches(1.5)
        
        # Add spacing
        doc.add_paragraph()
        
        # Add assessment and homework
        if lesson.assessments:
            p = doc.add_paragraph()
            run_label = p.add_run("Assessment: ")
            run_label.bold = True
            p.add_run(lesson.assessments[0].question)
        
        # Add homework placeholder
        p = doc.add_paragraph()
        run_label = p.add_run("Homework: ")
        run_label.bold = True
        p.add_run("Complete the exercises in your exercise book")